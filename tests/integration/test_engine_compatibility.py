"""The engine against the finite workflow and the downstream stages (revised phase 6).

* §57 - the same sheets read by the existing finite path (*Add Folder ->
  Process All*, driven headlessly through exactly the Scan stage's sequence)
  and by the engine produce the same durable outcome;
* §58 - batch partitioning (one unit, many units, arrivals over time, a restart
  between units) never changes the examination's semantics;
* §39 / §40 - Attendance and scoring over the session's effective set agree;
* §65 / §17 - rejection and *Reprocess All* (supersession) count each sheet once,
  however many rows and attempts history holds.
"""

from __future__ import annotations

from collections.abc import Iterator, Sequence
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import select
from tests.engine_rig import (
    OPTIONS,
    EngineRig,
    KillAt,
    digest,
    durable_view,
    readable_sheets,
)

from omr_scanner.database.models import BatchScan
from omr_scanner.domain.processing import UnitPolicy
from omr_scanner.domain.scan_lifecycle import RejectionReason
from omr_scanner.services import (
    BatchOptions,
    BatchRecorder,
    batch_store,
    candidate_import,
    create_project,
    intake,
    process_batch,
    project_sets,
    reconciliation_store,
    scan_lifecycle,
    scan_recovery,
    scan_sessions,
    scoring_store,
    session_population,
    set_attendance,
)
from omr_scanner.services.answer_key import plan_for, read_key

SHEETS = 40
OPERATOR = "Exam Office"


@pytest.fixture
def projects(workspace) -> Iterator[list]:
    opened: list = []
    yield opened
    for project in opened:
        project.close()


def new_rig(projects: list, workspace: Path, name: str) -> EngineRig:
    project = create_project(workspace, name)
    projects.append(project)
    return EngineRig(project)


def run_finite(rig: EngineRig, sheets: Sequence[bytes]) -> str:
    """*Add Folder -> Process All*, headless: exactly the Scan stage's steps.

    ``ScanPage._start_batch`` registers the batch (``start_batch``), marks the
    run's sheets queued and the batch running, records through a
    ``BatchRecorder`` with the template (the work unit); ``BatchWorker.run``
    first records the manual intake source and links exact duplicates; at the
    end ``_generate_conflicts`` completes the batch-scope review state and
    ``_settle_batch_state`` finalises.
    """
    folder = rig.project.project.layout.scans_original_dir / "finite"
    folder.mkdir(parents=True, exist_ok=True)
    paths = []
    for index, data in enumerate(sheets):
        path = folder / f"{index:06d}.png"
        path.write_bytes(data)
        paths.append(path)
    database = rig.database
    identity = batch_store.BatchIdentity.of(rig.template)
    batch_id = scan_sessions.start_batch(
        database, paths, identity=identity, scan_session_id=rig.session_id,
        settings={scan_recovery.WORK_UNIT_SETTING: True}, started_by=OPERATOR,
    )
    batch_store.mark_queued(database, batch_id, paths)
    batch_store.set_batch_status(database, batch_id, batch_store.BatchStatus.RUNNING)
    recorder = BatchRecorder(
        database=database, batch_id=batch_id, template=rig.template, commit_when_idle=True
    )
    intake.record_manual_batch(database, batch_id)
    linked = scan_lifecycle.link_exact_duplicates(database, batch_id, paths)
    intake.mirror_duplicates(database, batch_id)
    skipped = {str(item.path) for item in linked}
    run = [path for path in paths if str(path) not in skipped]
    process_batch(
        run, rig.template, options=BatchOptions(recognition=OPTIONS),
        on_result=recorder.record, workers=1,
    )
    assert recorder.flush()
    scan_recovery.complete_batch_review_state(database, batch_id)
    batch_store.finalise_batch(database, batch_id)
    return batch_id


def run_engine(
    rig: EngineRig,
    sheets: Sequence[bytes],
    *,
    unit: int,
    sources: str = "abc",
    waves: int = 1,
    kill_between_units: bool = False,
) -> None:
    """The engine over the same sheets, spread round robin over ``sources``."""
    for name in sources:
        rig.source(name)
    per_wave = -(-len(sheets) // waves)
    for wave in range(waves):
        chunk = list(enumerate(sheets))[wave * per_wave:(wave + 1) * per_wave]
        for index, data in chunk:
            rig.write(sources[index % len(sources)], [(f"{index:06d}.png", data)])
        hooks = KillAt("finalised", 1) if kill_between_units and wave == 0 else None
        engine = rig.new_engine(unit_policy=UnitPolicy(max_unit_size=unit, trickle_seconds=0),
                                hooks=hooks)
        rig.make_ready()
        if hooks is not None:
            rig.run_until_killed()
            engine = rig.new_engine(unit_policy=UnitPolicy(max_unit_size=unit,
                                                           trickle_seconds=0))
            rig.make_ready()
        rig.run()
        engine.shutdown()


def _differences(actual: dict, expected: dict) -> list[str]:
    """Readable per-sheet differences between two durable views (for the assertion message)."""
    import json

    found: list[str] = []
    for key in sorted(set(actual["results"]) | set(expected["results"])):
        left, right = actual["results"].get(key), expected["results"].get(key)
        if left == right:
            continue
        if left is None or right is None:
            found.append(f"{key[:8]}: only in {'expected' if left is None else 'actual'}")
            continue
        if left[:5] != right[:5]:
            found.append(f"{key[:8]}: {left[:5]} != {right[:5]}")
        a, b = json.loads(left[5]), json.loads(right[5])
        found += [f"{key[:8]}.{field}" for field in sorted(set(a) | set(b)) if a.get(field) != b.get(field)]
    for name in ("effective", "conflicts", "unread_duplicates", "pending"):
        if actual[name] != expected[name]:
            found.append(f"{name}: {sorted(set(map(str, actual[name])) ^ set(map(str, expected[name])))[:4]}")
    return found


@pytest.fixture(scope="module")
def sheets() -> tuple[bytes, ...]:
    return readable_sheets(48)[:SHEETS]


class TestFiniteEquivalence:
    def test_engine_matches_add_folder_process_all(self, projects, workspace, sheets):
        finite = new_rig(projects, workspace, "Finite")
        run_finite(finite, sheets)
        expected = durable_view(finite.database, finite.session_id)
        engine = new_rig(projects, workspace, "Engine")
        run_engine(engine, sheets, unit=8)
        actual = durable_view(engine.database, engine.session_id)
        assert _differences(actual, expected) == []
        assert actual == expected
        # Every unique sheet was read exactly once in both.
        assert set(engine.recognise.calls.values()) == {1}

    def test_partitioning_never_changes_the_outcome(self, projects, workspace, sheets):
        views = {}
        for name, options in {
            "one_unit": {"unit": SHEETS, "sources": "a"},
            "ten_units": {"unit": 4, "sources": "a"},
            "three_sources": {"unit": 5, "sources": "abc"},
            "arrivals": {"unit": 6, "sources": "ab", "waves": 4},
            "restart_between": {"unit": 7, "sources": "ab", "waves": 2,
                                "kill_between_units": True},
        }.items():
            rig = new_rig(projects, workspace, name)
            run_engine(rig, sheets, **options)  # type: ignore[arg-type]
            views[name] = durable_view(rig.database, rig.session_id)
        first = views["one_unit"]
        for name, view in views.items():
            assert view == first, name


def _roster(path: Path, candidates: Sequence[str]) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for column, header in enumerate(("Sl.No.", "Roll No.", "Name", "Total", "Merit"), start=1):
        sheet.cell(row=3, column=column, value=header)
    for offset, candidate in enumerate(candidates):
        sheet.cell(row=4 + offset, column=1, value=offset + 1)
        sheet.cell(row=4 + offset, column=2, value=candidate)
        sheet.cell(row=4 + offset, column=3, value=f"CANDIDATE {candidate}")
    workbook.save(path)
    workbook.close()
    return path


def downstream(rig: EngineRig, tmp: Path, candidates: dict[str, list[str]]) -> dict:
    """Attendance and scores for every set, by sheet content (ids differ per project)."""
    database = rig.database
    with database.session() as session:
        content = dict(session.execute(select(BatchScan.scan_id, BatchScan.content_sha256)).all())
    key_batch = session_population.session_population(database, rig.session_id).key_batch_id
    plan = plan_for(rig.template)
    out: dict = {}
    for code, rolls in sorted(candidates.items()):
        set_id = project_sets.add_set(database, code, f"Paper {code}").set_id
        path = _roster(tmp / f"{rig.project.root.name}-{code}.xlsx", rolls)
        assignment = set_attendance.assign_attendance_workbook(
            database, set_id, path, candidate_import.read_roster(path), imported_by=OPERATOR
        )
        stored = scoring_store.save_key(
            database, read_key("A" * plan.question_count, plan, code).to_key()
        )
        scoring_store.verify_key(database, stored.key_id, verified_by=OPERATOR)
        roster = assignment.roster_id
        reconciliation_store.reconcile_batch(database, roster, key_batch)
        out[f"attendance-{code}"] = sorted(
            (entry.candidate_id, tuple(sorted(content[v.script.scan_id] for v in entry.scripts)))
            for entry in reconciliation_store.list_entries(database, roster, key_batch)
        )
        scoring_store.score_batch(database, roster, key_batch, rig.template, computed_by=OPERATOR)
        out[f"results-{code}"] = sorted(
            (item.candidate_id, str(item.status), item.final_score,
             content.get(item.scan_id) if item.scan_id is not None else None)
            for item in scoring_store.list_results(database, roster, key_batch, rig.template)
        )
    return out


class TestDownstream:
    def test_attendance_and_scores_match_the_finite_run(self, projects, workspace, sheets, tmp_path):
        finite = new_rig(projects, workspace, "FiniteDown")
        run_finite(finite, sheets)
        identifiers = session_population.effective_identifiers(
            finite.database,
            session_population.session_population(finite.database, finite.session_id),
        )
        codes = session_population.effective_set_codes(
            finite.database,
            session_population.session_population(finite.database, finite.session_id),
        )
        candidates: dict[str, list[str]] = {}
        for scan_id, item in identifiers.items():
            code = codes.get(scan_id)
            if item.value and "?" not in item.value and code is not None and code.value:
                candidates.setdefault(code.value, []).append(item.value)
        candidates = {code: sorted(set(rolls)) + ["999999"] for code, rolls in candidates.items()}
        assert candidates, "the dataset produced no readable candidate"
        expected = downstream(finite, tmp_path, candidates)

        engine = new_rig(projects, workspace, "EngineDown")
        run_engine(engine, sheets, unit=6, sources="abc", waves=2, kill_between_units=True)
        assert downstream(engine, tmp_path, candidates) == expected


class TestEffectiveMembership:
    def test_rejection_and_reprocess_count_each_sheet_once(self, projects, workspace, sheets):
        rig = new_rig(projects, workspace, "Effective")
        run_engine(rig, sheets[:20], unit=10, sources="a")
        population = session_population.session_population(rig.database, rig.session_id)
        effective = len(population.effective)
        assert effective > 0
        units = [item.batch_id for item in scan_sessions.batches_of(rig.database, rig.session_id)]
        assert len(units) == 2

        # A rejection takes a sheet out; a restart does not bring it back.
        some = sorted(population.effective)[0]
        scan_lifecycle.reject_scan(rig.database, some, reviewer=OPERATOR,
                                   reason=RejectionReason.FOLDED)
        rig.new_engine()
        rig.run()
        rig.engine.shutdown()  # type: ignore[union-attr]
        assert len(session_population.session_population(
            rig.database, rig.session_id).effective) == effective - 1

        # Reprocess All of the second unit: a superseding batch, read by the
        # engine; the superseded original stays, and counts nowhere.
        reprocess = scan_sessions.start_reprocess_batch(
            rig.database, units[1], identity=batch_store.BatchIdentity.of(rig.template),
            settings={scan_recovery.WORK_UNIT_SETTING: True}, started_by=OPERATOR,
        )
        calls_before = sum(rig.recognise.calls.values())
        for _restart in range(2):
            engine = rig.new_engine()
            rig.run()
            engine.shutdown()
        reread = sum(rig.recognise.calls.values()) - calls_before
        with rig.database.session() as session:
            rows = session.scalars(
                select(BatchScan).where(BatchScan.batch_id == reprocess)
            ).all()
        assert reread == sum(1 for row in rows if row.status != "duplicate")
        after = session_population.session_population(rig.database, rig.session_id)
        assert len(after.effective) == effective - 1
        assert units[1] not in after.live_batch_ids
        with rig.database.session() as session:
            kept = session.scalars(select(BatchScan).where(BatchScan.batch_id == units[1])).all()
        assert kept, "superseded rows are retained"
        assert not (set(after.effective) & {row.scan_id for row in kept})
        # History rows grew; the effective count did not.
        assert session_population.count_in_session(rig.database, rig.session_id) > effective


def test_source_outage_does_not_stop_registered_work(projects, workspace, sheets):
    """§48: registered project copies are read while their source is offline."""
    rig = new_rig(projects, workspace, "Outage")
    rig.source("a")
    rig.source("b")
    rig.write("a", [(f"{i:06d}.png", data) for i, data in enumerate(sheets[:6])])
    rig.write("b", [(f"{i:06d}.png", data) for i, data in enumerate(sheets[6:12])])
    engine = rig.new_engine(unit_policy=UnitPolicy(max_unit_size=6, trickle_seconds=0))
    rig.make_ready()
    engine.form_units()
    engine.form_units()
    rig.fs.unreachable.add(rig.root("a"))
    rig.write("b", [("000099.png", sheets[12])])
    rig.make_ready()
    rig.run()
    status = engine.status()
    assert status.caught_up
    assert status.pending == 0
    reached = {item.source_id: item.reachability.value for item in intake.list_sources(rig.database)}
    assert reached[rig.sources["a"]] == "unreachable"
    view = durable_view(rig.database, rig.session_id)
    assert {digest(item) for item in sheets[:13]} <= set(view["results"]) | set(
        view["unread_duplicates"]
    )
    engine.shutdown()
