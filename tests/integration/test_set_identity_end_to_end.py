"""Set identity end to end (phase 0.1.1-A, ACCEPTANCE_CRITERIA A4; defect 5).

Two examinations driven through the real services, stage by stage:

    definition -> recognition (stored readings) -> Resolve -> Attendance ->
    Answer Key -> scoring -> Results -> Reports (and the CSV export)

* **Case A - case-insensitive identity.** Set ``A`` is defined; every sheet
  reads ``a``; the answer key is typed as ``a``. Before this phase the
  sheets were "undefined" on Resolve and the key could not be found in
  Results (defect 5).
* **Case B - physical mark.** Set ``10`` is printed on the sheet as ``A``;
  every sheet reads ``A``. Every stage after Resolve must say Set 10, while
  the raw ``A`` stays on the scan as provenance.

"Recognition" is the stored reading a batch run leaves in ``batch_scan`` -
the same rows the real engine writes - registered directly, as the other
multi-set integration tests do. A real recognition run of a mapped sheet is
``tests/integration/test_set_identity_synthetic.py``.
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import pytest
from tests.conftest import build_answer_sheet_template

from omr_scanner.database.models import BatchScan, ScanBatch
from omr_scanner.domain.review import ConflictType
from omr_scanner.services import (
    batch_store,
    candidate_import,
    create_project,
    project_sets,
    reconciliation_store,
    report_store,
    review_store,
    scoring_store,
    set_attendance,
    set_identity,
)
from omr_scanner.services.answer_key import SetCodeCheck, check_sheet_set, plan_for, read_key
from omr_scanner.services.batch_processor import ProcessedScan
from omr_scanner.services.recognition_models import (
    AnswerView,
    FieldView,
    RecognitionOutcome,
    RegistrationStatus,
    ScanResult,
)
from omr_scanner.services.scan_export import LOGICAL_SET_COLUMN, render_scan_results

OPERATOR = "Set Identity E2E"
ROLLS = ("30001", "30002", "30003")
CORRECT = {"30001": 5, "30002": 3, "30003": 0}


def _attendance(path: Path, rolls: tuple[str, ...]) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for column, text in enumerate(("Sl.No.", "Roll No.", "Name", "Total"), start=1):
        sheet.cell(row=1, column=column, value=text)
    for offset, roll in enumerate(rolls):
        sheet.cell(row=2 + offset, column=1, value=offset + 1)
        sheet.cell(row=2 + offset, column=2, value=roll)
        sheet.cell(row=2 + offset, column=3, value=f"CANDIDATE {roll}")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def _results(folder: Path, set_read: str, numbers) -> list[ScanResult]:
    ordered = list(numbers)
    return [
        ScanResult(
            source_path=folder / f"sheet_{index:03d}.png",
            outcome=RecognitionOutcome.COMPLETE,
            registration=RegistrationStatus.REGISTERED,
            fields=(
                FieldView(
                    zone_id="roll_number", label="Roll", field_type="numeric",
                    value=roll, status="complete", needs_review=False, characters=(),
                ),
                FieldView(
                    zone_id="set_code", label="Set", field_type="set_code",
                    value=set_read, status="complete", needs_review=False, characters=(),
                ),
            ),
            answers=tuple(
                AnswerView(
                    number=number, zone_id="q",
                    value="A" if position < CORRECT[roll] else "B",
                    status="resolved", needs_review=False,
                    top_fill=0.9, margin=0.4, confidence=0.9,
                )
                for position, number in enumerate(ordered)
            ),
            identifier_zone_id="roll_number", set_code_zone_id="set_code",
        )
        for index, roll in enumerate(ROLLS)
    ]


def _register(database, folder: Path, results: list[ScanResult]) -> str:
    now = datetime.now(UTC)
    batch_id = batch_store.new_batch_id()
    with database.session() as session:
        session.add(
            ScanBatch(
                batch_id=batch_id, created_at=now, updated_at=now,
                source_folder=str(folder), status="completed", total_scans=len(results),
            )
        )
        session.flush()
        for index, result in enumerate(results):
            session.add(
                BatchScan(
                    batch_id=batch_id, batch_index=index,
                    source_path=str(result.source_path), filename=result.source_path.name,
                    status="completed", identifier_value=result.identifier_value,
                    set_code_value=result.set_code_value,
                    result_json=json.dumps(result.to_dict()),
                )
            )
    return batch_id


@dataclass
class Outcome:
    """What every stage said about the examination."""

    set_code: str
    undefined_conflicts: int
    effective: dict[int, tuple[str, str]]
    in_set: int
    other_sets: dict[str, int]
    key_found: bool
    results: dict[str, tuple[str, str, str]]
    by_set: dict[str, int]
    report_ok: bool
    report_status: str
    raw_readings: list[str]
    csv_header: list[str]
    csv_rows: list[dict[str, str]]


def _run(
    tmp_path: Path, *, code: str, mark: str, sheet_reads: str, key_typed_as: str
) -> Outcome:
    template = build_answer_sheet_template()
    plan = plan_for(template)
    session = create_project(tmp_path / "workspace", "Set Identity E2E")
    try:
        database = session.database
        # 1. Definition (Project Configuration -> Sets).
        exam_set = project_sets.add_set(
            database, code, "The examination's only set", physical_mark=mark, template=template
        )
        # 2. Recognition: what the engine read off each sheet.
        results = _results(tmp_path / "scans", sheet_reads, plan.numbers)
        batch_id = _register(database, tmp_path / "scans", results)
        # 3. Resolve: is the set read off each sheet one of the project's?
        undefined = review_store.sync_undefined_set_codes(database, batch_id)
        effective = {
            scan_id: (item.value, item.as_read)
            for scan_id, item in review_store.effective_set_codes(database, batch_id).items()
        }
        # 4. Attendance: the set's own roster, reconciled against its scripts.
        attendance = _attendance(tmp_path / "attendance.xlsx", ROLLS)
        assignment = set_attendance.assign_attendance_workbook(
            database,
            exam_set.set_id,
            attendance,
            candidate_import.read_roster(attendance),
            imported_by=OPERATOR,
        )
        assert assignment.template_adopted, assignment.template_blocker
        reconciliation_store.reconcile_batch(database, assignment.roster_id, batch_id)
        scope = reconciliation_store.script_scope(database, assignment.roster_id, batch_id)
        # 5. Answer Key: the solution sheet's mark agrees with the chosen set,
        # and the key is stored as the operator typed it.
        verdict = check_sheet_set(
            template,
            selected=exam_set.code,
            read=sheet_reads,
            identity=set_identity.load(database),
        )
        assert verdict.check is SetCodeCheck.MATCH, verdict.message
        stored = scoring_store.save_key(
            database, read_key("A" * plan.question_count, plan, key_typed_as).to_key()
        )
        scoring_store.verify_key(database, stored.key_id, verified_by=OPERATOR)
        key_found = scoring_store.verified_key(database, exam_set.code) is not None
        # 6. Scoring and 7. Results.
        counts = scoring_store.score_batch(
            database, assignment.roster_id, batch_id, template, computed_by=OPERATOR
        )
        listed = {
            item.candidate_id: (item.set_code, item.status.value, str(item.final_score))
            for item in scoring_store.list_results(
                database, assignment.roster_id, batch_id, template
            )
        }
        # 8. Reports: the set's own result workbook, final.
        report = report_store.generate_for_set(
            database, exam_set.set_id, batch_id, template,
            project_name="SetIdentityE2E", output_dir=tmp_path / "exports",
            computed_by=OPERATOR, final=True,
        )
        # The scan-results CSV, with the logical set beside the set as read.
        processed = [
            ProcessedScan(result=result, output_name=result.source_path.name)
            for result in results
        ]
        text = render_scan_results(
            processed, template, identity=set_identity.load(database)
        )
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
        with database.session() as orm:
            raw = sorted(row.set_code_value for row in orm.query(BatchScan).all())
        return Outcome(
            set_code=exam_set.code,
            undefined_conflicts=undefined,
            effective=effective,
            in_set=scope.in_set,
            other_sets=dict(scope.other_sets),
            key_found=key_found,
            results=listed,
            by_set=dict(counts.by_set),
            report_ok=report.ok,
            report_status=report.status,
            raw_readings=raw,
            csv_header=list(reader.fieldnames or ()),
            csv_rows=rows,
        )
    finally:
        session.close()


@pytest.fixture
def case_a(tmp_path: Path) -> Outcome:
    return _run(tmp_path / "case_a", code="A", mark="", sheet_reads="a", key_typed_as="a")


@pytest.fixture
def case_b(tmp_path: Path) -> Outcome:
    return _run(tmp_path / "case_b", code="10", mark="A", sheet_reads="A", key_typed_as="10")


class TestCaseALowerCaseReadingOfSetA:
    def test_resolve_does_not_call_it_an_undefined_set(self, case_a: Outcome) -> None:
        assert case_a.undefined_conflicts == 0
        assert set(case_a.effective.values()) == {("A", "a")}

    def test_attendance_places_every_script_in_set_a(self, case_a: Outcome) -> None:
        assert case_a.in_set == len(ROLLS)
        assert case_a.other_sets == {}

    def test_the_key_typed_as_a_is_found_for_set_a(self, case_a: Outcome) -> None:
        assert case_a.key_found

    def test_results_mark_everyone_against_set_a(self, case_a: Outcome) -> None:
        statuses = {roll: status for roll, (_set, status, _mark) in case_a.results.items()}
        assert statuses == dict.fromkeys(ROLLS, "scored")
        assert {item[0] for item in case_a.results.values()} == {"A"}
        assert case_a.by_set == {"A": len(ROLLS)}

    def test_the_report_is_generated(self, case_a: Outcome) -> None:
        assert case_a.report_ok, case_a.report_status

    def test_the_raw_reading_is_kept(self, case_a: Outcome) -> None:
        assert case_a.raw_readings == ["a"] * len(ROLLS)

    def test_an_unmapped_project_exports_exactly_the_old_columns(self, case_a: Outcome) -> None:
        assert LOGICAL_SET_COLUMN not in case_a.csv_header


class TestCaseBSet10PrintedAsA:
    def test_resolve_translates_the_mark_and_keeps_it(self, case_b: Outcome) -> None:
        assert case_b.undefined_conflicts == 0
        assert set(case_b.effective.values()) == {("10", "A")}

    def test_attendance_places_every_script_in_set_10(self, case_b: Outcome) -> None:
        assert case_b.in_set == len(ROLLS)
        assert case_b.other_sets == {}

    def test_the_key_belongs_to_logical_set_10(self, case_b: Outcome) -> None:
        assert case_b.key_found

    def test_results_say_set_10(self, case_b: Outcome) -> None:
        assert {item[0] for item in case_b.results.values()} == {"10"}
        assert {item[1] for item in case_b.results.values()} == {"scored"}
        assert case_b.by_set == {"10": len(ROLLS)}

    def test_marks_are_the_same_as_an_unmapped_examination(
        self, case_a: Outcome, case_b: Outcome
    ) -> None:
        """The mapping changes which set a sheet is, never what it scores."""
        assert {roll: item[2] for roll, item in case_b.results.items()} == {
            roll: item[2] for roll, item in case_a.results.items()
        }

    def test_the_report_is_generated(self, case_b: Outcome) -> None:
        assert case_b.report_ok, case_b.report_status

    def test_the_raw_a_is_retained_as_provenance(self, case_b: Outcome) -> None:
        assert case_b.raw_readings == ["A"] * len(ROLLS)

    def test_the_export_shows_set_as_read_and_set(self, case_b: Outcome) -> None:
        assert LOGICAL_SET_COLUMN in case_b.csv_header
        assert {(row["set_code"], row[LOGICAL_SET_COLUMN]) for row in case_b.csv_rows} == {
            ("A", "10")
        }


def test_a_reading_naming_no_set_is_still_undefined(tmp_path: Path) -> None:
    """The translation never invents a set: "D" stays "D" and is flagged."""
    template = build_answer_sheet_template()
    plan = plan_for(template)
    session = create_project(tmp_path / "workspace", "Undefined")
    try:
        database = session.database
        project_sets.add_set(database, "10", physical_mark="A")
        batch_id = _register(database, tmp_path, _results(tmp_path, "D", plan.numbers))
        assert review_store.sync_undefined_set_codes(database, batch_id) == len(ROLLS)
        effective = review_store.effective_set_codes(database, batch_id)
        assert {item.value for item in effective.values()} == {"D"}
        conflicts = review_store.list_conflicts(database, batch_id)
        assert {item.conflict_type for item in conflicts} == {ConflictType.SET_CODE_UNDEFINED}
        assert all("10 (printed A)" in item.observation.detail for item in conflicts)
    finally:
        session.close()
