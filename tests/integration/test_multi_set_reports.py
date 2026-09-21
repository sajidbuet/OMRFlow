"""A whole multi-Set examination, generated end to end (Part 2, §25, §26, §27).

Scope:
    One synthetic project carrying the brief's own example - *Recruitment
    Exam, Bangladesh Submarine Cable Regulatory Authority*, divided into Set
    10 (Electrical), Set 11 (Civil) and Set 12 (Mechanical) - built through
    the real services and generated through
    :func:`report_store.generate_for_set`. Every assertion reads the produced
    `.xlsx` back with openpyxl, because that file is what an examination
    office actually files.

Size:
    §27 specifies 500 / 350 / 725 candidates. Scoring and generating three
    cohorts that size takes minutes, which is too slow to run on every
    change, so the test that runs **by default** uses 40 / 30 / 50 - small
    enough to be quick, large enough to contain every case §27 names
    (overlapping rolls, absentees, a tie, a zero, a full mark, an unknown
    script, a duplicate script, an absent-with-script). The full
    500 / 350 / 725 project runs as a single ``stress``-marked test, excluded
    from the default run and invoked with ``-m stress``.

What this file deliberately does not cover:
    Workbook *formatting* preservation - merges, fonts, fills, borders,
    widths, heights, page setup, print titles, headers/footers and the logo -
    which ``tests/unit/test_meritwise_workbook.py`` already asserts in depth
    against a workbook built to contain every one of them. Here the subject
    is multi-Set correctness and isolation.

Why one batch per Set:
    A Set's scripts are its own scanning run, and reconciliation matches one
    batch against one roster. Giving each Set its own batch is both the
    realistic workflow and the arrangement that keeps "Set 10's result" a
    question with one unambiguous answer.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import TYPE_CHECKING

import openpyxl
import pytest
from tests.conftest import build_answer_sheet_template
from tests.report_fixtures import sha256_of

from omr_scanner.domain.reporting import ABSENT_MARK_DISPLAY
from omr_scanner.reporting import excel as rx
from omr_scanner.services import (
    batch_store,
    candidate_import,
    create_project,
    project_sets,
    reconciliation_store,
    report_store,
    scoring_store,
    set_attendance,
)
from omr_scanner.services.answer_key import plan_for, read_key
from omr_scanner.services.recognition_models import (
    AnswerView,
    FieldView,
    RecognitionOutcome,
    RegistrationStatus,
    ScanResult,
)

if TYPE_CHECKING:
    from pathlib import Path

    from omr_scanner.database import ProjectDatabase
    from omr_scanner.domain.exam_sets import ExamSet
    from omr_scanner.services import ProjectSession

OPERATOR = "Dr. Rahman"

EXAM_NAME = "Recruitment Exam, Bangladesh Submarine Cable Regulatory Authority"

HEADERS: tuple[str, ...] = ("Sl.No.", "Roll No.", "Name", "Total", "Merit")
HEADER_ROW = 3
ROLL_COLUMN = 2
MARKS_COLUMN = 4
MERIT_COLUMN = 5

DEFAULT_SIZES: dict[str, int] = {"10": 40, "11": 30, "12": 50}
STRESS_SIZES: dict[str, int] = {"10": 500, "11": 350, "12": 725}

SHARED_ROLL = "10001"
"""Registered in Set 10 *and* Set 11, with a different mark in each - the §19
case a set-blind generator would get wrong without anything looking amiss."""


# ----------------------------------------------------------------------
# Building the synthetic examination
# ----------------------------------------------------------------------
@dataclass
class SetPlan:
    """One Set's intended contents, before anything is written."""

    code: str
    description: str
    rolls: list[str]
    absent: set[str] = field(default_factory=set)
    """Recorded absent in the attendance workbook."""
    correct_by_roll: dict[str, int] = field(default_factory=dict)
    """How many questions each present candidate answers correctly.

    The mark itself is never predicted here - the scoring policy decides that
    - but a candidate who gets more right necessarily scores higher, and two
    candidates who get the same number right necessarily tie. That is enough
    to assert ordering and tie handling without duplicating the scorer."""
    extra_scripts: list[tuple[str, int]] = field(default_factory=list)
    """``(roll, correct_count)`` scripts beyond one per present candidate -
    duplicates, unknown candidates and absent-with-script cases."""


def build_plans(sizes: dict[str, int], question_count: int) -> list[SetPlan]:
    """Lay out three Sets containing every case §27 names.

    Marks arise from *how a script answered*, and are never predicted here:
    the answer key is all ``A``, so a script whose first ``n`` answers are
    ``A`` and whose rest are ``B`` gets exactly ``n`` right. A full mark, a
    zero and a deliberate tie are therefore stated as answer counts, and
    whatever the scoring policy turns them into is read back from the
    database when it is asserted.
    """
    plans: list[SetPlan] = []

    def spread(index: int) -> int:
        """A repeatable, varied number of correct answers."""
        return (index * 7) % (question_count + 1)

    # ---- Set 10: clean, and the owner of the shared roll ----
    ten_rolls = [SHARED_ROLL] + [f"10{n:03d}" for n in range(2, sizes["10"] + 1)]
    ten = SetPlan(
        code="10",
        description="Name of Post: Assistant Engineer (Electrical)",
        rolls=ten_rolls,
        absent={roll for index, roll in enumerate(ten_rolls, start=1) if index % 10 == 0},
    )
    for index, roll in enumerate(ten.rolls):
        if roll in ten.absent:
            continue
        if roll == SHARED_ROLL:
            ten.correct_by_roll[roll] = question_count       # full marks
        elif index == 1:
            ten.correct_by_roll[roll] = question_count       # ties with it
        elif index == 2:
            ten.correct_by_roll[roll] = 0                    # nothing right
        else:
            ten.correct_by_roll[roll] = spread(index)
    plans.append(ten)

    # ---- Set 11: clean, and registers the same roll as Set 10 ----
    eleven_rolls = [SHARED_ROLL] + [f"11{n:03d}" for n in range(2, sizes["11"] + 1)]
    eleven = SetPlan(
        code="11",
        description="Name of Post: Assistant Engineer (Civil)",
        rolls=eleven_rolls,
        absent={
            roll for index, roll in enumerate(eleven_rolls, start=1) if index % 10 == 0
        },
    )
    for index, roll in enumerate(eleven.rolls):
        if roll in eleven.absent:
            continue
        # The shared roll answers *differently* here - one correct against
        # Set 10's full marks - so the two Sets' marks for it genuinely
        # disagree, which is what makes the §19 assertion meaningful.
        eleven.correct_by_roll[roll] = 1 if roll == SHARED_ROLL else spread(index + 3)
    plans.append(eleven)

    # ---- Set 12: carries the exception cases ----
    twelve_rolls = [f"12{n:03d}" for n in range(1, sizes["12"] + 1)]
    twelve = SetPlan(
        code="12",
        description="Name of Post: Assistant Engineer (Mechanical)",
        rolls=twelve_rolls,
        absent={
            roll for index, roll in enumerate(twelve_rolls, start=1) if index % 10 == 0
        },
    )
    for index, roll in enumerate(twelve.rolls):
        if roll in twelve.absent:
            continue
        twelve.correct_by_roll[roll] = spread(index + 5)
    twelve.extra_scripts = [
        ("12002", question_count),   # duplicate script for a registered candidate
        ("99999", question_count),   # unknown: registered in no Set
        ("12010", question_count),   # absent-with-script (12010 is absent)
    ]
    plans.append(twelve)
    return plans


def write_attendance_workbook(path: Path, plan: SetPlan) -> Path:
    """Write one Set's attendance workbook, decorated the way a real one is."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = f"Attendance Set {plan.code}"
    sheet.cell(row=1, column=1, value="Bangladesh Submarine Cable Regulatory Authority")
    sheet.cell(row=2, column=1, value=plan.description)
    for column, text in enumerate(HEADERS, start=1):
        sheet.cell(row=HEADER_ROW, column=column, value=text)
    for offset, roll in enumerate(plan.rolls):
        row = HEADER_ROW + 1 + offset
        sheet.cell(row=row, column=1, value=offset + 1)
        sheet.cell(row=row, column=ROLL_COLUMN, value=roll)
        sheet.cell(row=row, column=3, value=f"CANDIDATE {roll}")
        if roll in plan.absent:
            sheet.cell(row=row, column=MARKS_COLUMN, value="ABSENT")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def register_batch(
    database: ProjectDatabase, tmp_path: Path, plan: SetPlan, numbers
) -> str:
    """Register this Set's own batch of scripts and return its batch id."""
    from omr_scanner.database.models import BatchScan, ScanBatch

    ordered_numbers = list(numbers)
    scripts: list[tuple[str, int]] = list(plan.correct_by_roll.items())
    scripts.extend(plan.extra_scripts)

    now = datetime.now(UTC)
    batch_id = batch_store.new_batch_id()
    with database.session() as session:
        session.add(
            ScanBatch(
                batch_id=batch_id, created_at=now, updated_at=now,
                source_folder=str(tmp_path), status="completed",
                total_scans=len(scripts),
            )
        )
        session.flush()
        for index, (roll, correct) in enumerate(scripts):
            name = f"set{plan.code}_{index:04d}.png"
            # The key is all "A": the first `correct` answers match it, the
            # rest do not, so the script earns exactly `correct` right.
            letters = {
                number: ("A" if position < correct else "B")
                for position, number in enumerate(ordered_numbers)
            }
            result = ScanResult(
                source_path=tmp_path / name,
                outcome=RecognitionOutcome.COMPLETE,
                registration=RegistrationStatus.REGISTERED,
                fields=(
                    FieldView(
                        zone_id="roll_number", label="Roll", field_type="numeric",
                        value=roll, status="complete", needs_review=False, characters=(),
                    ),
                    FieldView(
                        zone_id="set_code", label="Set", field_type="set_code",
                        value=plan.code, status="complete", needs_review=False,
                        characters=(),
                    ),
                ),
                answers=tuple(
                    AnswerView(
                        number=number, zone_id="q", value=letters[number],
                        status="resolved", needs_review=False,
                        top_fill=0.9, margin=0.4, confidence=0.9,
                    )
                    for number in ordered_numbers
                ),
                identifier_zone_id="roll_number", set_code_zone_id="set_code",
            )
            session.add(
                BatchScan(
                    batch_id=batch_id, batch_index=index,
                    source_path=str(tmp_path / name), filename=name,
                    status="completed", identifier_value=roll,
                    set_code_value=plan.code,
                    result_json=json.dumps(result.to_dict()),
                )
            )
    return batch_id


@dataclass
class BuiltSet:
    """One Set, fully built: defined, attended, scanned, keyed and scored."""

    plan: SetPlan
    exam_set: ExamSet
    attendance_path: Path
    roster_id: int
    batch_id: str


def build_examination(
    session: ProjectSession, tmp_path: Path, sizes: dict[str, int]
) -> tuple[list[BuiltSet], object]:
    """Build the whole synthetic examination and return its Sets."""
    template = build_answer_sheet_template()
    plan = plan_for(template)
    built: list[BuiltSet] = []

    for set_plan in build_plans(sizes, plan.question_count):
        exam_set = project_sets.add_set(
            session.database, set_plan.code, set_plan.description
        )
        attendance = write_attendance_workbook(
            tmp_path / f"set{set_plan.code}_attendance.xlsx", set_plan
        )
        validation = candidate_import.read_roster(attendance)
        assignment = set_attendance.assign_attendance_workbook(
            session.database, exam_set.set_id, attendance, validation,
            imported_by=OPERATOR,
        )
        assert assignment.template_adopted, assignment.template_blocker

        batch_id = register_batch(session.database, tmp_path, set_plan, plan.numbers)
        reconciliation_store.reconcile_batch(
            session.database, assignment.roster_id, batch_id
        )

        stored_key = scoring_store.save_key(
            session.database,
            read_key("A" * plan.question_count, plan, set_plan.code).to_key(),
        )
        scoring_store.verify_key(
            session.database, stored_key.key_id, verified_by=OPERATOR
        )
        scoring_store.score_batch(
            session.database, assignment.roster_id, batch_id, template,
            computed_by=OPERATOR,
        )
        built.append(
            BuiltSet(
                plan=set_plan, exam_set=exam_set, attendance_path=attendance,
                roster_id=assignment.roster_id, batch_id=batch_id,
            )
        )
    return built, template


# ----------------------------------------------------------------------
# Reading a generated workbook back
# ----------------------------------------------------------------------
def sheet_rolls(sheet, first_row: int) -> list[str]:
    """Every non-empty Roll No. on a sheet, from ``first_row`` down."""
    rolls: list[str] = []
    for row in range(first_row, sheet.max_row + 1):
        value = sheet.cell(row=row, column=ROLL_COLUMN).value
        if value not in (None, ""):
            rolls.append(str(value))
    return rolls


def marks_by_roll(sheet, first_row: int) -> dict[str, object]:
    """Roll -> whatever its marks cell holds."""
    found: dict[str, object] = {}
    for row in range(first_row, sheet.max_row + 1):
        roll = sheet.cell(row=row, column=ROLL_COLUMN).value
        if roll in (None, ""):
            continue
        found[str(roll)] = sheet.cell(row=row, column=MARKS_COLUMN).value
    return found


def stored_marks(
    database: ProjectDatabase, built: BuiltSet, template
) -> dict[str, float]:
    """What Phase 8 actually stored, read back - never recomputed here."""
    results = scoring_store.list_results(
        database, built.roster_id, built.batch_id, template
    )
    return {
        item.candidate_id: float(item.final_score)
        for item in results
        if item.final_score is not None
    }


def every_cell_text(path: Path) -> set[str]:
    """Every string in the whole workbook, for the contamination check."""
    workbook = openpyxl.load_workbook(path)
    try:
        found: set[str] = set()
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if value is not None:
                        found.add(str(value))
        return found
    finally:
        workbook.close()


# ----------------------------------------------------------------------
# Fixtures
# ----------------------------------------------------------------------
@pytest.fixture
def examination(workspace: Path, tmp_path: Path):
    """The default-size synthetic examination, with its session open."""
    session = create_project(workspace, "BSCRA Recruitment", exam_name=EXAM_NAME)
    try:
        built, template = build_examination(session, tmp_path, DEFAULT_SIZES)
        yield session, built, template
    finally:
        if not session.is_closed:
            session.close()


@pytest.fixture
def generated(examination, tmp_path: Path):
    """Every Set generated, keyed by set code."""
    session, built, template = examination
    exports = tmp_path / "exports"
    outcomes = {}
    for item in built:
        # Set 12 deliberately carries unresolved exceptions, so a *final*
        # export of it is correctly refused - §8's rule that an incomplete
        # report is never presented as finalised. A preview still produces
        # the workbook, which is what these assertions read.
        outcomes[item.plan.code] = report_store.generate_for_set(
            session.database, item.exam_set.set_id, item.batch_id, template,
            project_name="BSCRA Recruitment", output_dir=exports,
            computed_by=OPERATOR, final=item.plan.code != "12",
        )
    return session, built, template, outcomes


# ----------------------------------------------------------------------
# §27 - three independent workbooks
# ----------------------------------------------------------------------
class TestEachSetProducesItsOwnWorkbook:
    def test_every_set_generates(self, generated):
        _session, _built, _template, outcomes = generated
        for code, outcome in outcomes.items():
            assert outcome.output_path is not None, (code, outcome.warnings)
            assert outcome.output_path.is_file()

    def test_the_clean_sets_pass_a_final_export(self, generated):
        _session, _built, _template, outcomes = generated
        assert outcomes["10"].ok, outcomes["10"].warnings
        assert outcomes["11"].ok, outcomes["11"].warnings

    def test_a_set_with_unresolved_exceptions_is_refused_a_final_export(
        self, examination, tmp_path: Path
    ):
        """§8/§15: an incomplete report is never presented as finalised."""
        session, built, template = examination
        twelve = next(item for item in built if item.plan.code == "12")
        outcome = report_store.generate_for_set(
            session.database, twelve.exam_set.set_id, twelve.batch_id, template,
            project_name="BSCRA Recruitment", output_dir=tmp_path / "final",
            final=True,
        )
        assert outcome.status == "blocked"
        assert outcome.output_path is None
        assert outcome.warnings

    def test_the_three_workbooks_are_three_separate_files(self, generated):
        _session, _built, _template, outcomes = generated
        paths = {outcome.output_path for outcome in outcomes.values()}
        assert len(paths) == 3

    def test_each_workbook_names_its_own_set(self, generated):
        _session, _built, _template, outcomes = generated
        for code, outcome in outcomes.items():
            assert outcome.set_code == code
            # `default_output_stem` builds "<project>_Set<code>_Result".
            assert f"Set{code}" in outcome.output_path.name


# ----------------------------------------------------------------------
# §25 Rollwise
# ----------------------------------------------------------------------
class TestRollwise:
    def _sheet(self, session, built_item, outcome) -> tuple[object, object]:
        workbook = openpyxl.load_workbook(outcome.output_path)
        association = report_store.resolve_set_sources(
            session.database, built_item.exam_set.set_id
        ).association
        return workbook, workbook[association.sheet_name]

    def test_every_registered_candidate_is_present_including_absentees(self, generated):
        session, built, _template, outcomes = generated
        for item in built:
            workbook, sheet = self._sheet(session, item, outcomes[item.plan.code])
            try:
                rolls = sheet_rolls(sheet, HEADER_ROW + 1)
                assert rolls == item.plan.rolls
                assert item.plan.absent <= set(rolls)
            finally:
                workbook.close()

    def test_marks_match_the_stored_results(self, generated):
        """Read back from `candidate_result`, never recomputed in the test."""
        session, built, template, outcomes = generated
        for item in built:
            expected = stored_marks(session.database, item, template)
            workbook, sheet = self._sheet(session, item, outcomes[item.plan.code])
            try:
                written = marks_by_roll(sheet, HEADER_ROW + 1)
            finally:
                workbook.close()
            for roll, mark in expected.items():
                if roll in item.plan.absent:
                    continue
                assert written[roll] == pytest.approx(mark), (item.plan.code, roll)

    def test_absent_candidates_show_the_projects_own_absence_marker(self, generated):
        session, built, _template, outcomes = generated
        for item in built:
            workbook, sheet = self._sheet(session, item, outcomes[item.plan.code])
            try:
                written = marks_by_roll(sheet, HEADER_ROW + 1)
            finally:
                workbook.close()
            for roll in item.plan.absent:
                assert written[roll] == ABSENT_MARK_DISPLAY, (item.plan.code, roll)

    def test_no_absent_candidate_is_given_a_number(self, generated):
        session, built, _template, outcomes = generated
        for item in built:
            workbook, sheet = self._sheet(session, item, outcomes[item.plan.code])
            try:
                written = marks_by_roll(sheet, HEADER_ROW + 1)
            finally:
                workbook.close()
            for roll in item.plan.absent:
                assert not isinstance(written[roll], (int, float))

    def test_the_candidate_order_is_the_offices_own(self, generated):
        """Rollwise never reorders: the attendance sheet's order is the record."""
        session, built, _template, outcomes = generated
        for item in built:
            workbook, sheet = self._sheet(session, item, outcomes[item.plan.code])
            try:
                assert sheet_rolls(sheet, HEADER_ROW + 1) == item.plan.rolls
            finally:
                workbook.close()


# ----------------------------------------------------------------------
# §25 Meritwise
# ----------------------------------------------------------------------
class TestMeritwise:
    def test_the_sheet_exists_and_is_named_meritwise(self, generated):
        _session, built, _template, outcomes = generated
        for item in built:
            workbook = openpyxl.load_workbook(outcomes[item.plan.code].output_path)
            try:
                assert rx.MERITWISE_SHEET_NAME in workbook.sheetnames
                assert rx.MERITWISE_SHEET_NAME == "meritwise"
            finally:
                workbook.close()

    def test_absent_candidates_are_removed(self, generated):
        _session, built, _template, outcomes = generated
        for item in built:
            workbook = openpyxl.load_workbook(outcomes[item.plan.code].output_path)
            try:
                rolls = set(
                    sheet_rolls(workbook[rx.MERITWISE_SHEET_NAME], HEADER_ROW + 1)
                )
            finally:
                workbook.close()
            assert not (rolls & item.plan.absent), item.plan.code

    def test_every_scored_present_candidate_remains(self, generated):
        session, built, template, outcomes = generated
        for item in built:
            expected = {
                roll for roll in stored_marks(session.database, item, template)
                if roll not in item.plan.absent and roll in set(item.plan.rolls)
            }
            workbook = openpyxl.load_workbook(outcomes[item.plan.code].output_path)
            try:
                rolls = set(
                    sheet_rolls(workbook[rx.MERITWISE_SHEET_NAME], HEADER_ROW + 1)
                )
            finally:
                workbook.close()
            assert expected <= rolls, item.plan.code

    def test_rows_are_ordered_by_descending_mark(self, generated):
        _session, built, _template, outcomes = generated
        for item in built:
            workbook = openpyxl.load_workbook(outcomes[item.plan.code].output_path)
            try:
                sheet = workbook[rx.MERITWISE_SHEET_NAME]
                marks = [
                    sheet.cell(row=row, column=MARKS_COLUMN).value
                    for row in range(HEADER_ROW + 1, sheet.max_row + 1)
                    if sheet.cell(row=row, column=ROLL_COLUMN).value not in (None, "")
                ]
            finally:
                workbook.close()
            numeric = [value for value in marks if isinstance(value, (int, float))]
            assert numeric == sorted(numeric, reverse=True), item.plan.code

    def test_a_tie_is_broken_by_ascending_roll(self, generated):
        """Tied candidates keep the same mark; only their row order is decided."""
        _session, built, _template, outcomes = generated
        ten = next(item for item in built if item.plan.code == "10")
        workbook = openpyxl.load_workbook(outcomes["10"].output_path)
        try:
            sheet = workbook[rx.MERITWISE_SHEET_NAME]
            rows = [
                (
                    str(sheet.cell(row=row, column=ROLL_COLUMN).value),
                    sheet.cell(row=row, column=MARKS_COLUMN).value,
                )
                for row in range(HEADER_ROW + 1, sheet.max_row + 1)
                if sheet.cell(row=row, column=ROLL_COLUMN).value not in (None, "")
            ]
        finally:
            workbook.close()
        del ten
        by_mark: dict[object, list[str]] = {}
        for roll, mark in rows:
            by_mark.setdefault(mark, []).append(roll)
        tied = [rolls for rolls in by_mark.values() if len(rolls) > 1]
        assert tied, "the fixture is meant to contain a tie"
        for rolls in tied:
            assert rolls == sorted(rolls)

    def test_the_merit_column_holds_a_rank_formula(self, generated):
        """§9: ranking stays an Excel formula, so ties share a rank."""
        _session, _built, _template, outcomes = generated
        workbook = openpyxl.load_workbook(outcomes["10"].output_path)
        try:
            sheet = workbook[rx.MERITWISE_SHEET_NAME]
            value = sheet.cell(row=HEADER_ROW + 1, column=MERIT_COLUMN).value
        finally:
            workbook.close()
        assert isinstance(value, str)
        assert "RANK.EQ(" in value

    def test_it_keeps_the_templates_own_heading(self, generated):
        """§11: meritwise derives from the completed rollwise sheet."""
        _session, built, _template, outcomes = generated
        for item in built:
            workbook = openpyxl.load_workbook(outcomes[item.plan.code].output_path)
            try:
                sheet = workbook[rx.MERITWISE_SHEET_NAME]
                assert sheet.cell(row=1, column=1).value == (
                    "Bangladesh Submarine Cable Regulatory Authority"
                )
                assert sheet.cell(row=2, column=1).value == item.plan.description
            finally:
                workbook.close()


# ----------------------------------------------------------------------
# §19 / §23 - the critical isolation requirement
# ----------------------------------------------------------------------
class TestNoCrossSetContamination:
    def test_no_candidate_from_another_set_appears_anywhere_in_a_workbook(
        self, generated
    ):
        _session, built, _template, outcomes = generated
        by_code = {item.plan.code: set(item.plan.rolls) for item in built}
        for item in built:
            mine = by_code[item.plan.code]
            foreign = set().union(
                *(rolls for code, rolls in by_code.items() if code != item.plan.code)
            ) - mine
            text = every_cell_text(outcomes[item.plan.code].output_path)
            leaked = foreign & text
            assert not leaked, (item.plan.code, sorted(leaked)[:5])

    def test_the_shared_roll_carries_each_sets_own_mark(self, generated):
        """The sharpest form of §19: one roll, two Sets, two different marks."""
        session, built, template, outcomes = generated
        marks_written: dict[str, object] = {}
        expected: dict[str, float] = {}
        for code in ("10", "11"):
            item = next(entry for entry in built if entry.plan.code == code)
            expected[code] = stored_marks(session.database, item, template)[SHARED_ROLL]
            association = report_store.resolve_set_sources(
                session.database, item.exam_set.set_id
            ).association
            workbook = openpyxl.load_workbook(outcomes[code].output_path)
            try:
                marks_written[code] = marks_by_roll(
                    workbook[association.sheet_name], HEADER_ROW + 1
                )[SHARED_ROLL]
            finally:
                workbook.close()

        assert expected["10"] != expected["11"], (
            "the fixture must give the shared roll different marks in each Set"
        )
        assert marks_written["10"] == pytest.approx(expected["10"])
        assert marks_written["11"] == pytest.approx(expected["11"])

    def test_each_workbook_was_built_from_its_own_attendance_template(self, generated):
        session, built, _template, _outcomes = generated
        for item in built:
            sources = report_store.resolve_set_sources(
                session.database, item.exam_set.set_id
            )
            assert sources.association.template_path == str(item.attendance_path)
            assert sources.roster_id == item.roster_id

    def test_each_workbook_carries_its_own_posts_name(self, generated):
        _session, built, _template, outcomes = generated
        for item in built:
            text = every_cell_text(outcomes[item.plan.code].output_path)
            assert item.plan.description in text
            for other in built:
                if other.plan.code != item.plan.code:
                    assert other.plan.description not in text


# ----------------------------------------------------------------------
# §26 - the source workbooks are never modified
# ----------------------------------------------------------------------
class TestOriginalAttendanceFilesAreUntouched:
    def test_every_source_workbook_hashes_the_same_after_generation(
        self, examination, tmp_path: Path
    ):
        session, built, template = examination
        before = {
            item.plan.code: sha256_of(item.attendance_path) for item in built
        }
        for item in built:
            report_store.generate_for_set(
                session.database, item.exam_set.set_id, item.batch_id, template,
                project_name="BSCRA Recruitment", output_dir=tmp_path / "exports",
                final=False,
            )
        after = {item.plan.code: sha256_of(item.attendance_path) for item in built}
        assert after == before

    def test_regenerating_twice_still_leaves_them_untouched(
        self, examination, tmp_path: Path
    ):
        session, built, template = examination
        ten = next(item for item in built if item.plan.code == "10")
        before = sha256_of(ten.attendance_path)
        for _ in range(2):
            report_store.generate_for_set(
                session.database, ten.exam_set.set_id, ten.batch_id, template,
                project_name="BSCRA Recruitment", output_dir=tmp_path / "exports",
            )
        assert sha256_of(ten.attendance_path) == before

    def test_the_output_is_a_different_file_from_the_template(self, generated):
        _session, built, _template, outcomes = generated
        for item in built:
            assert outcomes[item.plan.code].output_path != item.attendance_path


# ----------------------------------------------------------------------
# §27 at full size
# ----------------------------------------------------------------------
@pytest.mark.stress
class TestFullScaleExamination:
    """500 / 350 / 725 candidates, as §27 specifies.

    Excluded from the default run by the ``stress`` marker; run with
    ``-m stress``. Everything it asserts is asserted at the smaller default
    size too - this exists to show the same guarantees hold at the size an
    examination office actually works at.
    """

    def test_three_full_size_sets_generate_without_cross_contamination(
        self, workspace: Path, tmp_path: Path
    ):
        session = create_project(
            workspace, "BSCRA Recruitment", exam_name=EXAM_NAME
        )
        try:
            built, template = build_examination(session, tmp_path, STRESS_SIZES)
            exports = tmp_path / "exports"
            outcomes = {
                item.plan.code: report_store.generate_for_set(
                    session.database, item.exam_set.set_id, item.batch_id, template,
                    project_name="BSCRA Recruitment", output_dir=exports,
                    final=item.plan.code != "12",
                )
                for item in built
            }

            for item in built:
                outcome = outcomes[item.plan.code]
                assert outcome.output_path is not None, outcome.warnings

            assert [len(item.plan.rolls) for item in built] == [500, 350, 725]

            by_code = {item.plan.code: set(item.plan.rolls) for item in built}
            for item in built:
                mine = by_code[item.plan.code]
                foreign = set().union(
                    *(
                        rolls for code, rolls in by_code.items()
                        if code != item.plan.code
                    )
                ) - mine
                text = every_cell_text(outcomes[item.plan.code].output_path)
                assert not (foreign & text), item.plan.code

                association = report_store.resolve_set_sources(
                    session.database, item.exam_set.set_id
                ).association
                workbook = openpyxl.load_workbook(outcomes[item.plan.code].output_path)
                try:
                    rollwise = workbook[association.sheet_name]
                    assert sheet_rolls(rollwise, HEADER_ROW + 1) == item.plan.rolls
                    merit = set(
                        sheet_rolls(workbook[rx.MERITWISE_SHEET_NAME], HEADER_ROW + 1)
                    )
                finally:
                    workbook.close()
                assert not (merit & item.plan.absent)
        finally:
            if not session.is_closed:
                session.close()
