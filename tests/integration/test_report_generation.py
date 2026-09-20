"""End-to-end result-report generation (Phase 9).

Scope:
    :mod:`omr_scanner.services.report_store` against a real project database
    and real Phase 7/8 services - reconciliation, scoring - exactly as the
    application wires them, plus real openpyxl workbooks built by
    ``tests/report_fixtures.py``.

Why this level carries the weight:
    A generated report is what an examination office files. Every assertion
    here reads the *generated workbook itself* back with openpyxl, not
    internal state - the same thing an operator opening the file would see.
"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from fractions import Fraction
from typing import TYPE_CHECKING

import openpyxl
import pytest
from tests.conftest import build_answer_sheet_template
from tests.report_fixtures import build_result_template, sha256_of

from omr_scanner.database import open_project_database
from omr_scanner.domain.reconciliation import AttendanceState, CandidateRecord
from omr_scanner.domain.reporting import ReadinessIssueKind
from omr_scanner.domain.scoring import NegativeMarking, ScoringPolicy
from omr_scanner.reporting import excel as rx
from omr_scanner.services import batch_store, reconciliation_store, report_store, scoring_store
from omr_scanner.services.answer_key import plan_for, read_key
from omr_scanner.services.candidate_import import ColumnMapping, RosterValidation
from omr_scanner.services.recognition_models import (
    AnswerView,
    FieldView,
    RecognitionOutcome,
    RegistrationStatus,
    ScanResult,
)
from omr_scanner.services.report_template import ReportColumnMapping, preview_template

if TYPE_CHECKING:
    from collections.abc import Iterator
    from pathlib import Path

    from omr_scanner.database import ProjectDatabase

OPERATOR = "Dr. Rahman"
PRESENT = AttendanceState.PRESENT
ABSENT = AttendanceState.ABSENT

ROSTER_A = [
    ("10001", "CAND A", PRESENT),
    ("10002", "CAND B", PRESENT),
    ("10003", "CAND C", ABSENT),
    ("10004", "CAND D", PRESENT),
]


@pytest.fixture(scope="module")
def template():
    return build_answer_sheet_template()


@pytest.fixture(scope="module")
def plan(template):
    return plan_for(template)


@pytest.fixture
def database(tmp_path: Path) -> Iterator[ProjectDatabase]:
    handle = open_project_database(tmp_path / "database.sqlite", create=True)
    try:
        yield handle
    finally:
        handle.close()


def validation_of(candidates) -> RosterValidation:
    records = tuple(
        CandidateRecord(
            candidate_id=candidate_id, display_name=name, source_row=index + 2,
            imported_attendance=state,
            imported_value="ABSENT" if state is ABSENT else "55",
        )
        for index, (candidate_id, name, state) in enumerate(candidates)
    )
    return RosterValidation(
        candidates=records, issues=(), rows_read=len(records), blank_ids=0,
        duplicate_ids=0,
        expected_present=sum(1 for i in records if i.imported_attendance is PRESENT),
        expected_absent=sum(1 for i in records if i.imported_attendance is ABSENT),
        attendance_unknown=0, mapping=ColumnMapping(candidate_id=0, name=1, attendance=2),
        source_name="roster.csv",
    )


def make_result(path, roll: str, set_code: str, answers: dict[int, str], numbers):
    return ScanResult(
        source_path=path, outcome=RecognitionOutcome.COMPLETE,
        registration=RegistrationStatus.REGISTERED,
        fields=(
            FieldView(
                zone_id="roll_number", label="Roll", field_type="numeric", value=roll,
                status="complete", needs_review=False, characters=(),
            ),
            FieldView(
                zone_id="set_code", label="Set", field_type="set_code", value=set_code,
                status="complete", needs_review=False, characters=(),
            ),
        ),
        answers=tuple(
            AnswerView(
                number=number, zone_id="q", value=answers.get(number, "A"),
                status="resolved", needs_review=False, top_fill=0.9, margin=0.4,
                confidence=0.9,
            )
            for number in numbers
        ),
        identifier_zone_id="roll_number", set_code_zone_id="set_code",
    )


def build_batch(database, template, plan, tmp_path, rows, *, roster=None):
    """Reconcile a batch of ``(filename, roll, set_code, answers)`` rows."""
    roster_id = reconciliation_store.import_roster(
        database, validation_of(roster or ROSTER_A), imported_by=OPERATOR
    )
    from omr_scanner.database.models import BatchScan, ScanBatch

    now = datetime.now(UTC)
    batch_id = batch_store.new_batch_id()
    with database.session() as session:
        session.add(
            ScanBatch(
                batch_id=batch_id, created_at=now, updated_at=now,
                source_folder=str(tmp_path), status="completed", total_scans=len(rows),
            )
        )
        session.flush()
        for index, (name, roll, set_code, answers) in enumerate(rows):
            result = make_result(tmp_path / name, roll, set_code, answers, plan.numbers)
            session.add(
                BatchScan(
                    batch_id=batch_id, batch_index=index, source_path=str(tmp_path / name),
                    filename=name, status="completed", identifier_value=roll,
                    set_code_value=set_code, result_json=json.dumps(result.to_dict()),
                )
            )
    reconciliation_store.reconcile_batch(database, roster_id, batch_id)
    return roster_id, batch_id


def verified_set(database, plan, set_code: str, answers: str | None = None, wrong=()):
    stored = scoring_store.save_key(
        database,
        read_key(
            answers or "A" * plan.question_count, plan, set_code, wrong_questions=list(wrong)
        ).to_key(),
    )
    return scoring_store.verify_key(database, stored.key_id, verified_by=OPERATOR)


@pytest.fixture
def association_mapping():
    """The column mapping matching ``build_result_template``'s default headers."""
    return ReportColumnMapping(roll=1, marks=3, serial=0, name=2, rank=4)


def associate(database, set_code, path, mapping=None):
    from tests.report_fixtures import SAMPLE_HEADERS

    from omr_scanner.services.report_template import suggest_mapping

    used_mapping = mapping or suggest_mapping(list(SAMPLE_HEADERS)).to_mapping()
    preview = preview_template(path)
    return report_store.associate_template(
        database, set_code, path, used_mapping, sheet_name=preview.sheet,
        updated_by=OPERATOR,
    )


def results_by_id(database, roster_id, batch_id, template):
    return {
        item.candidate_id: item
        for item in scoring_store.list_results(database, roster_id, batch_id, template)
    }


# ----------------------------------------------------------------------
class TestAcceptanceScenarioA:
    """Scenario A of the phase brief (§54): a normal, complete result."""

    @pytest.fixture
    def ready_set(self, database, template, plan, tmp_path):
        rows = [
            ("s1.png", "10001", "A", dict.fromkeys(plan.numbers, "A")),
            ("s2.png", "10002", "A", dict.fromkeys(plan.numbers[:3], "B")),
            ("s4.png", "10004", "A", dict.fromkeys(plan.numbers, "A")),
        ]
        roster_id, batch_id = build_batch(database, template, plan, tmp_path, rows)
        verified_set(database, plan, "A")
        scoring_store.score_batch(database, roster_id, batch_id, template)

        template_path = build_result_template(
            tmp_path / "set_a_template.xlsx", candidate_count=4,
            roll_prefix="1000", roll_digits=5,
            # Position 3 (roll 10003) is the ROSTER_A candidate recorded
            # absent - a real absentee sheet already shows this before the
            # exam, so the fixture must agree with reconciliation or the
            # readiness check correctly (and rightly) flags a mismatch.
            names=["CAND A", "CAND B", "CAND C", "CAND D"], absent_every=3,
        )
        associate(database, "A", template_path)
        return roster_id, batch_id

    def test_readiness_is_clear(self, database, template, ready_set):
        roster_id, batch_id = ready_set
        report = report_store.check_readiness(
            database, roster_id, batch_id, template, "A", for_final_export=True
        )
        assert report.is_ready, report.describe()

    def test_rollwise_contains_every_candidate(self, database, template, ready_set, tmp_path):
        roster_id, batch_id = ready_set
        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A",
            project_name="Test Exam", output_dir=tmp_path / "exports", computed_by=OPERATOR,
        )
        assert outcome.ok, outcome.warnings
        workbook = openpyxl.load_workbook(outcome.output_path)
        sheet = workbook["5.AD-E-1.Rollwise(All)"]
        rolls = [sheet.cell(row=r, column=2).value for r in range(2, 6)]
        assert rolls == ["10001", "10002", "10003", "10004"]

    def test_the_absent_candidate_remains_in_position(
        self, database, template, ready_set, tmp_path
    ):
        roster_id, batch_id = ready_set
        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A",
            project_name="Test Exam", output_dir=tmp_path / "exports",
        )
        workbook = openpyxl.load_workbook(outcome.output_path)
        sheet = workbook["5.AD-E-1.Rollwise(All)"]
        assert sheet.cell(row=4, column=2).value == "10003"  # position unchanged
        assert sheet.cell(row=4, column=4).value == "ABSENT"
        assert sheet.cell(row=4, column=5).value == "---"

    def test_present_candidates_have_numeric_scores(
        self, database, template, ready_set, tmp_path
    ):
        roster_id, batch_id = ready_set
        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A",
            project_name="Test Exam", output_dir=tmp_path / "exports",
        )
        workbook = openpyxl.load_workbook(outcome.output_path)
        sheet = workbook["5.AD-E-1.Rollwise(All)"]
        assert sheet.cell(row=2, column=4).value == plan_for(template).question_count

    def test_rank_formulas_are_present_for_scored_candidates(
        self, database, template, ready_set, tmp_path
    ):
        roster_id, batch_id = ready_set
        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A",
            project_name="Test Exam", output_dir=tmp_path / "exports",
        )
        workbook = openpyxl.load_workbook(outcome.output_path)
        sheet = workbook["5.AD-E-1.Rollwise(All)"]
        assert "RANK.EQ" in str(sheet.cell(row=2, column=5).value)

    def test_meritwise_contains_only_scored_candidates(
        self, database, template, ready_set, tmp_path
    ):
        roster_id, batch_id = ready_set
        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A",
            project_name="Test Exam", output_dir=tmp_path / "exports",
        )
        workbook = openpyxl.load_workbook(outcome.output_path)
        sheet = workbook[rx.MERITWISE_SHEET_NAME]
        rolls = [
            sheet.cell(row=r, column=2).value
            for r in range(2, sheet.max_row + 1)
            if sheet.cell(row=r, column=2).value
        ]
        assert set(rolls) == {"10001", "10002", "10004"}

    def test_support_sheets_all_exist(self, database, template, ready_set, tmp_path):
        roster_id, batch_id = ready_set
        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A",
            project_name="Test Exam", output_dir=tmp_path / "exports",
        )
        workbook = openpyxl.load_workbook(outcome.output_path)
        for name in (
            "5.AD-E-1.Rollwise(All)", rx.MERITWISE_SHEET_NAME, rx.SUMMARY_SHEET_NAME,
            rx.ANSWER_KEY_SHEET_NAME, rx.PROCESSING_LOG_SHEET_NAME,
        ):
            assert name in workbook.sheetnames

    def test_generation_is_recorded_on_the_audit_trail(
        self, database, template, ready_set, tmp_path
    ):
        roster_id, batch_id = ready_set
        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A",
            project_name="Test Exam", output_dir=tmp_path / "exports",
        )
        assert outcome.report_id > 0
        assert outcome.status == "success"


# ----------------------------------------------------------------------
class TestAcceptanceScenarioB:
    """Tied marks (§54 Scenario B): 88, 88, 85 -> 1, 1, 3."""

    def test_application_rank_matches_excel_semantics(self, database, template, plan, tmp_path):
        rows = [
            ("s1.png", "10001", "A", dict.fromkeys(plan.numbers, "A")),  # full marks
            ("s2.png", "10002", "A", dict.fromkeys(plan.numbers, "A")),  # full marks, tied
            ("s4.png", "10004", "A", dict.fromkeys(plan.numbers[:3], "B")),  # fewer
        ]
        roster_id, batch_id = build_batch(database, template, plan, tmp_path, rows)
        verified_set(database, plan, "A")
        scoring_store.score_batch(database, roster_id, batch_id, template)
        results = results_by_id(database, roster_id, batch_id, template)

        from omr_scanner.domain.reporting import compute_ranks

        scores = [
            (candidate_id, item.final_score if item.has_mark else None)
            for candidate_id, item in results.items()
            if candidate_id in ("10001", "10002", "10004")
        ]
        ranks = compute_ranks(scores)
        assert ranks["10001"] == ranks["10002"] == 1
        assert ranks["10004"] == 3


# ----------------------------------------------------------------------
class TestAcceptanceScenarioC:
    """A candidate in the template but not the project (§54 Scenario C)."""

    def test_a_missing_candidate_blocks_final_export(self, database, template, plan, tmp_path):
        roster_id, batch_id = build_batch(
            database, template, plan, tmp_path,
            [("s1.png", "10001", "A", dict.fromkeys(plan.numbers, "A"))],
            roster=[("10001", "CAND A", PRESENT)],
        )
        verified_set(database, plan, "A")
        scoring_store.score_batch(database, roster_id, batch_id, template)

        template_path = build_result_template(
            tmp_path / "t.xlsx", candidate_count=2, roll_prefix="1000", roll_digits=5,
            names=["CAND A", "GHOST"],
        )
        associate(database, "A", template_path)

        report = report_store.check_readiness(
            database, roster_id, batch_id, template, "A", for_final_export=True
        )
        assert not report.is_ready
        assert report.by_kind(ReadinessIssueKind.CANDIDATE_NOT_IN_PROJECT)

    def test_final_export_is_refused_not_silently_missing_the_row(
        self, database, template, plan, tmp_path
    ):
        roster_id, batch_id = build_batch(
            database, template, plan, tmp_path,
            [("s1.png", "10001", "A", dict.fromkeys(plan.numbers, "A"))],
            roster=[("10001", "CAND A", PRESENT)],
        )
        verified_set(database, plan, "A")
        scoring_store.score_batch(database, roster_id, batch_id, template)
        template_path = build_result_template(
            tmp_path / "t.xlsx", candidate_count=2, roll_prefix="1000", roll_digits=5,
            names=["CAND A", "GHOST"],
        )
        associate(database, "A", template_path)

        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=tmp_path / "exports", final=True,
        )
        assert outcome.status == "blocked"
        assert outcome.output_path is None


# ----------------------------------------------------------------------
class TestAcceptanceScenarioD:
    """An existing output file is never silently overwritten (§54 Scenario D)."""

    def test_a_second_generation_produces_a_new_filename(
        self, database, template, plan, tmp_path
    ):
        roster_id, batch_id = build_batch(
            database, template, plan, tmp_path,
            [("s1.png", "10001", "A", dict.fromkeys(plan.numbers, "A"))],
            roster=[("10001", "CAND A", PRESENT)],
        )
        verified_set(database, plan, "A")
        scoring_store.score_batch(database, roster_id, batch_id, template)
        template_path = build_result_template(
            tmp_path / "t.xlsx", candidate_count=1, roll_prefix="1000", roll_digits=5,
            names=["CAND A"],
        )
        associate(database, "A", template_path)

        outdir = tmp_path / "exports"
        first = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=outdir,
        )
        second = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=outdir,
        )
        assert first.output_path != second.output_path
        assert first.output_path.exists()
        assert second.output_path.exists()
        assert second.output_path.name.endswith("_1.xlsx")


# ----------------------------------------------------------------------
class TestAcceptanceScenarioE:
    """Two sets, two templates, two keys - no contamination (§54 Scenario E, §39)."""

    @pytest.fixture
    def two_sets(self, database, template, plan, tmp_path):
        count = plan.question_count
        rows = [
            ("s1.png", "10001", "A", dict.fromkeys(plan.numbers, "A")),
            ("s2.png", "20001", "B", dict.fromkeys(plan.numbers, "D")),
        ]
        roster_id, batch_id = build_batch(
            database, template, plan, tmp_path, rows,
            roster=[("10001", "CAND A", PRESENT), ("20001", "CAND X", PRESENT)],
        )
        verified_set(database, plan, "A", "A" * count)
        verified_set(database, plan, "B", "D" * count)
        scoring_store.score_batch(database, roster_id, batch_id, template)

        template_a = build_result_template(
            tmp_path / "set_a.xlsx", candidate_count=1, roll_prefix="1000",
            roll_digits=5, names=["CAND A"], sheet_name="Set A Sheet",
        )
        template_b = build_result_template(
            tmp_path / "set_b.xlsx", candidate_count=1, roll_prefix="2000",
            roll_digits=5, names=["CAND X"], sheet_name="Set B Sheet",
        )
        associate(database, "A", template_a)
        associate(database, "B", template_b)
        return roster_id, batch_id, tmp_path

    def test_set_a_workbook_contains_only_set_a_candidates(self, database, template, two_sets):
        roster_id, batch_id, tmp_path = two_sets
        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=tmp_path / "exports",
        )
        workbook = openpyxl.load_workbook(outcome.output_path)
        sheet = workbook["Set A Sheet"]
        rolls = [sheet.cell(row=r, column=2).value for r in range(2, sheet.max_row + 1)]
        assert "10001" in rolls
        assert "20001" not in rolls

    def test_set_b_workbook_contains_only_set_b_candidates(self, database, template, two_sets):
        roster_id, batch_id, tmp_path = two_sets
        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "B", project_name="P",
            output_dir=tmp_path / "exports",
        )
        workbook = openpyxl.load_workbook(outcome.output_path)
        sheet = workbook["Set B Sheet"]
        rolls = [sheet.cell(row=r, column=2).value for r in range(2, sheet.max_row + 1)]
        assert "20001" in rolls
        assert "10001" not in rolls

    def test_set_a_answer_key_sheet_shows_only_set_as_key(self, database, template, two_sets):
        roster_id, batch_id, tmp_path = two_sets
        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=tmp_path / "exports",
        )
        workbook = openpyxl.load_workbook(outcome.output_path)
        sheet = workbook[rx.ANSWER_KEY_SHEET_NAME]
        assert "Set A" in sheet.cell(row=1, column=1).value
        assert "A" * plan_for(template).question_count in sheet.cell(row=3, column=1).value

    def test_set_bs_candidate_scores_full_marks_only_against_its_own_key(
        self, database, template, two_sets
    ):
        # If Set A's key were used against Set B's candidate (all "D"), the
        # mark would be zero instead of full - a wrong-key swap could not go
        # unnoticed.
        roster_id, batch_id, _tmp_path = two_sets
        results = results_by_id(database, roster_id, batch_id, template)
        assert results["20001"].final_score == plan_for(template).question_count
        assert results["10001"].final_score == plan_for(template).question_count

    def test_generated_report_audit_rows_reference_the_correct_set(
        self, database, template, two_sets
    ):
        roster_id, batch_id, tmp_path = two_sets
        report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=tmp_path / "exports",
        )
        report_store.generate_xlsx(
            database, roster_id, batch_id, template, "B", project_name="P",
            output_dir=tmp_path / "exports",
        )
        from sqlalchemy import select

        from omr_scanner.database.models import GeneratedReport

        with database.session() as session:
            rows = session.scalars(select(GeneratedReport)).all()
            by_set = {row.set_code for row in rows}
        assert by_set == {"A", "B"}


# ----------------------------------------------------------------------
class TestAcceptanceScenarioF:
    """Changed scoring configuration regenerates rather than patches (§54 F)."""

    def test_a_regeneration_reflects_the_new_policy(self, database, template, plan, tmp_path):
        roster_id, batch_id = build_batch(
            database, template, plan, tmp_path,
            [("s1.png", "10001", "A", {plan.numbers[0]: "B"})],
            roster=[("10001", "CAND A", PRESENT)],
        )
        verified_set(database, plan, "A")
        scoring_store.score_batch(database, roster_id, batch_id, template)
        template_path = build_result_template(
            tmp_path / "t.xlsx", candidate_count=1, roll_prefix="1000", roll_digits=5,
            names=["CAND A"],
        )
        associate(database, "A", template_path)

        first = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=tmp_path / "exports1",
        )
        first_mark = openpyxl.load_workbook(first.output_path)[
            "5.AD-E-1.Rollwise(All)"
        ].cell(row=2, column=4).value

        scoring_store.save_policy(
            database,
            ScoringPolicy(
                correct_mark=Fraction(1), incorrect_penalty=Fraction(1, 4),
                mode=NegativeMarking.FIXED,
            ),
        )
        scoring_store.score_batch(database, roster_id, batch_id, template)
        second = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=tmp_path / "exports2",
        )
        second_mark = openpyxl.load_workbook(second.output_path)[
            "5.AD-E-1.Rollwise(All)"
        ].cell(row=2, column=4).value

        assert first_mark != second_mark
        assert second_mark == pytest.approx(
            plan.question_count - 1 - 0.25
        )


# ----------------------------------------------------------------------
class TestTemplatePreservation:
    def test_the_template_file_is_never_modified_by_generation(
        self, database, template, plan, tmp_path
    ):
        roster_id, batch_id = build_batch(
            database, template, plan, tmp_path,
            [("s1.png", "10001", "A", dict.fromkeys(plan.numbers, "A"))],
            roster=[("10001", "CAND A", PRESENT)],
        )
        verified_set(database, plan, "A")
        scoring_store.score_batch(database, roster_id, batch_id, template)
        template_path = build_result_template(
            tmp_path / "t.xlsx", candidate_count=1, roll_prefix="1000", roll_digits=5,
            names=["CAND A"],
        )
        associate(database, "A", template_path)
        before = sha256_of(template_path)

        report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=tmp_path / "exports",
        )
        after = sha256_of(template_path)
        assert before == after


# ----------------------------------------------------------------------
class TestPdfExport:
    def test_pdf_generation_uses_a_freshly_generated_workbook(
        self, database, template, plan, tmp_path
    ):
        from omr_scanner.reporting.pdf import PdfExportResult

        roster_id, batch_id = build_batch(
            database, template, plan, tmp_path,
            [("s1.png", "10001", "A", dict.fromkeys(plan.numbers, "A"))],
            roster=[("10001", "CAND A", PRESENT)],
        )
        verified_set(database, plan, "A")
        scoring_store.score_batch(database, roster_id, batch_id, template)
        template_path = build_result_template(
            tmp_path / "t.xlsx", candidate_count=1, roll_prefix="1000", roll_digits=5,
            names=["CAND A"],
        )
        associate(database, "A", template_path)

        class RecordingExporter:
            """Read the sheet names *during* export.

            The source file lives in a temporary directory report_store
            cleans up as soon as export() returns, so verifying it
            afterwards would be reading a file that no longer exists, not a
            defect in report_store.
            """

            name = "fake"

            def __init__(self) -> None:
                self.observed_sheetnames: list[list[str]] = []

            def is_available(self) -> bool:
                return True

            def export(self, workbook_path, output_pdf_path) -> PdfExportResult:
                # Closed explicitly: a read-only openpyxl handle left open
                # keeps the file locked on Windows, and report_store deletes
                # this temporary file the instant export() returns.
                source = openpyxl.load_workbook(workbook_path, read_only=True)
                try:
                    self.observed_sheetnames.append(list(source.sheetnames))
                finally:
                    source.close()
                output_pdf_path.write_bytes(b"%PDF-1.4\n%%EOF")
                return PdfExportResult(output_path=output_pdf_path, engine="fake")

        exporter = RecordingExporter()
        outcome = report_store.generate_pdf(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=tmp_path / "exports", sheet_name="5.AD-E-1.Rollwise(All)",
            report_label="Rollwise", pdf_exporter=exporter,
        )
        assert outcome.ok
        assert outcome.output_path.suffix == ".pdf"
        assert outcome.output_path.exists()
        # The exported source contained only the Rollwise sheet.
        assert exporter.observed_sheetnames == [["5.AD-E-1.Rollwise(All)"]]

    def test_pdf_export_fails_gracefully_when_no_engine_is_available(
        self, database, template, plan, tmp_path
    ):
        from omr_scanner.reporting.pdf import UnavailablePdfExporter

        roster_id, batch_id = build_batch(
            database, template, plan, tmp_path,
            [("s1.png", "10001", "A", dict.fromkeys(plan.numbers, "A"))],
            roster=[("10001", "CAND A", PRESENT)],
        )
        verified_set(database, plan, "A")
        scoring_store.score_batch(database, roster_id, batch_id, template)
        template_path = build_result_template(
            tmp_path / "t.xlsx", candidate_count=1, roll_prefix="1000", roll_digits=5,
            names=["CAND A"],
        )
        associate(database, "A", template_path)

        outcome = report_store.generate_pdf(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=tmp_path / "exports", sheet_name="5.AD-E-1.Rollwise(All)",
            report_label="Rollwise", pdf_exporter=UnavailablePdfExporter(),
        )
        assert outcome.status == "failed"
        assert outcome.output_path is None
        # The XLSX itself was still generated successfully underneath.
        assert list((tmp_path / "exports").glob("*.xlsx"))


# ----------------------------------------------------------------------
class TestMigrationOntoAnExistingPhase8Project:
    """A project scored before Phase 9 must keep everything and gain the rest."""

    def test_a_phase_8_project_upgrades_and_becomes_reportable(self, tmp_path):
        import sqlite3

        from sqlalchemy import text

        from omr_scanner.database.migrations import SCHEMA_VERSION

        db_path = tmp_path / "database.sqlite"
        handle = open_project_database(db_path, create=True)
        roster_id = reconciliation_store.import_roster(
            handle, validation_of([("10001", "CAND A", PRESENT)]), imported_by=OPERATOR
        )
        handle.close()

        connection = sqlite3.connect(db_path)
        try:
            connection.executescript(
                """
                DROP TABLE IF EXISTS generated_report;
                DROP TABLE IF EXISTS report_layout_config;
                DROP TABLE IF EXISTS report_template_association;
                DELETE FROM schema_migration WHERE version >= 6;
                """
            )
            connection.commit()
        finally:
            connection.close()

        reopened = open_project_database(db_path)
        try:
            assert reopened.schema_version == SCHEMA_VERSION
            with reopened.session() as session:
                names = {
                    row[0]
                    for row in session.execute(
                        text("SELECT name FROM sqlite_master WHERE type='table'")
                    ).all()
                }
            assert {
                "report_template_association", "report_layout_config", "generated_report",
            } <= names
            # Phase 7's data is untouched.
            assert (
                len(reconciliation_store.roster_candidates(reopened, roster_id)) == 1
            )
            # ...and a template can now be associated.
            template_path = build_result_template(
                tmp_path / "t.xlsx", candidate_count=1, roll_prefix="1000",
                roll_digits=5, names=["CAND A"],
            )
            stored = associate(reopened, "A", template_path)
            assert stored.set_code == "A"
        finally:
            reopened.close()


class TestReadinessBlocking:
    def test_no_template_blocks_export(self, database, template, plan, tmp_path):
        roster_id, batch_id = build_batch(
            database, template, plan, tmp_path,
            [("s1.png", "10001", "A", dict.fromkeys(plan.numbers, "A"))],
            roster=[("10001", "CAND A", PRESENT)],
        )
        verified_set(database, plan, "A")
        scoring_store.score_batch(database, roster_id, batch_id, template)
        report = report_store.check_readiness(
            database, roster_id, batch_id, template, "A", for_final_export=True
        )
        assert not report.is_ready
        assert report.by_kind(ReadinessIssueKind.NO_TEMPLATE)

    def test_present_without_score_blocks_export(self, database, template, plan, tmp_path):
        roster_id, batch_id = build_batch(
            database, template, plan, tmp_path, [],
            roster=[("10001", "CAND A", PRESENT)],
        )
        template_path = build_result_template(
            tmp_path / "t.xlsx", candidate_count=1, roll_prefix="1000", roll_digits=5,
            names=["CAND A"],
        )
        associate(database, "A", template_path)
        report = report_store.check_readiness(
            database, roster_id, batch_id, template, "A", for_final_export=True
        )
        assert not report.is_ready
        assert report.by_kind(ReadinessIssueKind.PRESENT_WITHOUT_SCORE)

    def test_a_preview_still_generates_with_warnings(self, database, template, plan, tmp_path):
        roster_id, batch_id = build_batch(
            database, template, plan, tmp_path, [],
            roster=[("10001", "CAND A", PRESENT)],
        )
        template_path = build_result_template(
            tmp_path / "t.xlsx", candidate_count=1, roll_prefix="1000", roll_digits=5,
            names=["CAND A"],
        )
        associate(database, "A", template_path)
        outcome = report_store.generate_xlsx(
            database, roster_id, batch_id, template, "A", project_name="P",
            output_dir=tmp_path / "exports", final=False,
        )
        assert outcome.status == "success"
        assert outcome.warnings
