"""Reject & Rescan, end to end through the services, on a real project.

Scope:
    One batch holding Sets 1, 2 and 3 - the same shape as the combined
    multi-set scenario - with real files in the project's ``scans_original``
    folder, content hashes, per-set attendance workbooks (so result templates
    exist), verified keys and scoring. Every invariant the lifecycle promises
    is asserted where it is enforced: in the services, not the GUI.

Rescans are stored rows rather than rendered sheets, registered with
:func:`~omr_scanner.services.batch_store.add_scans_to_batch` exactly as the
Scan stage now registers them, and given a recognition result the way the
coordinator records one. ``tests/gui/test_reject_rescan_gui.py`` drives real
recognition of real rendered rescans through the Scan stage.

Privacy:
    Every identifier and name here is fictional.
"""

from __future__ import annotations

import json
import os
import sqlite3
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import TYPE_CHECKING

import openpyxl
import pytest
from sqlalchemy import select, text, update
from tests.conftest import build_answer_sheet_template

from omr_scanner.database.migrations import SCHEMA_VERSION
from omr_scanner.database.models import AuditEvent, BatchScan, ReviewConflict, ScanBatch
from omr_scanner.domain.reconciliation import ReconciliationStatus
from omr_scanner.domain.reporting import ReadinessIssueKind
from omr_scanner.domain.review import ConflictState, ConflictType
from omr_scanner.domain.scan_lifecycle import (
    EXPORT_ENTITY,
    LIFECYCLE_ENTITY,
    FileState,
    LifecycleAction,
    LifecycleState,
    PurgeMode,
    RejectionReason,
)
from omr_scanner.domain.scoring import BlockReason, ResultStatus, StaleReason
from omr_scanner.services import (
    batch_store,
    candidate_import,
    create_project,
    open_project,
    project_health,
    project_sets,
    reconciliation_store,
    report_store,
    review_store,
    scan_lifecycle,
    scan_provenance,
    scoring_store,
    set_attendance,
)
from omr_scanner.services.answer_key import plan_for, read_key
from omr_scanner.services.recognition_models import (
    AnswerView,
    FieldView,
    RecognitionOutcome,
    RegistrationStatus,
    ScanResult,
)
from omr_scanner.services.review_store import ReviewError
from omr_scanner.services.scan_lifecycle import LifecycleError

if TYPE_CHECKING:
    from collections.abc import Iterator
    from pathlib import Path

    from omr_scanner.services import ProjectSession

OPERATOR = "Dr. A. Rahman"
TEMPLATE = build_answer_sheet_template()
PLAN = plan_for(TEMPLATE)

HEADERS = ("Sl.No.", "Roll No.", "Name", "Total", "Merit")
HEADER_ROW = 3

ROSTERS: dict[str, list[tuple[str, bool]]] = {
    "1": [("100121", False), ("100122", False)],
    "2": [("200121", False), ("200122", False), ("200123", False), ("200124", True)],
    "3": [("300121", False), ("300122", False)],
}
"""Per set: ``(roll, absent)``."""

SCRIPTS: tuple[tuple[str, str, str, int], ...] = (
    ("s1a.png", "100121", "1", 12),
    ("s1b.png", "100122", "1", 13),
    ("s2a.png", "200121", "2", 14),
    ("s2b.png", "200122", "2", 15),  # the poor-quality scan rejected below
    ("s2c.png", "20123", "2", 16),  # 200123 wrote five digits: one missing
    ("s3a.png", "300121", "3", 17),
    ("s3b.png", "300122", "3", 18),
    ("s_x.png", "200124", "7", 19),  # a set this project does not define
    ("s_blur.png", "2?0?2?", "?", 5),  # unreadable identity and set
)
"""``(file, roll as read, set as read, questions answered correctly)``."""


def make_result(path: Path, roll: str, set_code: str, correct: int) -> ScanResult:
    """A ScanResult whose first ``correct`` answers match an all-``A`` key."""
    return ScanResult(
        source_path=path,
        outcome=RecognitionOutcome.COMPLETE,
        registration=RegistrationStatus.REGISTERED,
        fields=(
            FieldView(
                zone_id="roll_number", label="Roll", field_type="numeric",
                value=roll, status="complete", needs_review=False, characters=(),
            ),
            FieldView(
                zone_id="set_code", label="Set", field_type="set_code",
                value=set_code, status="complete", needs_review=False, characters=(),
            ),
        ),
        answers=tuple(
            AnswerView(
                number=number, zone_id="q", value="A" if index < correct else "B",
                status="resolved", needs_review=False,
                top_fill=0.9, margin=0.4, confidence=0.9,
            )
            for index, number in enumerate(PLAN.numbers)
        ),
        identifier_zone_id="roll_number",
        set_code_zone_id="set_code",
    )


def write_attendance(path: Path, code: str) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = f"Set {code}"
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(row=HEADER_ROW, column=column, value=header)
    for offset, (roll, absent) in enumerate(ROSTERS[code]):
        row = HEADER_ROW + 1 + offset
        sheet.cell(row=row, column=1, value=offset + 1)
        sheet.cell(row=row, column=2, value=roll)
        sheet.cell(row=row, column=3, value=f"CANDIDATE {roll}")
        if absent:
            sheet.cell(row=row, column=4, value="ABSENT")
    workbook.save(path)
    workbook.close()
    return path


@dataclass
class World:
    """The project under test, and handles on everything in it."""

    session: ProjectSession
    batch_id: str
    set_ids: dict[str, str]
    rosters: dict[str, int]
    ids: dict[str, int] = field(default_factory=dict)

    @property
    def database(self):
        return self.session.database

    @property
    def scans_dir(self) -> Path:
        return self.session.project.layout.scans_original_dir

    def reconcile(self, code: str) -> dict[str, object]:
        reconciliation_store.reconcile_batch(self.database, self.rosters[code], self.batch_id)
        return {
            entry.candidate_id: entry
            for entry in reconciliation_store.list_entries(
                self.database, self.rosters[code], self.batch_id
            )
        }

    def score(self, code: str) -> dict[str, object]:
        scoring_store.score_batch(
            self.database, self.rosters[code], self.batch_id, TEMPLATE, computed_by=OPERATOR
        )
        return self.results(code)

    def results(self, code: str) -> dict[str, object]:
        return {
            item.candidate_id: item
            for item in scoring_store.list_results(
                self.database, self.rosters[code], self.batch_id, TEMPLATE
            )
        }

    def coordinator_passes(self) -> None:
        """What the Scan stage runs after every batch, in this order."""
        scan_lifecycle.sync_reimports(self.database, self.batch_id)
        review_store.sync_duplicate_identifiers(self.database, self.batch_id)
        review_store.sync_undefined_set_codes(self.database, self.batch_id)

    def add_processed_scan(
        self, name: str, roll: str, set_code: str, correct: int, *, content: bytes | None = None
    ) -> int:
        """Register a new file in the batch and record it as read."""
        path = self.scans_dir / name
        path.write_bytes(content if content is not None else f"scan:{name}".encode())
        assert batch_store.add_scans_to_batch(self.database, self.batch_id, [path]) == 1
        scan_id = batch_store.scan_ids_by_path(self.database, self.batch_id)[path]
        result = make_result(path, roll, set_code, correct)
        with self.database.session() as session:
            session.execute(
                update(BatchScan)
                .where(BatchScan.scan_id == scan_id)
                .values(
                    status="completed", identifier_value=roll, set_code_value=set_code,
                    result_json=json.dumps(result.to_dict()),
                    content_sha256=scan_provenance.hash_file(path),
                )
            )
        self.ids[name] = scan_id
        self.coordinator_passes()
        return scan_id


def build_world(session: ProjectSession, tmp_path: Path) -> World:
    set_ids = {
        code: project_sets.add_set(session.database, code, f"Paper {code}").set_id
        for code in "123"
    }
    rosters: dict[str, int] = {}
    for code, set_id in set_ids.items():
        path = write_attendance(tmp_path / f"set{code}.xlsx", code)
        assignment = set_attendance.assign_attendance_workbook(
            session.database, set_id, path, candidate_import.read_roster(path),
            imported_by=OPERATOR,
        )
        assert assignment.template_adopted, assignment.template_blocker
        rosters[code] = assignment.roster_id

    scans_dir = session.project.layout.scans_original_dir
    now = datetime.now(UTC)
    batch_id = batch_store.new_batch_id()
    with session.database.session() as db:
        db.add(
            ScanBatch(
                batch_id=batch_id, created_at=now, updated_at=now,
                source_folder=str(scans_dir), status="completed", total_scans=len(SCRIPTS),
            )
        )
        db.flush()
        for index, (name, roll, code, correct) in enumerate(SCRIPTS):
            path = scans_dir / name
            path.write_bytes(f"scan:{name}".encode())
            db.add(
                BatchScan(
                    batch_id=batch_id, batch_index=index, source_path=str(path),
                    filename=name, status="completed", identifier_value=roll,
                    set_code_value=code,
                    result_json=json.dumps(make_result(path, roll, code, correct).to_dict()),
                )
            )
    scan_provenance.compute_hashes_for_batch(session.database, batch_id)
    world = World(
        session=session, batch_id=batch_id, set_ids=set_ids, rosters=rosters,
        ids={
            path.name: scan_id
            for path, scan_id in batch_store.scan_ids_by_path(
                session.database, batch_id
            ).items()
        },
    )
    review_store.sync_duplicate_identifiers(session.database, batch_id)
    review_store.sync_undefined_set_codes(session.database, batch_id)
    for code in "123":
        stored = scoring_store.save_key(
            session.database, read_key("A" * PLAN.question_count, PLAN, code).to_key()
        )
        scoring_store.verify_key(session.database, stored.key_id, verified_by=OPERATOR)
        world.reconcile(code)
    return world


@pytest.fixture
def world(workspace: Path, tmp_path: Path) -> Iterator[World]:
    session = create_project(workspace, "Reject and Rescan")
    try:
        yield build_world(session, tmp_path)
    finally:
        if not session.is_closed:
            session.close()


def reject(world: World, name: str, **kwargs: object) -> None:
    kwargs.setdefault("reason", RejectionReason.FOLDED)
    scan_lifecycle.reject_scan(world.database, world.ids[name], reviewer=OPERATOR, **kwargs)


def lifecycle_events(world: World, name: str) -> list[AuditEvent]:
    with world.database.session() as session:
        return list(
            session.scalars(
                select(AuditEvent)
                .where(AuditEvent.entity_type == LIFECYCLE_ENTITY)
                .where(AuditEvent.entity_id == str(world.ids[name]))
                .order_by(AuditEvent.event_id)
            ).all()
        )


# ----------------------------------------------------------------------
# Rejecting
# ----------------------------------------------------------------------
class TestRejecting:
    def test_an_active_scan_can_be_rejected(self, world):
        reject(world, "s2b.png", note="Folded through the answer grid")
        case = scan_lifecycle.get_case(world.database, world.ids["s2b.png"])
        assert case is not None
        assert case.state is LifecycleState.REJECTED_PENDING_RESCAN
        assert case.state.label == "REJECTED — RESCAN REQUIRED"
        assert case.reason is RejectionReason.FOLDED
        assert case.identity == "200122"
        assert case.set_code == "2"
        assert case.rejected_by == OPERATOR
        assert world.ids["s2b.png"] in scan_lifecycle.ineligible_scan_ids(
            world.database, world.batch_id
        )

    def test_nothing_is_deleted_by_rejecting(self, world):
        path = world.scans_dir / "s2b.png"
        before = path.read_bytes()
        reject(world, "s2b.png")
        assert path.read_bytes() == before
        with world.database.session() as session:
            row = session.get(BatchScan, world.ids["s2b.png"])
            assert row is not None and row.result_json and row.status == "completed"

    def test_the_rejection_is_audited(self, world):
        reject(world, "s2b.png", reason=RejectionReason.OTHER, note="Coffee stain")
        events = lifecycle_events(world, "s2b.png")
        assert [event.action for event in events] == [LifecycleAction.REJECTED.value]
        event = events[0]
        assert event.reviewer == OPERATOR
        assert event.previous_value == "active"
        assert event.new_value == "rejected_pending_rescan"
        assert event.reason_code == "other"
        assert event.reason_text == "Coffee stain"
        assert event.batch_id == world.batch_id and event.scan_id == world.ids["s2b.png"]
        assert "200122" in event.detail

    def test_rejecting_needs_a_reviewer_and_a_note_for_other(self, world):
        with pytest.raises(ReviewError):
            scan_lifecycle.reject_scan(
                world.database, world.ids["s2b.png"], reviewer=" ",
                reason=RejectionReason.FOLDED,
            )
        with pytest.raises(LifecycleError):
            reject(world, "s2b.png", reason=RejectionReason.OTHER)
        assert scan_lifecycle.state_of(world.database, world.ids["s2b.png"]) is (
            LifecycleState.ACTIVE
        )

    def test_a_scan_cannot_be_rejected_twice(self, world):
        reject(world, "s2b.png")
        with pytest.raises(LifecycleError):
            reject(world, "s2b.png")

    def test_the_rejected_scan_is_in_the_rescan_queue(self, world):
        reject(world, "s2b.png")
        cases = scan_lifecycle.list_cases(world.database, world.batch_id)
        assert [case.source_name for case in cases] == ["s2b.png"]
        assert scan_lifecycle.count_cases(world.database, world.batch_id).outstanding == 1

    def test_the_rejected_scan_stops_being_a_valid_script(self, world):
        assert world.reconcile("2")["200122"].status is ReconciliationStatus.MATCHED
        reject(world, "s2b.png")
        entries = world.reconcile("2")
        entry = entries["200122"]
        assert entry.status is ReconciliationStatus.RESCAN_REQUIRED
        assert entry.status.label == "Script received but rejected — rescan required"
        assert entry.script_count == 0
        assert [view.script.rejected for view in entry.scripts] == [True]
        # Received and rejected is not the same as never received.
        assert ReconciliationStatus.PRESENT_WITHOUT_SCRIPT not in {
            item.status for key, item in entries.items() if key == "200122"
        }
        scope = reconciliation_store.script_scope(
            world.database, world.rosters["2"], world.batch_id
        )
        assert scope.rescan_required == 1
        assert scope.in_set == 2  # s2a and s2c; s2b no longer counts

    def test_the_stored_reconciliation_is_refreshed_by_the_rejection(self, world):
        reject(world, "s2b.png")
        entries = {
            entry.candidate_id: entry
            for entry in reconciliation_store.list_entries(
                world.database, world.rosters["2"], world.batch_id
            )
        }
        assert entries["200122"].status is ReconciliationStatus.RESCAN_REQUIRED

    def test_other_sets_are_untouched(self, world):
        before = {code: set(world.reconcile(code)) for code in "13"}
        reject(world, "s2b.png")
        after = {code: set(world.reconcile(code)) for code in "13"}
        assert before == after
        assert not any(key.startswith("2") for key in after["1"] | after["3"])

    def test_the_rejection_persists_after_reopening(self, world):
        reject(world, "s2b.png")
        root = world.session.root
        scan_id = world.ids["s2b.png"]
        world.session.close()
        with open_project(root) as reopened:
            assert scan_lifecycle.state_of(reopened.database, scan_id) is (
                LifecycleState.REJECTED_PENDING_RESCAN
            )

    def test_the_rejection_survives_reprocessing_and_resume(self, world):
        reject(world, "s2b.png")
        scan_id = world.ids["s2b.png"]
        assert batch_store.mark_for_reprocessing(
            world.database, world.batch_id, [scan_id], reason="template edited"
        ) == 1
        assert batch_store.recover_interrupted(world.database) == (0, 0)
        paths = batch_store.resumable_scans(world.database, world.batch_id)
        assert world.scans_dir / "s2b.png" in paths
        assert scan_lifecycle.state_of(world.database, scan_id) is (
            LifecycleState.REJECTED_PENDING_RESCAN
        )


class TestScoringAndResults:
    def test_a_rejected_scan_is_excluded_from_scoring(self, world):
        reject(world, "s2b.png")
        world.reconcile("2")
        results = world.score("2")
        assert results["200122"].status is ResultStatus.BLOCKED
        assert results["200122"].scan_id is None
        inputs = scoring_store.gather_inputs(
            world.database, world.rosters["2"], world.batch_id, TEMPLATE
        )
        assert world.ids["s2b.png"] not in inputs.answers

    def test_a_mark_made_from_a_scan_goes_stale_when_it_is_rejected(self, world):
        results = world.score("2")
        assert results["200122"].status is ResultStatus.SCORED
        reject(world, "s2b.png")
        stale = world.results("2")["200122"]
        assert StaleReason.RECONCILIATION in stale.stale_reasons

    def test_final_export_of_a_set_with_an_outstanding_rescan_is_refused(self, world, tmp_path):
        world.score("3")
        reject(world, "s3b.png")
        world.score("3")
        outcome = report_store.generate_for_set(
            world.database, world.set_ids["3"], world.batch_id, TEMPLATE,
            project_name="RR", output_dir=tmp_path / "out", computed_by=OPERATOR,
        )
        assert outcome.status == "blocked"
        assert outcome.readiness is not None
        issues = outcome.readiness.by_kind(ReadinessIssueKind.RESCAN_OUTSTANDING)
        assert [item.roll for item in issues] == ["300122"]
        assert outcome.readiness.only_acknowledgeable_blocks

    def test_an_acknowledged_incomplete_export_is_marked_and_audited(self, world, tmp_path):
        world.score("3")
        reject(world, "s3b.png")
        world.score("3")
        outcome = report_store.generate_for_set(
            world.database, world.set_ids["3"], world.batch_id, TEMPLATE,
            project_name="RR", output_dir=tmp_path / "out", computed_by=OPERATOR,
            acknowledge_incomplete=True,
        )
        assert outcome.ok, outcome.warnings
        assert outcome.warnings[0].startswith("INCOMPLETE RESULTS")
        # The workbook itself says so, in its Processing Log.
        workbook = openpyxl.load_workbook(outcome.output_path)
        try:
            log = "\n".join(
                str(cell)
                for row in workbook[workbook.sheetnames[-1]].iter_rows(values_only=True)
                for cell in row
                if cell
            )
        finally:
            workbook.close()
        assert "INCOMPLETE RESULTS" in log
        # The rejected candidate has no mark in it.
        exports = scan_lifecycle.incomplete_exports(world.database)
        assert len(exports) == 1
        assert exports[0].reviewer == OPERATOR
        with world.database.session() as session:
            event = session.scalars(
                select(AuditEvent).where(AuditEvent.entity_type == EXPORT_ENTITY)
            ).one()
        assert event.action == "export_incomplete"
        assert event.entity_id == "3"
        assert event.new_value == "1"

    def test_the_acknowledgement_needs_a_named_operator(self, world, tmp_path):
        reject(world, "s3b.png")
        world.score("3")
        outcome = report_store.generate_for_set(
            world.database, world.set_ids["3"], world.batch_id, TEMPLATE,
            project_name="RR", output_dir=tmp_path / "out", computed_by="",
            acknowledge_incomplete=True,
        )
        assert outcome.status == "blocked"
        assert scan_lifecycle.incomplete_exports(world.database) == ()

    def test_acknowledging_does_not_lift_any_other_block(self, world, tmp_path):
        reject(world, "s2b.png")
        world.score("2")
        outcome = report_store.generate_for_set(
            world.database, world.set_ids["2"], world.batch_id, TEMPLATE,
            project_name="RR", output_dir=tmp_path / "out", computed_by=OPERATOR,
            acknowledge_incomplete=True,
        )
        # Set 2 also has a candidate with no script (200123) - still blocking.
        assert outcome.status == "blocked"
        assert not outcome.readiness.only_acknowledgeable_blocks

    def test_a_rejected_scan_with_unknown_identity_is_reported_for_every_set(self, world):
        reject(world, "s_blur.png")
        case = scan_lifecycle.get_case(world.database, world.ids["s_blur.png"])
        assert case is not None and case.identity == "" and case.set_code == ""
        for code in "123":
            assert [item.source_name for item in scan_lifecycle.outstanding_for_set(
                world.database, world.batch_id, code
            )] == ["s_blur.png"]
        readiness = report_store.check_readiness(
            world.database, world.rosters["3"], world.batch_id, TEMPLATE, "3",
            for_final_export=True,
        )
        messages = [
            item.message for item in readiness.by_kind(ReadinessIssueKind.RESCAN_OUTSTANDING)
        ]
        assert any("s_blur.png" in item for item in messages)


class TestDeclaredIdentity:
    def test_a_declared_student_id_files_the_case_under_that_candidate(self, world):
        # s2c.png was read as 20123 - one digit short. The operator can read
        # 200123 off the paper.
        assert world.reconcile("2")["20123"].status is ReconciliationStatus.UNKNOWN_ID
        reject(world, "s2c.png", reason=RejectionReason.ID_UNREADABLE,
               declared_candidate_id="200123")
        entries = world.reconcile("2")
        assert "20123" not in entries
        assert entries["200123"].status is ReconciliationStatus.RESCAN_REQUIRED

    def test_a_declared_identity_is_not_the_scans_student_id(self, world):
        reject(world, "s2c.png", declared_candidate_id="200123")
        identifiers = review_store.effective_identifiers(world.database, world.batch_id)
        assert identifiers[world.ids["s2c.png"]].value == "20123"
        case = scan_lifecycle.get_case(world.database, world.ids["s2c.png"])
        assert case is not None
        assert case.identity == "200123"
        assert case.recognised_candidate_id == "20123"

    def test_a_declared_set_must_be_a_defined_set(self, world):
        with pytest.raises(LifecycleError):
            reject(world, "s_blur.png", declared_set_code="9")
        reject(world, "s_blur.png", declared_candidate_id="300122", declared_set_code="3")
        case = scan_lifecycle.get_case(world.database, world.ids["s_blur.png"])
        assert case is not None and case.set_code == "3"


# ----------------------------------------------------------------------
# Replacement
# ----------------------------------------------------------------------
class TestReplacement:
    def test_a_rescan_is_suggested_by_student_id(self, world):
        reject(world, "s2b.png")
        rescan = world.add_processed_scan("IMG_0042.jpg", "200122", "2", 20)
        found = scan_lifecycle.replacement_candidates(world.database, world.ids["s2b.png"])
        assert [item.scan_id for item in found] == [rescan]
        assert found[0].set_code_agrees is True

    def test_file_names_are_never_evidence(self, world):
        reject(world, "s2b.png")
        # Named like the original, but another candidate's sheet.
        world.add_processed_scan("s2b_rescan.png", "200121", "2", 20)
        assert scan_lifecycle.replacement_candidates(
            world.database, world.ids["s2b.png"]
        ) == ()

    def test_a_rescan_is_never_linked_silently(self, world):
        reject(world, "s2b.png")
        world.add_processed_scan("IMG_0042.jpg", "200122", "2", 20)
        world.reconcile("2")
        assert scan_lifecycle.state_of(world.database, world.ids["s2b.png"]) is (
            LifecycleState.REJECTED_PENDING_RESCAN
        )
        assert [event.action for event in lifecycle_events(world, "s2b.png")] == [
            "rejected"
        ]

    def test_confirming_makes_the_replacement_the_only_valid_script(self, world):
        reject(world, "s2b.png")
        rescan = world.add_processed_scan("IMG_0042.jpg", "200122", "2", 20)
        scan_lifecycle.confirm_replacement(
            world.database, world.ids["s2b.png"], rescan, reviewer=OPERATOR
        )
        entry = world.reconcile("2")["200122"]
        assert entry.status is ReconciliationStatus.MATCHED
        assert [view.script.scan_id for view in entry.scripts] == [rescan]
        result = world.score("2")["200122"]
        assert result.status is ResultStatus.SCORED
        assert result.scan_id == rescan
        # The replacement's answers, not the original's.
        assert result.correct_count == 20

    def test_the_confirmed_replacement_is_not_an_unresolved_duplicate(self, world):
        reject(world, "s2b.png")
        rescan = world.add_processed_scan("IMG_0042.jpg", "200122", "2", 20)
        scan_lifecycle.confirm_replacement(
            world.database, world.ids["s2b.png"], rescan, reviewer=OPERATOR
        )
        duplicates = [
            item
            for item in review_store.list_conflicts(world.database, world.batch_id)
            if item.conflict_type is ConflictType.IDENTIFIER_DUPLICATE
        ]
        assert duplicates == []
        assert ReconciliationStatus.DUPLICATE_SCRIPT not in {
            entry.status for entry in world.reconcile("2").values()
        }

    def test_the_replacement_is_audited_on_both_scans(self, world):
        reject(world, "s2b.png")
        rescan = world.add_processed_scan("IMG_0042.jpg", "200122", "2", 20)
        scan_lifecycle.confirm_replacement(
            world.database, world.ids["s2b.png"], rescan, reviewer=OPERATOR
        )
        original = lifecycle_events(world, "s2b.png")
        assert [event.action for event in original] == ["rejected", "replaced"]
        assert original[-1].previous_value == "rejected_pending_rescan"
        assert original[-1].new_value == "superseded_by_replacement"
        assert "IMG_0042.jpg" in original[-1].detail
        linked = lifecycle_events(world, "IMG_0042.jpg")
        assert [event.action for event in linked] == ["linked_replacement"]

    def test_the_replacement_persists_after_reopening(self, world):
        reject(world, "s2b.png")
        rescan = world.add_processed_scan("IMG_0042.jpg", "200122", "2", 20)
        scan_lifecycle.confirm_replacement(
            world.database, world.ids["s2b.png"], rescan, reviewer=OPERATOR
        )
        root, original = world.session.root, world.ids["s2b.png"]
        world.session.close()
        with open_project(root) as reopened:
            case = scan_lifecycle.get_case(reopened.database, original)
            assert case is not None
            assert case.state is LifecycleState.SUPERSEDED_BY_REPLACEMENT
            assert case.replacement_scan_id == rescan
            assert case.replacement_name == "IMG_0042.jpg"

    @pytest.mark.parametrize(
        "problem", ["itself", "same_bytes", "unread", "already", "rejected"]
    )
    def test_confirmation_refuses_what_cannot_be_a_rescan(self, world, problem):
        reject(world, "s2b.png")
        original = world.ids["s2b.png"]
        if problem == "itself":
            target = original
        elif problem == "same_bytes":
            target = world.add_processed_scan(
                "copy.png", "200122", "2", 15, content=b"scan:s2b.png"
            )
        elif problem == "unread":
            path = world.scans_dir / "pending.png"
            path.write_bytes(b"pending")
            batch_store.add_scans_to_batch(world.database, world.batch_id, [path])
            target = batch_store.scan_ids_by_path(world.database, world.batch_id)[path]
        elif problem == "already":
            reject(world, "s2a.png")
            target = world.add_processed_scan("r.png", "200121", "2", 20)
            scan_lifecycle.confirm_replacement(
                world.database, world.ids["s2a.png"], target, reviewer=OPERATOR
            )
        else:
            target = world.ids["s3a.png"]
            reject(world, "s3a.png")
        with pytest.raises(LifecycleError):
            scan_lifecycle.confirm_replacement(
                world.database, original, target, reviewer=OPERATOR
            )
        assert scan_lifecycle.state_of(world.database, original) is (
            LifecycleState.REJECTED_PENDING_RESCAN
        )

    def test_an_unknown_identity_case_is_linked_by_hand(self, world):
        reject(world, "s_blur.png")
        assert scan_lifecycle.replacement_candidates(
            world.database, world.ids["s_blur.png"]
        ) == ()
        rescan = world.add_processed_scan("blur_again.png", "300122", "3", 9)
        choices = scan_lifecycle.association_choices(world.database, world.ids["s_blur.png"])
        assert rescan in {item.scan_id for item in choices}
        scan_lifecycle.confirm_replacement(
            world.database, world.ids["s_blur.png"], rescan, reviewer=OPERATOR
        )
        assert scan_lifecycle.state_of(world.database, world.ids["s_blur.png"]) is (
            LifecycleState.SUPERSEDED_BY_REPLACEMENT
        )

    def test_an_unread_scan_in_another_batch_is_still_refused(self, world, tmp_path):
        # Another batch is fine now (tests/integration/
        # test_reject_and_rescan_cross_batch.py); an unread scan never is.
        reject(world, "s2b.png")
        other = batch_store.create_batch(
            world.database, [tmp_path / "elsewhere.png"], identity=batch_store.BatchIdentity()
        )
        target = batch_store.scan_ids_by_path(world.database, other)[tmp_path / "elsewhere.png"]
        with pytest.raises(LifecycleError):
            scan_lifecycle.confirm_replacement(
                world.database, world.ids["s2b.png"], target, reviewer=OPERATOR
            )


class TestReimports:
    def test_importing_the_rejected_bytes_again_does_not_resurrect_them(self, world):
        reject(world, "s2b.png")
        again = world.add_processed_scan(
            "s2b_again.png", "200122", "2", 15, content=b"scan:s2b.png"
        )
        assert scan_lifecycle.state_of(world.database, again) is (
            LifecycleState.REIMPORT_OF_REJECTED
        )
        entry = world.reconcile("2")["200122"]
        assert entry.status is ReconciliationStatus.RESCAN_REQUIRED
        assert entry.script_count == 0
        # Not offered as a replacement either.
        assert scan_lifecycle.replacement_candidates(
            world.database, world.ids["s2b.png"]
        ) == ()

    def test_a_reimport_keeps_the_existing_content_hash_provenance(self, world):
        reject(world, "s2b.png")
        again = world.add_processed_scan(
            "s2b_again.png", "200122", "2", 15, content=b"scan:s2b.png"
        )
        groups = scan_provenance.duplicate_groups(world.database, world.batch_id)
        assert any(
            set(group.scan_ids) == {world.ids["s2b.png"], again} for group in groups
        )
        events = lifecycle_events(world, "s2b_again.png")
        assert [event.action for event in events] == ["reimport_linked"]
        assert str(world.ids["s2b.png"]) in events[0].detail
        with pytest.raises(LifecycleError):
            scan_lifecycle.undo_reject(world.database, again, reviewer=OPERATOR)

    def test_a_reimport_after_replacement_does_not_reactivate_the_original(self, world):
        reject(world, "s2b.png")
        rescan = world.add_processed_scan("IMG_0042.jpg", "200122", "2", 20)
        scan_lifecycle.confirm_replacement(
            world.database, world.ids["s2b.png"], rescan, reviewer=OPERATOR
        )
        world.add_processed_scan("s2b_again.png", "200122", "2", 15, content=b"scan:s2b.png")
        entry = world.reconcile("2")["200122"]
        assert [view.script.scan_id for view in entry.scripts] == [rescan]
        assert scan_lifecycle.state_of(world.database, world.ids["s2b.png"]) is (
            LifecycleState.SUPERSEDED_BY_REPLACEMENT
        )


# ----------------------------------------------------------------------
# Undo
# ----------------------------------------------------------------------
class TestUndo:
    def test_undo_before_replacement_restores_the_scan_exactly(self, world):
        duplicate = world.add_processed_scan("s2a_dup.png", "200121", "2", 3)
        conflicts_before = {
            item.conflict_id: item.state
            for item in review_store.list_conflicts(
                world.database, world.batch_id,
                filters=review_store.ConflictFilter(scan_id=duplicate),
            )
        }
        assert conflicts_before, "a genuine duplicate is detected"
        assert world.reconcile("2")["200121"].status is ReconciliationStatus.DUPLICATE_SCRIPT

        reject(world, "s2a_dup.png")
        assert review_store.list_conflicts(
            world.database, world.batch_id,
            filters=review_store.ConflictFilter(scan_id=duplicate),
        ) == ()
        # ...but kept, not withdrawn.
        with world.database.session() as session:
            kept = session.scalars(
                select(ReviewConflict).where(ReviewConflict.scan_id == duplicate)
            ).all()
            assert {row.state for row in kept} == {ConflictState.OPEN.value}
        assert world.reconcile("2")["200121"].status is ReconciliationStatus.MATCHED

        scan_lifecycle.undo_reject(world.database, duplicate, reviewer=OPERATOR)
        after = {
            item.conflict_id: item.state
            for item in review_store.list_conflicts(
                world.database, world.batch_id,
                filters=review_store.ConflictFilter(scan_id=duplicate),
            )
        }
        assert after == conflicts_before
        assert world.reconcile("2")["200121"].status is ReconciliationStatus.DUPLICATE_SCRIPT
        assert [event.action for event in lifecycle_events(world, "s2a_dup.png")] == [
            "rejected", "reject_undone",
        ]

    def test_undo_after_replacement_is_refused(self, world):
        reject(world, "s2b.png")
        rescan = world.add_processed_scan("IMG_0042.jpg", "200122", "2", 20)
        scan_lifecycle.confirm_replacement(
            world.database, world.ids["s2b.png"], rescan, reviewer=OPERATOR
        )
        with pytest.raises(LifecycleError):
            scan_lifecycle.undo_reject(world.database, world.ids["s2b.png"], reviewer=OPERATOR)
        assert [
            view.script.scan_id
            for view in world.reconcile("2")["200122"].scripts
            if view.counts_as_a_script
        ] == [rescan]

    def test_removing_the_link_then_undoing_never_leaves_two_silent_scripts(self, world):
        reject(world, "s2b.png")
        rescan = world.add_processed_scan("IMG_0042.jpg", "200122", "2", 20)
        original = world.ids["s2b.png"]
        scan_lifecycle.confirm_replacement(world.database, original, rescan, reviewer=OPERATOR)
        scan_lifecycle.remove_replacement(world.database, original, reviewer=OPERATOR)
        # Back to awaiting a rescan - not reactivated.
        assert scan_lifecycle.state_of(world.database, original) is (
            LifecycleState.REJECTED_PENDING_RESCAN
        )
        scan_lifecycle.undo_reject(world.database, original, reviewer=OPERATOR)
        # Now two active scans share one ID: the ordinary duplicate machinery
        # says so, everywhere, and scoring refuses to choose.
        assert world.reconcile("2")["200122"].status is ReconciliationStatus.DUPLICATE_SCRIPT
        duplicates = {
            item.scan_id
            for item in review_store.list_conflicts(world.database, world.batch_id)
            if item.conflict_type is ConflictType.IDENTIFIER_DUPLICATE
        }
        assert {original, rescan} <= duplicates
        result = world.score("2")["200122"]
        assert result.status is ResultStatus.BLOCKED
        assert BlockReason.DUPLICATE_SCRIPT in {item.reason for item in result.blocks} or (
            result.scan_id is None
        )
        actions = [event.action for event in lifecycle_events(world, "s2b.png")]
        assert actions == ["rejected", "replaced", "replacement_removed", "reject_undone"]


# ----------------------------------------------------------------------
# Purge Rejects
# ----------------------------------------------------------------------
class TestPurge:
    def _replace(self, world: World) -> int:
        reject(world, "s2b.png")
        rescan = world.add_processed_scan("IMG_0042.jpg", "200122", "2", 20)
        scan_lifecycle.confirm_replacement(
            world.database, world.ids["s2b.png"], rescan, reviewer=OPERATOR
        )
        return rescan

    def test_a_scan_awaiting_rescan_is_never_eligible(self, world):
        reject(world, "s2b.png")
        plan = scan_lifecycle.plan_purge(world.database, world.session.root)
        assert plan.eligible == ()
        assert plan.pending == 1
        outcome = scan_lifecycle.execute_purge(
            world.database, world.session.root, mode=PurgeMode.DELETE, reviewer=OPERATOR
        )
        assert outcome.files_moved == 0
        assert (world.scans_dir / "s2b.png").exists()

    def test_a_superseded_original_becomes_eligible(self, world):
        self._replace(world)
        plan = scan_lifecycle.plan_purge(world.database, world.session.root)
        assert [item.case.source_name for item in plan.eligible] == ["s2b.png"]
        assert plan.eligible[0].files[0].role == "source"
        assert plan.size_bytes == len(b"scan:s2b.png")

    def test_quarantine_moves_only_the_original_and_keeps_history(self, world):
        rescan = self._replace(world)
        outcome = scan_lifecycle.execute_purge(
            world.database, world.session.root, mode=PurgeMode.QUARANTINE, reviewer=OPERATOR
        )
        assert outcome.processed == 1 and outcome.files_moved == 1
        assert not (world.scans_dir / "s2b.png").exists()
        assert (world.scans_dir / "IMG_0042.jpg").exists()
        moved = list(scan_lifecycle.quarantine_root(world.session.root).rglob("s2b.png"))
        assert len(moved) == 1 and moved[0].read_bytes() == b"scan:s2b.png"
        case = scan_lifecycle.get_case(world.database, world.ids["s2b.png"])
        assert case is not None
        assert case.state is LifecycleState.SUPERSEDED_BY_REPLACEMENT
        assert case.file_state is FileState.QUARANTINED
        assert case.replacement_scan_id == rescan
        assert case.source_name == "s2b.png" and case.content_sha256
        with world.database.session() as session:
            assert session.get(BatchScan, world.ids["s2b.png"]) is not None
        assert [event.action for event in lifecycle_events(world, "s2b.png")][-1] == (
            "image_quarantined"
        )
        # The replacement still counts.
        assert world.score("2")["200122"].scan_id == rescan

    def test_deleting_after_quarantine_purges_and_the_record_remains(self, world):
        self._replace(world)
        root = world.session.root
        scan_lifecycle.execute_purge(
            world.database, root, mode=PurgeMode.QUARANTINE, reviewer=OPERATOR
        )
        scan_lifecycle.execute_purge(
            world.database, root, mode=PurgeMode.DELETE, reviewer=OPERATOR
        )
        assert list(scan_lifecycle.quarantine_root(root).rglob("s2b.png")) == []
        case = scan_lifecycle.get_case(world.database, world.ids["s2b.png"])
        assert case is not None and case.file_state is FileState.PURGED
        history = scan_lifecycle.lifecycle_history(world.database, world.ids["s2b.png"])
        assert [item.action for item in history] == [
            LifecycleAction.REJECTED, LifecycleAction.REPLACED,
            LifecycleAction.IMAGE_QUARANTINED, LifecycleAction.IMAGE_PURGED,
        ]
        # Readable without the image, and nothing more to purge.
        assert "s2b.png" in history[-1].detail
        assert scan_lifecycle.plan_purge(world.database, root).eligible == ()
        # A purged image is not a missing source scan.
        report = project_health.full_check(world.database, root)
        assert "SOURCE_SCANS_MISSING" not in {item.code for item in report.issues}
        with pytest.raises(LifecycleError):
            scan_lifecycle.remove_replacement(
                world.database, world.ids["s2b.png"], reviewer=OPERATOR
            )

    def test_a_file_outside_project_storage_is_never_removed(self, world, tmp_path):
        external = tmp_path / "scanner_share" / "ext.png"
        external.parent.mkdir()
        external.write_bytes(b"external scan")
        with world.database.session() as session:
            session.execute(
                update(BatchScan)
                .where(BatchScan.scan_id == world.ids["s2b.png"])
                .values(
                    source_path=str(external),
                    content_sha256=scan_provenance.hash_file(external),
                )
            )
        self._replace(world)
        plan = scan_lifecycle.plan_purge(world.database, world.session.root)
        assert plan.eligible[0].files == ()
        assert "outside the project" in plan.eligible[0].skipped[0]
        scan_lifecycle.execute_purge(
            world.database, world.session.root, mode=PurgeMode.DELETE, reviewer=OPERATOR
        )
        assert external.exists()

    def test_a_file_another_scan_still_uses_is_not_removed(self, world):
        self._replace(world)
        with world.database.session() as session:
            session.execute(
                update(BatchScan)
                .where(BatchScan.scan_id == world.ids["s3a.png"])
                .values(output_path=str(world.scans_dir / "s2b.png"), copied=True)
            )
        plan = scan_lifecycle.plan_purge(world.database, world.session.root)
        assert plan.eligible[0].files == ()
        assert "still used by another scan" in plan.eligible[0].skipped[0]

    def test_a_changed_file_is_not_removed(self, world):
        self._replace(world)
        (world.scans_dir / "s2b.png").write_bytes(b"something else entirely")
        plan = scan_lifecycle.plan_purge(world.database, world.session.root)
        assert plan.eligible[0].files == ()
        assert "changed since" in plan.eligible[0].skipped[0]


class TestOwnedPath:
    def test_only_files_inside_managed_storage_are_owned(self, world, tmp_path):
        root = world.session.root
        inside = world.scans_dir / "s1a.png"
        assert scan_lifecycle.owned_path(root, str(inside)) == inside.resolve()
        outside = tmp_path / "loose.png"
        outside.write_bytes(b"x")
        assert scan_lifecycle.owned_path(root, str(outside)) is None
        # The project root itself is not scan storage - it holds the database.
        database_file = world.session.project.layout.database_file
        assert scan_lifecycle.owned_path(root, str(database_file)) is None
        assert scan_lifecycle.owned_path(root, str(world.scans_dir)) is None

    @pytest.mark.parametrize(
        "candidate",
        ["", "stress:1-000001", "relative/scan.png"],
    )
    def test_malformed_paths_are_refused(self, world, candidate):
        assert scan_lifecycle.owned_path(world.session.root, candidate) is None

    def test_parent_directory_escapes_are_refused(self, world, tmp_path):
        target = tmp_path / "victim.png"
        target.write_bytes(b"do not delete")
        sneaky = f"{world.scans_dir}{os.sep}..{os.sep}..{os.sep}{target.name}"
        assert scan_lifecycle.owned_path(world.session.root, sneaky) is None

    def test_a_link_inside_storage_pointing_outside_is_refused(self, world, tmp_path):
        target = tmp_path / "victim.png"
        target.write_bytes(b"do not delete")
        link = world.scans_dir / "link.png"
        try:
            link.symlink_to(target)
        except OSError:
            pytest.skip("creating a symbolic link needs a privilege this machine lacks")
        assert scan_lifecycle.owned_path(world.session.root, str(link)) is None

    def test_a_junction_inside_storage_pointing_outside_is_refused(self, world, tmp_path):
        winapi = pytest.importorskip("_winapi")
        outside = tmp_path / "elsewhere"
        outside.mkdir()
        victim = outside / "victim.png"
        victim.write_bytes(b"do not delete")
        junction = world.scans_dir / "junction"
        winapi.CreateJunction(str(outside), str(junction))
        try:
            assert (junction / "victim.png").exists()
            assert scan_lifecycle.owned_path(
                world.session.root, str(junction / "victim.png")
            ) is None
        finally:
            junction.rmdir()
        assert victim.exists()


# ----------------------------------------------------------------------
# Migration
# ----------------------------------------------------------------------
class TestMigrationOntoAnExistingProject:
    def test_a_version_10_project_upgrades_and_keeps_its_scans(self, world):
        root = world.session.root
        db_path = world.session.database.path
        world.session.close()
        connection = sqlite3.connect(db_path)
        try:
            connection.executescript(
                """
                DROP TABLE IF EXISTS scan_rejection;
                DELETE FROM schema_migration WHERE version >= 11;
                """
            )
            connection.commit()
            scans_before = connection.execute("SELECT COUNT(*) FROM batch_scan").fetchone()[0]
        finally:
            connection.close()

        with open_project(root) as reopened:
            assert reopened.database.schema_version == SCHEMA_VERSION
            with reopened.database.session() as session:
                tables = {
                    row[0]
                    for row in session.execute(
                        text("SELECT name FROM sqlite_master WHERE type='table'")
                    ).all()
                }
                scans_after = session.execute(text("SELECT COUNT(*) FROM batch_scan")).scalar()
            assert "scan_rejection" in tables
            assert scans_after == scans_before
            # Every existing scan is active: no row means never rejected.
            batch_id = batch_store.list_batches(reopened.database, limit=1)[0].batch_id
            assert scan_lifecycle.lifecycle_states(reopened.database, batch_id) == {}

    def test_the_migration_is_safe_to_rerun_after_a_failed_attempt(self, tmp_path):
        from sqlalchemy import create_engine

        from omr_scanner.database.migrations import _migration_011_reject_and_rescan

        engine = create_engine(f"sqlite:///{tmp_path / 'x.sqlite'}")
        from omr_scanner.database.models import Base

        with engine.begin() as connection:
            Base.metadata.create_all(
                connection,
                tables=[
                    Base.metadata.tables["scan_batch"],
                    Base.metadata.tables["batch_scan"],
                ],
            )
            _migration_011_reject_and_rescan(connection)
            _migration_011_reject_and_rescan(connection)
            names = {
                row[0]
                for row in connection.execute(
                    text("SELECT name FROM sqlite_master WHERE type='table'")
                ).all()
            }
        engine.dispose()
        assert "scan_rejection" in names
