"""Per-Set attendance, end to end (Part 2, §24).

Scope:
    The real services against a real project database and real `.xlsx` /
    `.csv` files on disk - `candidate_import`, `set_attendance`,
    `reconciliation_store` and `report_store` wired exactly as the
    application wires them. Nothing here hand-builds a
    :class:`ReconciliationEntry`; the classifications asserted below are the
    ones :func:`reconciliation_store.reconcile_batch` actually produced.

One test per §24 item, named so that a failure reads as the requirement it
broke rather than as the function that raised.

The attendance workbooks these tests build carry an institution title and a
post name above the header row, because that is what a real examination
office's sheet looks like - and because it is the shape that has to work for
the same file to serve as the result template afterwards.
"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from typing import TYPE_CHECKING

import openpyxl
import pytest
from tests.conftest import build_answer_sheet_template

from omr_scanner.domain.reconciliation import AttendanceState, ReconciliationIssue
from omr_scanner.services import (
    batch_store,
    candidate_import,
    create_project,
    open_project,
    project_sets,
    reconciliation_store,
    report_store,
    set_attendance,
)
from omr_scanner.services.answer_key import plan_for
from omr_scanner.services.recognition_models import (
    AnswerView,
    FieldView,
    RecognitionOutcome,
    RegistrationStatus,
    ScanResult,
)

if TYPE_CHECKING:
    from pathlib import Path

    from omr_scanner.database import ProjectDatabase

OPERATOR = "Dr. Rahman"

EXAM_NAME = "Recruitment Exam, Bangladesh Submarine Cable Regulatory Authority"

SET_DEFINITIONS: tuple[tuple[str, str], ...] = (
    ("10", "Name of Post: Assistant Engineer (Electrical)"),
    ("11", "Name of Post: Assistant Engineer (Civil)"),
    ("12", "Name of Post: Assistant Engineer (Mechanical)"),
)

HEADERS: tuple[str, ...] = ("Sl.No.", "Roll No.", "Name", "Total", "Merit")
HEADER_ROW = 3


# ----------------------------------------------------------------------
# Building the files an operator would actually choose
# ----------------------------------------------------------------------
def write_attendance_workbook(
    path: Path,
    rolls: list[str],
    *,
    post: str = "Name of Post: Assistant Engineer (Electrical)",
    absent: dict[str, str] | None = None,
) -> Path:
    """Write a decorated attendance workbook and return its path.

    Args:
        path: Where to write it.
        rolls: The registered candidates, in the office's own order.
        post: The post line above the table.
        absent: ``{roll: token}`` for candidates recorded absent - the token
            is written verbatim so a test can use ``ABS``, ``absent``,
            ``Abs`` and so on.

    A present candidate's marks cell is left **blank**, which is what a real
    attendance sheet looks like before the examination and is what
    :func:`candidate_import.attendance_from_cell` reads as "no absence was
    recorded".
    """
    marked_absent = absent or {}
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    sheet.cell(row=1, column=1, value="Bangladesh Submarine Cable Regulatory Authority")
    sheet.cell(row=2, column=1, value=post)
    for column, text in enumerate(HEADERS, start=1):
        sheet.cell(row=HEADER_ROW, column=column, value=text)
    for offset, roll in enumerate(rolls):
        row = HEADER_ROW + 1 + offset
        sheet.cell(row=row, column=1, value=offset + 1)
        sheet.cell(row=row, column=2, value=roll)
        sheet.cell(row=row, column=3, value=f"CANDIDATE {roll}")
        if roll in marked_absent:
            sheet.cell(row=row, column=4, value=marked_absent[roll])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def write_attendance_csv(path: Path, rolls: list[str]) -> Path:
    """Write the CSV equivalent - usable for reconciliation, never a template."""
    lines = ["Roll No.,Name,Total"]
    lines.extend(f"{roll},CANDIDATE {roll}," for roll in rolls)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def assign(
    database: ProjectDatabase, set_id: str, path: Path
) -> set_attendance.AttendanceAssignment:
    """Read a roster file and assign it to a set, exactly as the GUI does."""
    validation = candidate_import.read_roster(path)
    return set_attendance.assign_attendance_workbook(
        database, set_id, path, validation, imported_by=OPERATOR
    )


def make_scripts(
    database: ProjectDatabase,
    tmp_path: Path,
    scripts: list[tuple[str, str]],
) -> str:
    """Register a batch of ``(roll, set_code)`` scripts and return its batch id.

    Rows are written the way a completed batch leaves them, so that
    :func:`reconciliation_store.reconcile_batch` sees exactly what it would
    see after a real run.
    """
    from omr_scanner.database.models import BatchScan, ScanBatch

    template = build_answer_sheet_template()
    numbers = plan_for(template).numbers
    now = datetime.now(UTC)
    batch_id = batch_store.new_batch_id()
    with database.session() as session:
        session.add(
            ScanBatch(
                batch_id=batch_id, created_at=now, updated_at=now,
                source_folder=str(tmp_path), status="completed",
                total_scans=len(scripts),
            )
        )
        session.flush()
        for index, (roll, set_code) in enumerate(scripts):
            name = f"scan{index:03d}.png"
            result = ScanResult(
                source_path=tmp_path / name,
                outcome=RecognitionOutcome.COMPLETE,
                registration=RegistrationStatus.REGISTERED,
                fields=(
                    FieldView(
                        zone_id="roll_number", label="Roll", field_type="numeric",
                        value=roll, status="complete", needs_review=False, characters=(),
                    ),
                    FieldView(
                        zone_id="set_code", label="Set", field_type="set_code",
                        value=set_code, status="complete", needs_review=False,
                        characters=(),
                    ),
                ),
                answers=tuple(
                    AnswerView(
                        number=number, zone_id="q", value="A", status="resolved",
                        needs_review=False, top_fill=0.9, margin=0.4, confidence=0.9,
                    )
                    for number in numbers
                ),
                identifier_zone_id="roll_number", set_code_zone_id="set_code",
            )
            session.add(
                BatchScan(
                    batch_id=batch_id, batch_index=index,
                    source_path=str(tmp_path / name), filename=name,
                    status="completed", identifier_value=roll,
                    set_code_value=set_code, result_json=json.dumps(result.to_dict()),
                )
            )
    return batch_id


def issues_for(entries, candidate_id: str) -> set[ReconciliationIssue]:
    """Every issue reconciliation recorded against one candidate."""
    for entry in entries:
        if entry.candidate is not None and entry.candidate.candidate_id == candidate_id:
            return set(entry.issues)
    return set()


# ----------------------------------------------------------------------
# Fixtures
# ----------------------------------------------------------------------
@pytest.fixture
def three_sets(workspace: Path):
    """A project defining Sets 10, 11 and 12, with its session open."""
    session = create_project(workspace, "BSCRA Recruitment", exam_name=EXAM_NAME)
    try:
        sets = {
            code: project_sets.add_set(session.database, code, description)
            for code, description in SET_DEFINITIONS
        }
        yield session, sets
    finally:
        if not session.is_closed:
            session.close()


# ----------------------------------------------------------------------
# §24.1 - three Sets, three independent attendance files
# ----------------------------------------------------------------------
class TestThreeSetsAcceptThreeFiles:
    def test_each_set_accepts_its_own_attendance_file(self, three_sets, tmp_path: Path):
        session, sets = three_sets
        for code, exam_set in sets.items():
            path = write_attendance_workbook(
                tmp_path / f"set{code}_attendance.xlsx",
                [f"{code}00{n}" for n in range(1, 4)],
            )
            result = assign(session.database, exam_set.set_id, path)
            assert result.set_code == code
            assert result.candidate_count == 3
            assert result.source_name == f"set{code}_attendance.xlsx"

    def test_each_set_gets_its_own_roster_row(self, three_sets, tmp_path: Path):
        session, sets = three_sets
        roster_ids = set()
        for code, exam_set in sets.items():
            path = write_attendance_workbook(
                tmp_path / f"set{code}.xlsx", [f"{code}001"]
            )
            roster_ids.add(assign(session.database, exam_set.set_id, path).roster_id)
        assert len(roster_ids) == 3

    def test_the_overview_lists_every_set_with_its_own_file(
        self, three_sets, tmp_path: Path
    ):
        session, sets = three_sets
        for code, exam_set in sets.items():
            assign(
                session.database,
                exam_set.set_id,
                write_attendance_workbook(tmp_path / f"s{code}.xlsx", [f"{code}001"]),
            )
        overview = set_attendance.attendance_overview(session.database)
        assert [status.exam_set.code for status in overview] == ["10", "11", "12"]
        assert [status.attendance_file for status in overview] == [
            "s10.xlsx", "s11.xlsx", "s12.xlsx"
        ]

    def test_a_workbook_is_adopted_as_that_sets_result_template(
        self, three_sets, tmp_path: Path
    ):
        session, sets = three_sets
        result = assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "set10.xlsx", ["10001"]),
        )
        assert result.template_adopted is True
        assert result.template_blocker == ""

    def test_a_csv_is_usable_for_attendance_but_not_as_a_template(
        self, three_sets, tmp_path: Path
    ):
        """§16: CSV stays usable for reconciliation, and says why it is not a template."""
        session, sets = three_sets
        result = assign(
            session.database,
            sets["10"].set_id,
            write_attendance_csv(tmp_path / "set10.csv", ["10001", "10002"]),
        )
        assert result.candidate_count == 2
        assert result.template_adopted is False
        assert result.template_blocker == set_attendance.CSV_TEMPLATE_BLOCKER


# ----------------------------------------------------------------------
# §24.2 - the association survives a close and reopen
# ----------------------------------------------------------------------
class TestAssociationSurvivesReopen:
    def test_each_file_is_still_with_its_own_set_after_reopening(
        self, three_sets, tmp_path: Path
    ):
        session, sets = three_sets
        expected = {}
        for code, exam_set in sets.items():
            path = write_attendance_workbook(
                tmp_path / f"set{code}_attendance.xlsx", [f"{code}001", f"{code}002"]
            )
            assign(session.database, exam_set.set_id, path)
            expected[exam_set.set_id] = path.name
        root = session.root
        session.close()

        with open_project(root) as reopened:
            for exam_set in project_sets.list_sets(reopened.database):
                roster = reconciliation_store.active_roster(
                    reopened.database, exam_set.set_id
                )
                assert roster is not None
                assert roster.source_name == expected[exam_set.set_id]

    def test_the_result_template_is_still_with_its_own_set_after_reopening(
        self, three_sets, tmp_path: Path
    ):
        session, sets = three_sets
        for code, exam_set in sets.items():
            assign(
                session.database,
                exam_set.set_id,
                write_attendance_workbook(tmp_path / f"t{code}.xlsx", [f"{code}001"]),
            )
        root = session.root
        session.close()

        with open_project(root) as reopened:
            overview = set_attendance.attendance_overview(reopened.database)
            assert [status.template_file for status in overview] == [
                "t10.xlsx", "t11.xlsx", "t12.xlsx"
            ]

    def test_the_roster_is_bound_by_set_id_not_by_position(
        self, three_sets, tmp_path: Path
    ):
        """Reordering the sets must not move anybody's candidate list."""
        session, sets = three_sets
        for code, exam_set in sets.items():
            assign(
                session.database,
                exam_set.set_id,
                write_attendance_workbook(tmp_path / f"p{code}.xlsx", [f"{code}001"]),
            )
        # Put Set 12 first; the files must follow their sets, not the rows.
        project_sets.move_set(session.database, sets["12"].set_id, -2)
        root = session.root
        session.close()

        with open_project(root) as reopened:
            overview = set_attendance.attendance_overview(reopened.database)
            assert [status.exam_set.code for status in overview] == ["12", "10", "11"]
            assert [status.attendance_file for status in overview] == [
                "p12.xlsx", "p10.xlsx", "p11.xlsx"
            ]


# ----------------------------------------------------------------------
# §24.3 / §24.4 - candidate isolation between Sets
# ----------------------------------------------------------------------
class TestCandidatesAreSetScoped:
    def test_one_sets_candidates_are_not_reachable_through_another(
        self, three_sets, tmp_path: Path
    ):
        session, sets = three_sets
        assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "a.xlsx", ["10001", "10002"]),
        )
        assign(
            session.database,
            sets["11"].set_id,
            write_attendance_workbook(tmp_path / "b.xlsx", ["11001", "11002"]),
        )
        ten = reconciliation_store.active_roster(session.database, sets["10"].set_id)
        eleven = reconciliation_store.active_roster(session.database, sets["11"].set_id)
        assert ten is not None
        assert eleven is not None

        ten_ids = {
            item.candidate_id
            for item in reconciliation_store.roster_candidates(
                session.database, ten.roster_id
            )
        }
        eleven_ids = {
            item.candidate_id
            for item in reconciliation_store.roster_candidates(
                session.database, eleven.roster_id
            )
        }
        assert ten_ids == {"10001", "10002"}
        assert eleven_ids == {"11001", "11002"}
        assert not ten_ids & eleven_ids

    def test_the_same_roll_exists_independently_in_two_sets(
        self, three_sets, tmp_path: Path
    ):
        """§24.4, and the §19 correctness requirement it protects."""
        session, sets = three_sets
        shared = "10001"
        assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "a.xlsx", [shared, "10002"]),
        )
        assign(
            session.database,
            sets["11"].set_id,
            write_attendance_workbook(tmp_path / "b.xlsx", [shared, "20002"]),
        )
        ten = reconciliation_store.active_roster(session.database, sets["10"].set_id)
        eleven = reconciliation_store.active_roster(session.database, sets["11"].set_id)
        assert ten is not None
        assert eleven is not None
        assert ten.roster_id != eleven.roster_id

        ten_rows = reconciliation_store.roster_candidates(session.database, ten.roster_id)
        eleven_rows = reconciliation_store.roster_candidates(
            session.database, eleven.roster_id
        )
        assert shared in {item.candidate_id for item in ten_rows}
        assert shared in {item.candidate_id for item in eleven_rows}
        # Same roll, two rosters, two independent candidate records.
        assert {item.candidate_id for item in ten_rows} != {
            item.candidate_id for item in eleven_rows
        }

    def test_importing_for_one_set_does_not_deactivate_anothers_roster(
        self, three_sets, tmp_path: Path
    ):
        session, sets = three_sets
        assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "a.xlsx", ["10001"]),
        )
        first = reconciliation_store.active_roster(session.database, sets["10"].set_id)
        assign(
            session.database,
            sets["11"].set_id,
            write_attendance_workbook(tmp_path / "b.xlsx", ["11001"]),
        )
        still = reconciliation_store.active_roster(session.database, sets["10"].set_id)
        assert first is not None
        assert still is not None
        assert still.roster_id == first.roster_id

    def test_replacing_one_sets_file_supersedes_only_that_sets_roster(
        self, three_sets, tmp_path: Path
    ):
        session, sets = three_sets
        assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "old.xlsx", ["10001"]),
        )
        assign(
            session.database,
            sets["11"].set_id,
            write_attendance_workbook(tmp_path / "eleven.xlsx", ["11001"]),
        )
        assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "new.xlsx", ["10001", "10002"]),
        )
        ten = reconciliation_store.active_roster(session.database, sets["10"].set_id)
        eleven = reconciliation_store.active_roster(session.database, sets["11"].set_id)
        assert ten is not None
        assert eleven is not None
        assert ten.source_name == "new.xlsx"
        assert eleven.source_name == "eleven.xlsx"


# ----------------------------------------------------------------------
# §24.5 / §24.6 - absence tokens, case-insensitively
# ----------------------------------------------------------------------
class TestAbsenceTokens:
    @pytest.mark.parametrize("token", ["ABS", "abs", "Abs", " ABS "])
    def test_abs_is_recognised_case_insensitively(
        self, three_sets, tmp_path: Path, token: str
    ):
        session, sets = three_sets
        path = write_attendance_workbook(
            tmp_path / f"abs_{token.strip()}_{len(token)}.xlsx",
            ["10001", "10002"],
            absent={"10002": token},
        )
        result = assign(session.database, sets["10"].set_id, path)
        rows = reconciliation_store.roster_candidates(
            session.database, result.roster_id
        )
        states = {item.candidate_id: item.imported_attendance for item in rows}
        assert states["10002"] is AttendanceState.ABSENT
        assert states["10001"] is AttendanceState.PRESENT

    @pytest.mark.parametrize("token", ["ABSENT", "absent", "Absent", " Absent "])
    def test_absent_is_recognised_case_insensitively(
        self, three_sets, tmp_path: Path, token: str
    ):
        session, sets = three_sets
        path = write_attendance_workbook(
            tmp_path / f"absent_{token.strip()}_{len(token)}.xlsx",
            ["10001", "10002"],
            absent={"10002": token},
        )
        result = assign(session.database, sets["10"].set_id, path)
        rows = reconciliation_store.roster_candidates(
            session.database, result.roster_id
        )
        states = {item.candidate_id: item.imported_attendance for item in rows}
        assert states["10002"] is AttendanceState.ABSENT

    def test_a_word_merely_containing_abs_is_not_an_absence(
        self, three_sets, tmp_path: Path
    ):
        """``ABSENTEE`` is not a declaration of absence - see ABSENT_TOKENS."""
        session, sets = three_sets
        path = write_attendance_workbook(
            tmp_path / "absentee.xlsx", ["10001"], absent={"10001": "ABSENTEE"}
        )
        result = assign(session.database, sets["10"].set_id, path)
        rows = reconciliation_store.roster_candidates(
            session.database, result.roster_id
        )
        assert rows[0].imported_attendance is AttendanceState.PRESENT

    def test_the_absence_count_is_recorded_per_set(self, three_sets, tmp_path: Path):
        session, sets = three_sets
        assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(
                tmp_path / "a.xlsx", ["10001", "10002", "10003"],
                absent={"10002": "abs", "10003": "ABSENT"},
            ),
        )
        assign(
            session.database,
            sets["11"].set_id,
            write_attendance_workbook(tmp_path / "b.xlsx", ["11001", "11002"]),
        )
        ten = reconciliation_store.active_roster(session.database, sets["10"].set_id)
        eleven = reconciliation_store.active_roster(session.database, sets["11"].set_id)
        assert ten is not None
        assert eleven is not None
        assert ten.expected_absent == 2
        assert eleven.expected_absent == 0


# ----------------------------------------------------------------------
# §24.7 - §24.10 - the Phase 7 exceptions still surface, now per Set
# ----------------------------------------------------------------------
class TestReconciliationExceptionsSurvive:
    """The classifications come from the real reconciler, not a hand-built entry."""

    @pytest.fixture
    def reconciled(self, three_sets, tmp_path: Path):
        """Set 10, with one of every exception §24.7-§24.10 names."""
        session, sets = three_sets
        path = write_attendance_workbook(
            tmp_path / "set10.xlsx",
            ["10001", "10002", "10003", "10004"],
            absent={"10004": "ABSENT"},
        )
        assignment = assign(session.database, sets["10"].set_id, path)
        batch_id = make_scripts(
            session.database,
            tmp_path,
            [
                ("10001", "10"),   # matched
                ("10002", "10"),   # duplicate: two scripts for one candidate
                ("10002", "10"),
                ("99999", "10"),   # unknown: on no roster
                ("10004", "10"),   # absent-with-script
                # 10003 sat nothing -> present-without-script
            ],
        )
        reconciliation_store.reconcile_batch(
            session.database, assignment.roster_id, batch_id
        )
        entries = reconciliation_store.list_entries(
            session.database, assignment.roster_id, batch_id
        )
        return session, sets, entries

    def test_unknown_candidates_remain_explicit_exceptions(self, reconciled):
        _session, _sets, entries = reconciled
        unknown = [
            entry for entry in entries
            if ReconciliationIssue.UNKNOWN_ID in entry.issues
        ]
        assert len(unknown) == 1
        assert unknown[0].candidate is None

    def test_duplicate_scripts_remain_explicit_exceptions(self, reconciled):
        _session, _sets, entries = reconciled
        assert ReconciliationIssue.DUPLICATE_SCRIPT in issues_for(entries, "10002")

    def test_present_without_script_remains_detectable(self, reconciled):
        _session, _sets, entries = reconciled
        assert ReconciliationIssue.PRESENT_WITHOUT_SCRIPT in issues_for(entries, "10003")

    def test_absent_with_script_remains_detectable(self, reconciled):
        _session, _sets, entries = reconciled
        assert ReconciliationIssue.ABSENT_WITH_SCRIPT in issues_for(entries, "10004")

    def test_a_clean_candidate_carries_no_exception(self, reconciled):
        _session, _sets, entries = reconciled
        assert issues_for(entries, "10001") == set()

    def test_nothing_is_dropped(self, reconciled):
        """Every registered candidate, plus the unplaceable script, is accounted for."""
        _session, _sets, entries = reconciled
        registered = {
            entry.candidate.candidate_id
            for entry in entries
            if entry.candidate is not None
        }
        assert registered == {"10001", "10002", "10003", "10004"}
        assert any(entry.candidate is None for entry in entries)

    def test_one_sets_scripts_do_not_reconcile_against_anothers_roster(
        self, three_sets, tmp_path: Path
    ):
        """The §19 case: the same roll in two Sets must not cross over."""
        session, sets = three_sets
        ten = assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "a.xlsx", ["10001"]),
        )
        eleven = assign(
            session.database,
            sets["11"].set_id,
            write_attendance_workbook(tmp_path / "b.xlsx", ["10001"]),
        )
        # One script, sat by Set 10's candidate.
        batch_id = make_scripts(session.database, tmp_path, [("10001", "10")])

        reconciliation_store.reconcile_batch(session.database, ten.roster_id, batch_id)
        ten_entries = reconciliation_store.list_entries(
            session.database, ten.roster_id, batch_id
        )
        assert issues_for(ten_entries, "10001") == set()

        # Set 11's roster is a different roster: reconciling it against the
        # same batch is a separate question with a separate answer, and the
        # two must not share stored entries.
        eleven_entries = reconciliation_store.list_entries(
            session.database, eleven.roster_id, batch_id
        )
        assert eleven_entries == ()


# ----------------------------------------------------------------------
# §24.11 / §24.12 - a missing workbook, and never a substitution
# ----------------------------------------------------------------------
class TestMissingAttendanceIsRefused:
    def test_a_set_with_no_attendance_has_no_roster(self, three_sets, tmp_path: Path):
        session, sets = three_sets
        assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "a.xlsx", ["10001"]),
        )
        assert (
            reconciliation_store.active_roster(session.database, sets["11"].set_id)
            is None
        )

    def test_resolving_its_sources_refuses_and_names_the_set(
        self, three_sets, tmp_path: Path
    ):
        session, sets = three_sets
        assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "a.xlsx", ["10001"]),
        )
        with pytest.raises(report_store.SetGenerationRefusedError) as excinfo:
            report_store.resolve_set_sources(session.database, sets["11"].set_id)
        assert "Set 11" in excinfo.value.user_message
        assert "attendance" in excinfo.value.user_message.lower()

    def test_generating_it_is_blocked_rather_than_borrowing(
        self, three_sets, tmp_path: Path
    ):
        """§15 in full: blocked, with a message, and no file written."""
        session, sets = three_sets
        assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "a.xlsx", ["10001"]),
        )
        batch_id = make_scripts(session.database, tmp_path, [("10001", "10")])
        exports = tmp_path / "exports"

        outcome = report_store.generate_for_set(
            session.database, sets["11"].set_id, batch_id,
            build_answer_sheet_template(),
            project_name="BSCRA Recruitment", output_dir=exports,
        )
        assert outcome.status == "blocked"
        assert outcome.output_path is None
        assert any("Set 11" in warning for warning in outcome.warnings)
        assert not exports.exists() or not list(exports.glob("*.xlsx"))

    def test_the_status_says_no_file_is_assigned(self, three_sets, tmp_path: Path):
        session, sets = three_sets
        status = set_attendance.set_attendance_status(
            session.database, sets["11"]
        )
        assert status.has_attendance is False
        assert status.can_generate is False
        assert status.describe() == "No attendance file assigned"


class TestNoSubstitutionBetweenSets:
    """§24.12: one Set's workbook can never stand in for another's."""

    def test_each_sets_template_is_its_own_file(self, three_sets, tmp_path: Path):
        session, sets = three_sets
        for code, exam_set in sets.items():
            assign(
                session.database,
                exam_set.set_id,
                write_attendance_workbook(tmp_path / f"{code}.xlsx", [f"{code}001"]),
            )
        for code, exam_set in sets.items():
            sources = report_store.resolve_set_sources(
                session.database, exam_set.set_id
            )
            assert sources.set_code == code
            assert sources.association.template_path.endswith(f"{code}.xlsx")

    def test_a_template_belonging_to_another_set_is_not_returned(
        self, three_sets, tmp_path: Path
    ):
        """Resolution is by set_id first, so a shared code cannot reach across."""
        session, sets = three_sets
        assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "ten.xlsx", ["10001"]),
        )
        found = report_store.get_template_association_for_set(
            session.database, sets["11"].set_id, "11"
        )
        assert found is None

    def test_resolving_by_the_wrong_set_id_never_yields_anothers_roster(
        self, three_sets, tmp_path: Path
    ):
        session, sets = three_sets
        ten = assign(
            session.database,
            sets["10"].set_id,
            write_attendance_workbook(tmp_path / "ten.xlsx", ["10001"]),
        )
        for other in ("11", "12"):
            roster = reconciliation_store.active_roster(
                session.database, sets[other].set_id
            )
            assert roster is None
        mine = reconciliation_store.active_roster(session.database, sets["10"].set_id)
        assert mine is not None
        assert mine.roster_id == ten.roster_id


# ----------------------------------------------------------------------
# §29 - a project that predates per-Set attendance
# ----------------------------------------------------------------------
class TestLegacyUnassignedRoster:
    def test_a_roster_imported_without_a_set_is_reported_unassigned(
        self, three_sets, tmp_path: Path
    ):
        """Migration deliberately guesses nothing; the roster waits to be claimed."""
        session, _sets = three_sets
        validation = candidate_import.read_roster(
            write_attendance_workbook(tmp_path / "legacy.xlsx", ["10001"])
        )
        reconciliation_store.import_roster(
            session.database, validation, imported_by=OPERATOR
        )
        unassigned = reconciliation_store.unassigned_rosters(session.database)
        assert [item.source_name for item in unassigned] == ["legacy.xlsx"]

    def test_it_is_not_visible_as_any_sets_attendance(self, three_sets, tmp_path: Path):
        session, sets = three_sets
        validation = candidate_import.read_roster(
            write_attendance_workbook(tmp_path / "legacy.xlsx", ["10001"])
        )
        reconciliation_store.import_roster(
            session.database, validation, imported_by=OPERATOR
        )
        for exam_set in sets.values():
            assert (
                reconciliation_store.active_roster(session.database, exam_set.set_id)
                is None
            )

    def test_an_operator_can_assign_it_to_a_set_explicitly(
        self, three_sets, tmp_path: Path
    ):
        session, sets = three_sets
        validation = candidate_import.read_roster(
            write_attendance_workbook(tmp_path / "legacy.xlsx", ["10001"])
        )
        roster_id = reconciliation_store.import_roster(
            session.database, validation, imported_by=OPERATOR
        )
        reconciliation_store.assign_roster_to_set(
            session.database, roster_id, sets["10"].set_id
        )
        claimed = reconciliation_store.active_roster(
            session.database, sets["10"].set_id
        )
        assert claimed is not None
        assert claimed.roster_id == roster_id
        assert reconciliation_store.unassigned_rosters(session.database) == ()

    def test_it_cannot_be_reassigned_to_a_second_set(self, three_sets, tmp_path: Path):
        session, sets = three_sets
        validation = candidate_import.read_roster(
            write_attendance_workbook(tmp_path / "legacy.xlsx", ["10001"])
        )
        roster_id = reconciliation_store.import_roster(
            session.database, validation, imported_by=OPERATOR
        )
        reconciliation_store.assign_roster_to_set(
            session.database, roster_id, sets["10"].set_id
        )
        with pytest.raises(reconciliation_store.ReconciliationError):
            reconciliation_store.assign_roster_to_set(
                session.database, roster_id, sets["11"].set_id
            )
