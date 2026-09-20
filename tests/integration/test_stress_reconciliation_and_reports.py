"""Reconciliation and Phase 9 report generation against a real stress batch.

Phase 10 validation brief §8/§40: process a stress dataset through the real
production pipeline, reconcile it against a real (synthetic) roster through
the real :mod:`omr_scanner.services.reconciliation_store`, and generate a
real Phase 9 Excel report from the results - at a scale large enough to
exercise the architecture, not merely its correctness on a handful of rows.

Runs at 600 sheets by default (fast enough for the ordinary suite while
still drawing every case kind - see
``tests/unit/test_stress_roster.py::TestReconciliationRelevantCaseKindsExist``
for the same guarantee at a different scale) and again at 10,000 sheets
under the ``stress`` marker, mirroring
``tests/integration/test_stress_kill_resume.py``'s own 100/1,000/10,000
scale ladder.
"""

from __future__ import annotations

from fractions import Fraction
from typing import TYPE_CHECKING

import openpyxl
import pytest
from tests.conftest import build_answer_sheet_template

from omr_scanner.database.engine import open_project_database
from omr_scanner.domain.reconciliation import ReconciliationStatus
from omr_scanner.evaluation import stress_dataset, stress_roster, stress_runner
from omr_scanner.evaluation.test_cases import FieldLayout
from omr_scanner.reporting import excel as rx
from omr_scanner.services import batch_store, reconciliation_store
from omr_scanner.services.candidate_import import ColumnMapping, RosterValidation
from omr_scanner.services.report_template import ReportColumnMapping, TemplateRoster, TemplateRow

if TYPE_CHECKING:
    from pathlib import Path

    from omr_scanner.database.engine import ProjectDatabase
    from omr_scanner.domain.template import OmrTemplate

@pytest.fixture(scope="module")
def template() -> OmrTemplate:
    return build_answer_sheet_template(roll_digits=7)


@pytest.fixture(scope="module")
def layout(template: OmrTemplate) -> FieldLayout:
    return FieldLayout.of(template)


@pytest.fixture
def database(tmp_path: Path) -> ProjectDatabase:
    handle = open_project_database(tmp_path / "database.sqlite", create=True)
    yield handle
    handle.close()


def _run_stress_batch(
    database: ProjectDatabase,
    template: OmrTemplate,
    spec: stress_dataset.StressDatasetSpec,
    tmp_path: Path,
    *,
    workers: int,
) -> str:
    tmp_path.mkdir(parents=True, exist_ok=True)
    batch_id = stress_runner.create_stress_batch(database, spec, template)
    recorder = batch_store.BatchRecorder(database=database, batch_id=batch_id)
    stress_runner.run_stress_batch(
        database,
        batch_id,
        template,
        spec,
        workers=workers,
        scratch_root=tmp_path,
        on_result=recorder.record,
    )
    recorder.flush()
    return batch_id


def _import_stress_roster(
    database: ProjectDatabase,
    spec: stress_dataset.StressDatasetSpec,
    layout: FieldLayout,
) -> int:
    candidates = stress_roster.generate_stress_roster(spec, layout)
    validation = RosterValidation(
        candidates=candidates,
        issues=(),
        rows_read=len(candidates),
        blank_ids=0,
        duplicate_ids=0,
        expected_present=sum(1 for c in candidates if c.imported_attendance.value == "present"),
        expected_absent=sum(1 for c in candidates if c.imported_attendance.value == "absent"),
        attendance_unknown=0,
        mapping=ColumnMapping(candidate_id=0, name=1, attendance=2),
        source_name="stress-roster.csv",
    )
    return reconciliation_store.import_roster(database, validation, imported_by="stress-test")


def _duplicate_or_unknown_indices(
    spec: stress_dataset.StressDatasetSpec,
) -> tuple[int | None, int | None]:
    """Find one duplicate-generating index and one unknown-candidate index."""
    duplicate_kinds = {
        stress_dataset.StressCaseKind.DUPLICATE_ID,
        stress_dataset.StressCaseKind.DUPLICATE_ROLL_DIFFERENT_IMAGE,
        stress_dataset.StressCaseKind.EXACT_DUPLICATE_SCAN,
    }
    duplicate_index = None
    unknown_index = None
    for index in range(spec.sheet_count):
        kind = spec.kind_for_index(index)
        if duplicate_index is None and kind in duplicate_kinds:
            duplicate_index = index
        if unknown_index is None and kind is stress_dataset.StressCaseKind.UNKNOWN_CANDIDATE:
            unknown_index = index
        if duplicate_index is not None and unknown_index is not None:
            break
    return duplicate_index, unknown_index


def _assert_reconciliation_is_correct(
    database: ProjectDatabase,
    roster_id: int,
    batch_id: str,
    spec: stress_dataset.StressDatasetSpec,
    layout: FieldLayout,
) -> None:
    counts = reconciliation_store.reconcile_batch(database, roster_id, batch_id)

    # "Nothing disappears": every roster candidate and every unmatched script
    # is accounted for in exactly one entry - see reconciliation.py's own
    # stated invariant.
    entries = reconciliation_store.list_entries(database, roster_id, batch_id)
    total_scripts_in_entries = sum(len(entry.scripts) for entry in entries)
    scan_count = batch_store.load_summary(database, batch_id).total
    assert total_scripts_in_entries == scan_count

    # The deliberate ABSENT_WITH_SCRIPT candidate.
    forced_absent_id = stress_roster.absent_with_script_candidate_id(spec, layout)
    assert forced_absent_id is not None
    forced_entry = next(e for e in entries if e.candidate_id == forced_absent_id)
    assert forced_entry.status is ReconciliationStatus.ABSENT_WITH_SCRIPT
    assert counts.absent_with_script >= 1

    # A duplicate-generating sheet produces a genuine absentee (its own
    # natural-roll candidate has no script) and a genuine duplicate (the
    # partner candidate has two).
    duplicate_index, unknown_index = _duplicate_or_unknown_indices(spec)
    assert duplicate_index is not None, "seed/sheet_count drew no duplicate-generating kind"
    own_roll = stress_dataset.natural_roll(spec, layout, duplicate_index)
    if own_roll != forced_absent_id:
        own_entry = next((e for e in entries if e.candidate_id == own_roll), None)
        assert own_entry is not None
        assert own_entry.status is ReconciliationStatus.PRESENT_WITHOUT_SCRIPT

    assert counts.duplicate_script >= 1

    assert unknown_index is not None, "seed/sheet_count drew no UNKNOWN_CANDIDATE kind"
    assert counts.unknown_id >= 1


def _build_aligned_report_template(
    output_path: Path,
    spec: stress_dataset.StressDatasetSpec,
    layout: FieldLayout,
) -> TemplateRoster:
    """A minimal, real .xlsx roster whose Roll No. column matches the stress roster.

    Built directly with openpyxl (mirroring ``tests/report_fixtures.py``'s own
    approach) rather than through ``report_template.read_template`` - what is
    under test here is Phase 9's *writer* (``reporting.excel``), not its
    reader, and building the file this way is what lets every row's roll be
    identical to a candidate this test's own reconciliation already reasoned
    about.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Rollwise"
    headers = ("Sl.No.", "Roll No.", "Name", "Total (90)", "Merit")
    for column, text in enumerate(headers, start=1):
        sheet.cell(row=1, column=column, value=text)

    rows: list[TemplateRow] = []
    for index in range(spec.sheet_count):
        roll = stress_dataset.natural_roll(spec, layout, index)
        sheet_row_number = index + 2
        sheet.cell(row=sheet_row_number, column=1, value=index + 1)
        sheet.cell(row=sheet_row_number, column=2, value=roll)
        sheet.cell(row=sheet_row_number, column=3, value=f"Stress Candidate {index:07d}")
        rows.append(TemplateRow(sheet_row_number=sheet_row_number, roll=roll))
    workbook.save(output_path)
    workbook.close()

    return TemplateRoster(
        path=output_path,
        sheet="Rollwise",
        mapping=ReportColumnMapping(roll=1, marks=3, serial=0, name=2, rank=4),
        rows=tuple(rows),
        header_row_number=1,
    )


def _decisions_from_real_results(
    database: ProjectDatabase,
    roster_id: int,
    batch_id: str,
) -> dict[str, rx.CandidateReportRow]:
    """One decision per roster candidate, derived from real recognition output."""
    entries = reconciliation_store.list_entries(database, roster_id, batch_id)
    results_by_scan = batch_store.results_by_scan(database, batch_id)

    decisions: dict[str, rx.CandidateReportRow] = {}
    for entry in entries:
        if not entry.is_registered:
            continue
        present_scripts = [view for view in entry.scripts if view.counts_as_a_script]
        if not present_scripts:
            decisions[entry.candidate_id] = rx.CandidateReportRow(
                roll=entry.candidate_id, is_absent=True
            )
            continue
        result = results_by_scan.get(present_scripts[0].script.scan_id)
        answered = sum(1 for answer in result.answers if answer.value) if result else 0
        decisions[entry.candidate_id] = rx.CandidateReportRow(
            roll=entry.candidate_id, is_absent=False, final_score=Fraction(answered)
        )
    return decisions


def _assert_report_is_correct(
    database: ProjectDatabase,
    roster_id: int,
    batch_id: str,
    spec: stress_dataset.StressDatasetSpec,
    layout: FieldLayout,
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "report_template.xlsx"
    roster = _build_aligned_report_template(template_path, spec, layout)
    assert len(roster.rows) == spec.sheet_count

    output_path = tmp_path / "report_output.xlsx"
    rx.copy_into(template_path, output_path)

    decisions = _decisions_from_real_results(database, roster_id, batch_id)
    result = rx.populate_rollwise(output_path, roster, decisions)

    assert result.rows_written == spec.sheet_count
    # Every roster row had a decision computed above; nothing should be left
    # unresolved purely because this test failed to supply one.
    assert result.unresolved_rows == ()
    assert result.absent_written == sum(1 for d in decisions.values() if d.is_absent)
    assert result.marks_written == sum(1 for d in decisions.values() if not d.is_absent)

    # Read the written workbook back and spot-check specific real rows,
    # rather than trusting the writer's own return value alone.
    check_workbook = openpyxl.load_workbook(output_path, data_only=False)
    try:
        sheet = check_workbook["Rollwise"]

        # A genuine absentee: a candidate no script matched at all (a
        # duplicate-generating sheet's own natural-roll candidate - see
        # `_assert_reconciliation_is_correct`). Note this is deliberately
        # *not* the ABSENT_WITH_SCRIPT candidate: that one still has a real
        # script, and a report generator correctly reports what was
        # actually recognised for them rather than what the roster's
        # attendance column claimed - reconciliation is where that
        # disagreement is surfaced as an exception, not the report.
        duplicate_index, _unknown_index = _duplicate_or_unknown_indices(spec)
        assert duplicate_index is not None
        genuine_absentee_id = stress_dataset.natural_roll(spec, layout, duplicate_index)
        forced_absent_id = stress_roster.absent_with_script_candidate_id(spec, layout)
        assert genuine_absentee_id != forced_absent_id, (
            "this seed/sheet_count made the duplicate-generating index's own "
            "natural roll collide with the deliberate ABSENT_WITH_SCRIPT "
            "candidate - pick a different seed"
        )
        assert decisions[genuine_absentee_id].is_absent is True
        absent_row = next(r for r in roster.rows if r.roll == genuine_absentee_id)
        assert sheet.cell(row=absent_row.sheet_row_number, column=4).value == "ABSENT"

        present_row = next(
            r for r in roster.rows if decisions[r.roll].is_absent is False
        )
        written_marks = sheet.cell(row=present_row.sheet_row_number, column=4).value
        assert written_marks == float(decisions[present_row.roll].final_score)
    finally:
        check_workbook.close()


class TestStressReconciliationAndReportGeneration:
    def test_at_600_sheets(self, database, template, layout, tmp_path: Path) -> None:
        spec = stress_dataset.StressDatasetSpec(seed=20260920, sheet_count=600)
        batch_id = _run_stress_batch(database, template, spec, tmp_path / "scratch", workers=2)
        roster_id = _import_stress_roster(database, spec, layout)
        _assert_reconciliation_is_correct(database, roster_id, batch_id, spec, layout)
        _assert_report_is_correct(database, roster_id, batch_id, spec, layout, tmp_path)

    @pytest.mark.stress
    def test_at_10000_sheets(self, database, template, layout, tmp_path: Path) -> None:
        spec = stress_dataset.StressDatasetSpec(seed=20260920, sheet_count=10_000)
        batch_id = _run_stress_batch(database, template, spec, tmp_path / "scratch", workers=4)
        roster_id = _import_stress_roster(database, spec, layout)
        _assert_reconciliation_is_correct(database, roster_id, batch_id, spec, layout)
        _assert_report_is_correct(database, roster_id, batch_id, spec, layout, tmp_path)
