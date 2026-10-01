r"""Build the committed schema-13 project fixtures - with the Phase-1 (schema-13) build.

Purpose:
    Migration 14 and the scan-session backfill must be tested from projects
    the previous build itself wrote, not from hand-made approximations. This
    script drives the schema-13 build's own services (``main`` at ``ce3f082``,
    0.1.1-alpha.0 phase 1) to write four projects, committed beside it:

    * ``one_batch/`` - one batch of Set A, attended, reconciled, keyed,
      scored and reported (an XLSX final export is generated).
    * ``two_unrelated/`` - two batches of one examination, no relationship.
    * ``linked_rescan/`` - a sheet of batch 1 rejected and replaced, confirmed,
      from batch 2 (created later): an unambiguous cross-batch rescan.
    * ``ambiguous_chain/`` - three batches chained by confirmed rescans
      (1 -> 2 and 2 -> 3): ambiguous, must not be grouped.

How to regenerate (only ever with the schema-13 build)::

    git worktree add ..\OMRflow-schema13 ce3f082
    $env:PYTHONPATH = "..\OMRflow-schema13\src"
    .venv\Scripts\python.exe tests\fixtures\schema13\make_schema13_projects.py `
        --old-checkout ..\OMRflow-schema13

The script refuses to run against a build whose schema is not 13.
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import UTC, datetime, timedelta
from pathlib import Path

HERE = Path(__file__).resolve().parent
OPERATOR = "Fixture Builder"
START = datetime(2026, 9, 30, 9, 0, tzinfo=UTC)


def _use_old_checkout(root: Path) -> None:
    sys.path.insert(0, str(root / "src"))
    sys.path.insert(1, str(root))


def _attendance(path: Path, rolls: list[str]) -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for column, text in enumerate(("Sl.No.", "Roll No.", "Name", "Total"), start=1):
        sheet.cell(row=1, column=column, value=text)
    for offset, roll in enumerate(rolls):
        sheet.cell(row=2 + offset, column=1, value=offset + 1)
        sheet.cell(row=2 + offset, column=2, value=roll)
        sheet.cell(row=2 + offset, column=3, value=f"CANDIDATE {roll}")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def _register(
    database, folder: Path, scripts, numbers, *, created: datetime
) -> tuple[str, list[int]]:
    """A completed batch of ``(roll, set read, correct answers)`` scripts."""
    from sqlalchemy import select

    from omr_scanner.database.models import BatchScan, ScanBatch
    from omr_scanner.services import batch_store
    from omr_scanner.services.recognition_models import (
        AnswerView,
        FieldView,
        RecognitionOutcome,
        RegistrationStatus,
        ScanResult,
    )

    batch_id = batch_store.new_batch_id()
    ordered = list(numbers)
    folder.mkdir(parents=True, exist_ok=True)
    with database.session() as session:
        session.add(
            ScanBatch(
                batch_id=batch_id, created_at=created, updated_at=created,
                source_folder=str(folder), status="completed", total_scans=len(scripts),
            )
        )
        session.flush()
        for index, (roll, set_read, correct) in enumerate(scripts):
            name = f"{batch_id[:6]}_{index:03d}.png"
            result = ScanResult(
                source_path=folder / name,
                outcome=RecognitionOutcome.COMPLETE,
                registration=RegistrationStatus.REGISTERED,
                fields=(
                    FieldView(zone_id="roll_number", label="Roll", field_type="numeric",
                              value=roll, status="complete", needs_review=False, characters=()),
                    FieldView(zone_id="set_code", label="Set", field_type="set_code",
                              value=set_read, status="complete", needs_review=False,
                              characters=()),
                ),
                answers=tuple(
                    AnswerView(number=number, zone_id="q",
                               value="A" if position < correct else "B", status="resolved",
                               needs_review=False, top_fill=0.9, margin=0.4, confidence=0.9)
                    for position, number in enumerate(ordered)
                ),
                identifier_zone_id="roll_number", set_code_zone_id="set_code",
            )
            session.add(
                BatchScan(
                    batch_id=batch_id, batch_index=index, source_path=str(folder / name),
                    filename=name, status="completed", identifier_value=roll,
                    set_code_value=set_read, result_json=json.dumps(result.to_dict()),
                    content_sha256=f"{batch_id}{index:04d}".ljust(64, "0")[:64],
                )
            )
        session.flush()
        ids = [
            int(item)
            for item in session.scalars(
                select(BatchScan.scan_id)
                .where(BatchScan.batch_id == batch_id)
                .order_by(BatchScan.batch_index)
            ).all()
        ]
    return batch_id, ids


def _replace(database, original: int, replacement: int) -> None:
    from omr_scanner.domain.scan_lifecycle import RejectionReason
    from omr_scanner.services import scan_lifecycle

    scan_lifecycle.reject_scan(
        database, original, reviewer=OPERATOR, reason=RejectionReason.FOLDED
    )
    scan_lifecycle.confirm_replacement(database, original, replacement, reviewer=OPERATOR)


def _project(workspace: Path, name: str) -> tuple[object, object]:
    from omr_scanner.services import create_project, project_sets

    session = create_project(workspace, name, exam_name=f"Schema 13 fixture: {name}")
    exam_set = project_sets.add_set(session.database, "A", "Set A")
    return session, exam_set


def _score(session, exam_set, rolls, batch_id, template, *, report: bool = False) -> None:
    from omr_scanner.services import (
        candidate_import,
        reconciliation_store,
        report_store,
        scoring_store,
        set_attendance,
    )
    from omr_scanner.services.answer_key import plan_for, read_key

    database = session.database
    plan = plan_for(template)
    attendance = _attendance(session.root / "inputs" / "attendance.xlsx", rolls)
    assignment = set_attendance.assign_attendance_workbook(
        database, exam_set.set_id, attendance, candidate_import.read_roster(attendance),
        imported_by=OPERATOR,
    )
    reconciliation_store.reconcile_batch(database, assignment.roster_id, batch_id)
    stored = scoring_store.save_key(
        database, read_key("A" * plan.question_count, plan, "A").to_key()
    )
    scoring_store.verify_key(database, stored.key_id, verified_by=OPERATOR)
    scoring_store.score_batch(database, assignment.roster_id, batch_id, template,
                              computed_by=OPERATOR)
    if report:
        outcome = report_store.generate_for_set(
            database, exam_set.set_id, batch_id, template, project_name=session.name,
            output_dir=session.root / "exports", computed_by=OPERATOR, final=True,
        )
        if not outcome.ok:
            raise SystemExit(f"report generation failed: {outcome.status} {outcome.warnings}")


def main() -> int:
    """Write the four fixtures beside this script."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--old-checkout", type=Path, required=True)
    arguments = parser.parse_args()
    _use_old_checkout(arguments.old_checkout.resolve())

    from tests.conftest import build_answer_sheet_template

    from omr_scanner import __version__
    from omr_scanner.database.migrations import SCHEMA_VERSION
    from omr_scanner.services.answer_key import plan_for

    if SCHEMA_VERSION != 13:
        print(f"Refusing: this build writes schema {SCHEMA_VERSION}, not 13", file=sys.stderr)
        return 2

    template = build_answer_sheet_template()
    numbers = plan_for(template).numbers
    workspace = HERE / "_build"
    shutil.rmtree(workspace, ignore_errors=True)
    workspace.mkdir()
    built: dict[str, Path] = {}

    session, exam_set = _project(workspace, "one_batch")
    rolls = ["30001", "30002", "30003"]
    batch, _ = _register(session.database, session.root / "scans", [
        (rolls[0], "A", 5), (rolls[1], "A", 3), (rolls[2], "A", 0),
    ], numbers, created=START)
    _score(session, exam_set, rolls, batch, template, report=True)
    built["one_batch"] = session.root
    session.close()

    session, exam_set = _project(workspace, "two_unrelated")
    _register(session.database, session.root / "s1", [("31001", "A", 4)], numbers, created=START)
    _register(session.database, session.root / "s2", [("31002", "A", 2)], numbers,
              created=START + timedelta(hours=2))
    built["two_unrelated"] = session.root
    session.close()

    session, exam_set = _project(workspace, "linked_rescan")
    first, first_ids = _register(session.database, session.root / "s1", [
        ("32001", "A", 5), ("32002", "A", 1),
    ], numbers, created=START)
    _second, second_ids = _register(session.database, session.root / "s2", [("32002", "A", 4)],
                                    numbers, created=START + timedelta(hours=1))
    _replace(session.database, first_ids[1], second_ids[0])
    _score(session, exam_set, ["32001", "32002"], first, template)
    built["linked_rescan"] = session.root
    session.close()

    session, exam_set = _project(workspace, "ambiguous_chain")
    _b1, ids1 = _register(session.database, session.root / "s1",
                          [("33001", "A", 5), ("33002", "A", 1)], numbers, created=START)
    _b2, ids2 = _register(session.database, session.root / "s2",
                          [("33002", "A", 4), ("33003", "A", 2)], numbers,
                          created=START + timedelta(hours=1))
    _b3, ids3 = _register(session.database, session.root / "s3", [("33003", "A", 3)], numbers,
                          created=START + timedelta(hours=2))
    _replace(session.database, ids1[1], ids2[0])
    _replace(session.database, ids2[1], ids3[0])
    built["ambiguous_chain"] = session.root
    session.close()

    for name, root in built.items():
        target = HERE / name
        shutil.rmtree(target, ignore_errors=True)
        shutil.copytree(root, target, ignore=shutil.ignore_patterns("backups", "logs"))
    shutil.rmtree(workspace, ignore_errors=True)
    (HERE / "PROVENANCE.json").write_text(
        json.dumps(
            {
                "written_by_version": __version__,
                "schema_version": SCHEMA_VERSION,
                "written_at": datetime.now(UTC).isoformat(timespec="seconds"),
                "projects": sorted(built),
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {', '.join(sorted(built))} with OMRFlow {__version__} (schema 13)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
