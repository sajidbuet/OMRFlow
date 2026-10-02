r"""Build the committed schema-15 project fixtures - with the schema-15 build.

Purpose:
    Migration 16 (intake sources and ledger) must be tested from projects the
    previous build itself wrote. This script drives the schema-15 build's own
    services (``main`` at ``fd063f8``, 0.1.1-alpha.0 phase 4) to write two
    projects, committed beside it:

    * ``finite/`` - the traditional finite workflow: one scan session, one
      batch of Set A, attended, reconciled, keyed, scored and reported, then
      the session closed and a final export generated.
    * ``duplicates/`` - one open scan session of two batches: the second batch
      repeats one file of the first byte for byte under another name, which
      the schema-15 build links as ``duplicate_content`` before recognition;
      and a second, closed session holding the same bytes again (another
      session's copy is never a duplicate).

How to regenerate (only ever with the schema-15 build)::

    .venv\Scripts\python.exe tests\fixtures\schema15\make_schema15_projects.py `
        --old-checkout <checkout of fd063f8>

The script refuses to run against a build whose schema is not 15.
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import UTC, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
OPERATOR = "Fixture Builder"


def _use_old_checkout(root: Path) -> None:
    sys.path.insert(0, str(root / "src"))
    sys.path.insert(1, str(root))


def _attendance(path: Path, rolls: list[str]) -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for column, text in enumerate(("Sl.No.", "Roll No.", "Name", "Total"), start=1):
        sheet.cell(row=1, column=column, value=text)
    for offset, roll in enumerate(rolls):
        sheet.cell(row=2 + offset, column=1, value=offset + 1)
        sheet.cell(row=2 + offset, column=2, value=roll)
        sheet.cell(row=2 + offset, column=3, value=f"CANDIDATE {roll}")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def _read_batch(database, files: list[tuple[Path, str, int]], template, *, session_id=None) -> str:
    """A Process All: hash, link exact duplicates (the worker's steps), then read."""
    from sqlalchemy import update
    from tests.integration.test_reject_and_rescan import make_result

    from omr_scanner.database.models import BatchScan
    from omr_scanner.services import batch_store, scan_lifecycle, scan_provenance, scan_sessions

    paths = [path for path, _roll, _correct in files]
    batch_id = scan_sessions.start_batch(
        database, paths, identity=batch_store.BatchIdentity.of(template), started_by=OPERATOR,
        scan_session_id=session_id,
    )
    scan_provenance.compute_hashes_for_batch(database, batch_id)
    linked = {str(item.path) for item in scan_lifecycle.link_exact_duplicates(database, batch_id)}
    ids = batch_store.scan_ids_by_path(database, batch_id)
    with database.session() as session:
        for path, roll, correct in files:
            if str(path) in linked:
                continue
            session.execute(
                update(BatchScan)
                .where(BatchScan.scan_id == ids[path])
                .values(
                    status="completed", identifier_value=roll, set_code_value="A",
                    result_json=json.dumps(make_result(path, roll, "A", correct).to_dict()),
                )
            )
    batch_store.finalise_batch(database, batch_id)
    return batch_id


def _write(folder: Path, name: str, content: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / name
    path.write_bytes(content.encode())
    return path


def _set_up(workspace: Path, name: str, rolls: list[str], template):
    from omr_scanner.services import (
        candidate_import,
        create_project,
        project_sets,
        scoring_store,
        set_attendance,
    )
    from omr_scanner.services.answer_key import plan_for, read_key

    session = create_project(workspace, name, exam_name=f"Schema 15 fixture: {name}")
    exam_set = project_sets.add_set(session.database, "A", "Set A")
    attendance = _attendance(session.root / "inputs" / "attendance.xlsx", rolls)
    assignment = set_attendance.assign_attendance_workbook(
        session.database, exam_set.set_id, attendance, candidate_import.read_roster(attendance),
        imported_by=OPERATOR,
    )
    plan = plan_for(template)
    stored = scoring_store.save_key(
        session.database, read_key("A" * plan.question_count, plan, "A").to_key()
    )
    scoring_store.verify_key(session.database, stored.key_id, verified_by=OPERATOR)
    return session, exam_set, assignment.roster_id


def main() -> int:  # noqa: PLR0915 - one linear fixture recipe
    """Write the two fixtures beside this script."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--old-checkout", type=Path, required=True)
    arguments = parser.parse_args()
    _use_old_checkout(arguments.old_checkout.resolve())

    from tests.conftest import build_answer_sheet_template

    from omr_scanner import __version__
    from omr_scanner.database.migrations import SCHEMA_VERSION
    from omr_scanner.services import (
        reconciliation_store,
        report_store,
        scan_sessions,
        scoring_store,
    )

    if SCHEMA_VERSION != 15:
        print(f"Refusing: this build writes schema {SCHEMA_VERSION}, not 15", file=sys.stderr)
        return 2

    template = build_answer_sheet_template()
    workspace = HERE / "_build"
    shutil.rmtree(workspace, ignore_errors=True)
    workspace.mkdir()
    built: dict[str, Path] = {}
    facts: dict[str, object] = {}

    # --- finite: the traditional one-batch workflow, closed and exported ----
    rolls = ["50001", "50002", "50003"]
    session, exam_set, roster = _set_up(workspace, "finite", rolls, template)
    scans = session.root / "scans"
    batch = _read_batch(session.database, [
        (_write(scans, "f1.png", "schema15:f1"), rolls[0], 5),
        (_write(scans, "f2.png", "schema15:f2"), rolls[1], 3),
        (_write(scans, "f3.png", "schema15:f3"), rolls[2], 0),
    ], template)
    reconciliation_store.reconcile_batch(session.database, roster, batch)
    scoring_store.score_batch(session.database, roster, batch, template, computed_by=OPERATOR)
    owner = scan_sessions.session_of_batch(session.database, batch)
    assert owner is not None
    scan_sessions.close_scan_session(session.database, owner, closed_by=OPERATOR)
    outcome = report_store.generate_for_set(
        session.database, exam_set.set_id, batch, template, project_name=session.name,
        output_dir=session.root / "exports", computed_by=OPERATOR, final=True,
    )
    if not outcome.ok:
        raise SystemExit(f"report generation failed: {outcome.status} {outcome.warnings}")
    facts["finite_batch"] = batch
    facts["finite_session"] = owner
    built["finite"] = session.root
    session.close()

    # --- duplicates: same bytes again in the session, and in another session --
    rolls = ["51001", "51002", "51003"]
    session, _exam_set, roster = _set_up(workspace, "duplicates", rolls, template)
    first = _read_batch(session.database, [
        (_write(session.root / "s1", "d1.png", "schema15:d1"), rolls[0], 5),
        (_write(session.root / "s1", "d2.png", "schema15:d2"), rolls[1], 2),
    ], template)
    second = _read_batch(session.database, [
        (_write(session.root / "s2", "d2-again.png", "schema15:d2"), rolls[1], 2),
        (_write(session.root / "s2", "d3.png", "schema15:d3"), rolls[2], 4),
    ], template)
    open_session = scan_sessions.session_of_batch(session.database, first)
    other = scan_sessions.create_scan_session(
        session.database, name="Other sitting", created_by=OPERATOR, activate=False
    )
    third = _read_batch(session.database, [
        (_write(session.root / "s3", "d1-other.png", "schema15:d1"), rolls[0], 5),
    ], template, session_id=other.scan_session_id)
    scan_sessions.close_scan_session(session.database, other.scan_session_id, closed_by=OPERATOR)
    facts.update(
        duplicates_first_batch=first,
        duplicates_second_batch=second,
        duplicates_open_session=open_session,
        duplicates_other_batch=third,
        duplicates_closed_session=other.scan_session_id,
    )
    built["duplicates"] = session.root
    session.close()

    for name, root in built.items():
        target = HERE / name
        shutil.rmtree(target, ignore_errors=True)
        shutil.copytree(root, target, ignore=shutil.ignore_patterns("backups", "logs"))
    shutil.rmtree(workspace, ignore_errors=True)
    (HERE / "PROVENANCE.json").write_text(
        json.dumps(
            {
                "written_by_version": __version__,
                "schema_version": SCHEMA_VERSION,
                "written_at": datetime.now(UTC).isoformat(timespec="seconds"),
                "projects": sorted(built),
                **facts,
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {', '.join(sorted(built))} with OMRFlow {__version__} (schema 15)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
