r"""Build the committed schema-12 project fixtures - with the *previous* application.

Purpose:
    Phase 0.1.1-A adds migration 13. Its upgrade tests must start from
    projects written by the build that came before it (schema 12,
    ``0.1.0-alpha.2`` + unreleased work up to ``0ed96ed``), not from a
    hand-written approximation of one. This script drives that build's own
    services to write two projects, which are committed beside it:

    * ``unique_sets/`` - sets ``10`` and ``11``; Set 10 attended, scanned,
      reconciled, keyed (verified) and scored.
    * ``colliding_sets/`` - sets ``A``, ``a`` and ``B``, which the old build
      allowed because it compared set codes exactly; Set ``A`` attended,
      scanned (one sheet read ``A``, one read ``a``), reconciled, keyed and
      scored. This is the collision migration 13 must keep, never merge.

How to regenerate (only ever with the schema-12 build)::

    git worktree add ../OMRflow-schema12 0ed96ed
    $env:PYTHONPATH = "..\OMRflow-schema12\src"
    .venv\Scripts\python.exe tests\fixtures\schema12\make_schema12_projects.py `
        --old-checkout ..\OMRflow-schema12

The script refuses to run against a build whose schema is not 12.

What does NOT belong here:
    Any assertion. The tests live in
    ``tests/integration/test_set_identity_migration.py``.
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import UTC, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
OPERATOR = "Fixture Builder"


def _use_old_checkout(root: Path) -> None:
    """Import OMRFlow, and its test helpers, from the schema-12 checkout."""
    sys.path.insert(0, str(root / "src"))
    sys.path.insert(1, str(root))


def _attendance(path: Path, rolls: list[str]) -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for column, text in enumerate(("Sl.No.", "Roll No.", "Name", "Total"), start=1):
        sheet.cell(row=1, column=column, value=text)
    for offset, roll in enumerate(rolls):
        sheet.cell(row=2 + offset, column=1, value=offset + 1)
        sheet.cell(row=2 + offset, column=2, value=roll)
        sheet.cell(row=2 + offset, column=3, value=f"CANDIDATE {roll}")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def _register(database, folder: Path, scripts: list[tuple[str, str, int]], numbers) -> str:
    """Register a batch of ``(roll, set code as read, correct answers)`` scripts."""
    from omr_scanner.database.models import BatchScan, ScanBatch
    from omr_scanner.services import batch_store
    from omr_scanner.services.recognition_models import (
        AnswerView,
        FieldView,
        RecognitionOutcome,
        RegistrationStatus,
        ScanResult,
    )

    now = datetime.now(UTC)
    batch_id = batch_store.new_batch_id()
    ordered = list(numbers)
    with database.session() as session:
        session.add(
            ScanBatch(
                batch_id=batch_id, created_at=now, updated_at=now,
                source_folder=str(folder), status="completed", total_scans=len(scripts),
            )
        )
        session.flush()
        for index, (roll, set_read, correct) in enumerate(scripts):
            name = f"sheet_{index:03d}.png"
            result = ScanResult(
                source_path=folder / name,
                outcome=RecognitionOutcome.COMPLETE,
                registration=RegistrationStatus.REGISTERED,
                fields=(
                    FieldView(
                        zone_id="roll_number", label="Roll", field_type="numeric",
                        value=roll, status="complete", needs_review=False, characters=(),
                    ),
                    FieldView(
                        zone_id="set_code", label="Set", field_type="set_code",
                        value=set_read, status="complete", needs_review=False,
                        characters=(),
                    ),
                ),
                answers=tuple(
                    AnswerView(
                        number=number, zone_id="q",
                        value="A" if position < correct else "B",
                        status="resolved", needs_review=False,
                        top_fill=0.9, margin=0.4, confidence=0.9,
                    )
                    for position, number in enumerate(ordered)
                ),
                identifier_zone_id="roll_number", set_code_zone_id="set_code",
            )
            session.add(
                BatchScan(
                    batch_id=batch_id, batch_index=index,
                    source_path=str(folder / name), filename=name, status="completed",
                    identifier_value=roll, set_code_value=set_read,
                    result_json=json.dumps(result.to_dict()),
                )
            )
    return batch_id


def _build(
    workspace: Path,
    name: str,
    sets: list[tuple[str, str]],
    scored_set: str,
    scripts: list[tuple[str, str, int]],
) -> Path:
    from tests.conftest import build_answer_sheet_template

    from omr_scanner.services import (
        candidate_import,
        create_project,
        project_sets,
        reconciliation_store,
        scoring_store,
        set_attendance,
    )
    from omr_scanner.services.answer_key import plan_for, read_key

    template = build_answer_sheet_template()
    plan = plan_for(template)
    session = create_project(workspace, name, exam_name=f"Schema 12 fixture: {name}")
    try:
        database = session.database
        defined = {code: project_sets.add_set(database, code, text) for code, text in sets}
        rolls = sorted({roll for roll, _set, _correct in scripts})
        attendance = _attendance(session.root / "inputs" / "attendance.xlsx", rolls)
        validation = candidate_import.read_roster(attendance)
        assignment = set_attendance.assign_attendance_workbook(
            database, defined[scored_set].set_id, attendance, validation,
            imported_by=OPERATOR,
        )
        batch_id = _register(database, session.root / "scans", scripts, plan.numbers)
        reconciliation_store.reconcile_batch(database, assignment.roster_id, batch_id)
        stored = scoring_store.save_key(
            database, read_key("A" * plan.question_count, plan, scored_set).to_key()
        )
        scoring_store.verify_key(database, stored.key_id, verified_by=OPERATOR)
        scoring_store.score_batch(
            database, assignment.roster_id, batch_id, template, computed_by=OPERATOR
        )
        root = session.root
    finally:
        session.close()
    return root


def main() -> int:
    """Write both fixtures beside this script."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--old-checkout", type=Path, required=True)
    arguments = parser.parse_args()
    _use_old_checkout(arguments.old_checkout.resolve())

    from omr_scanner import __version__
    from omr_scanner.database.migrations import SCHEMA_VERSION

    if SCHEMA_VERSION != 12:
        print(f"Refusing: this build writes schema {SCHEMA_VERSION}, not 12", file=sys.stderr)
        return 2

    workspace = HERE / "_build"
    shutil.rmtree(workspace, ignore_errors=True)
    workspace.mkdir()
    built = {
        "unique_sets": _build(
            workspace,
            "unique_sets",
            [("10", "Electrical"), ("11", "Civil")],
            "10",
            [("10001", "10", 5), ("10002", "10", 3)],
        ),
        "colliding_sets": _build(
            workspace,
            "colliding_sets",
            [("A", "Upper-case A"), ("a", "Lower-case a"), ("B", "Set B")],
            "A",
            [("20001", "A", 5), ("20002", "a", 4)],
        ),
    }
    for name, root in built.items():
        target = HERE / name
        shutil.rmtree(target, ignore_errors=True)
        shutil.copytree(root, target, ignore=shutil.ignore_patterns("backups", "logs"))
    shutil.rmtree(workspace, ignore_errors=True)
    (HERE / "PROVENANCE.json").write_text(
        json.dumps(
            {
                "written_by_version": __version__,
                "schema_version": SCHEMA_VERSION,
                "written_at": datetime.now(UTC).isoformat(timespec="seconds"),
                "projects": sorted(built),
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {', '.join(sorted(built))} with OMRFlow {__version__} (schema 12)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
