r"""Build the committed schema-14 project fixtures - with the schema-14 build.

Purpose:
    Migration 15 (session scope) and its upgrade step must be tested from
    projects the previous build itself wrote. This script drives the
    schema-14 build's own services (``main`` at ``0e94d67``, 0.1.1-alpha.0
    phase 3) to write two projects, committed beside it:

    * ``one_session/`` - one scan session, one batch of Set A, attended,
      reconciled, keyed, scored and reported (an XLSX final export).
    * ``split_state/`` - one scan session of two batches. Set A was reconciled
      and scored while the first batch was the newest; a second batch was then
      read into the same session and Set A reconciled and scored again - which
      the schema-14 build did against that newer batch alone. Its downstream
      state is therefore in **both** batches; the upgrade must keep the one
      the schema-14 build showed (the newer).

How to regenerate (only ever with the schema-14 build)::

    .venv\Scripts\python.exe tests\fixtures\schema14\make_schema14_projects.py `
        --old-checkout <checkout of 0e94d67>

The script refuses to run against a build whose schema is not 14.
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import UTC, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
OPERATOR = "Fixture Builder"


def _use_old_checkout(root: Path) -> None:
    sys.path.insert(0, str(root / "src"))
    sys.path.insert(1, str(root))


def _attendance(path: Path, rolls: list[str]) -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for column, text in enumerate(("Sl.No.", "Roll No.", "Name", "Total"), start=1):
        sheet.cell(row=1, column=column, value=text)
    for offset, roll in enumerate(rolls):
        sheet.cell(row=2 + offset, column=1, value=offset + 1)
        sheet.cell(row=2 + offset, column=2, value=roll)
        sheet.cell(row=2 + offset, column=3, value=f"CANDIDATE {roll}")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def _read_batch(database, folder: Path, scripts, template) -> str:
    """A Process All: a new batch of the active session (sealing the last), read."""
    from sqlalchemy import update
    from tests.integration.test_reject_and_rescan import make_result

    from omr_scanner.database.models import BatchScan
    from omr_scanner.services import batch_store, scan_sessions

    folder.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, _roll, _correct in scripts:
        path = folder / name
        path.write_bytes(f"schema14:{name}".encode())
        paths.append(path)
    batch_id = scan_sessions.start_batch(
        database, paths, identity=batch_store.BatchIdentity.of(template), started_by=OPERATOR
    )
    ids = batch_store.scan_ids_by_path(database, batch_id)
    with database.session() as session:
        for path, (_name, roll, correct) in zip(paths, scripts, strict=True):
            session.execute(
                update(BatchScan)
                .where(BatchScan.scan_id == ids[path])
                .values(
                    status="completed", identifier_value=roll, set_code_value="A",
                    result_json=json.dumps(make_result(path, roll, "A", correct).to_dict()),
                )
            )
    batch_store.finalise_batch(database, batch_id)
    return batch_id


def _reconcile_and_score(session, roster_id: int, batch_id: str, template) -> None:
    from omr_scanner.services import reconciliation_store, scoring_store

    reconciliation_store.reconcile_batch(session.database, roster_id, batch_id)
    scoring_store.score_batch(
        session.database, roster_id, batch_id, template, computed_by=OPERATOR
    )


def _set_up(workspace: Path, name: str, rolls: list[str], template):
    from omr_scanner.services import (
        candidate_import,
        create_project,
        project_sets,
        scoring_store,
        set_attendance,
    )
    from omr_scanner.services.answer_key import plan_for, read_key

    session = create_project(workspace, name, exam_name=f"Schema 14 fixture: {name}")
    exam_set = project_sets.add_set(session.database, "A", "Set A")
    attendance = _attendance(session.root / "inputs" / "attendance.xlsx", rolls)
    assignment = set_attendance.assign_attendance_workbook(
        session.database, exam_set.set_id, attendance, candidate_import.read_roster(attendance),
        imported_by=OPERATOR,
    )
    plan = plan_for(template)
    stored = scoring_store.save_key(
        session.database, read_key("A" * plan.question_count, plan, "A").to_key()
    )
    scoring_store.verify_key(session.database, stored.key_id, verified_by=OPERATOR)
    return session, exam_set, assignment.roster_id


def main() -> int:
    """Write the two fixtures beside this script."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--old-checkout", type=Path, required=True)
    arguments = parser.parse_args()
    _use_old_checkout(arguments.old_checkout.resolve())

    from tests.conftest import build_answer_sheet_template

    from omr_scanner import __version__
    from omr_scanner.database.migrations import SCHEMA_VERSION
    from omr_scanner.services import report_store

    if SCHEMA_VERSION != 14:
        print(f"Refusing: this build writes schema {SCHEMA_VERSION}, not 14", file=sys.stderr)
        return 2

    template = build_answer_sheet_template()
    workspace = HERE / "_build"
    shutil.rmtree(workspace, ignore_errors=True)
    workspace.mkdir()
    built: dict[str, Path] = {}

    rolls = ["40001", "40002", "40003"]
    session, exam_set, roster = _set_up(workspace, "one_session", rolls, template)
    batch = _read_batch(session.database, session.root / "scans", [
        ("a1.png", rolls[0], 5), ("a2.png", rolls[1], 3), ("a3.png", rolls[2], 0),
    ], template)
    _reconcile_and_score(session, roster, batch, template)
    outcome = report_store.generate_for_set(
        session.database, exam_set.set_id, batch, template, project_name=session.name,
        output_dir=session.root / "exports", computed_by=OPERATOR, final=True,
    )
    if not outcome.ok:
        raise SystemExit(f"report generation failed: {outcome.status} {outcome.warnings}")
    built["one_session"] = session.root
    session.close()

    rolls = ["41001", "41002", "41003"]
    session, exam_set, roster = _set_up(workspace, "split_state", rolls, template)
    first = _read_batch(session.database, session.root / "s1", [
        ("b1.png", rolls[0], 5), ("b2.png", rolls[1], 2),
    ], template)
    _reconcile_and_score(session, roster, first, template)
    second = _read_batch(session.database, session.root / "s2", [("b3.png", rolls[2], 4)], template)
    _reconcile_and_score(session, roster, second, template)
    built["split_state"] = session.root
    session.close()

    for name, root in built.items():
        target = HERE / name
        shutil.rmtree(target, ignore_errors=True)
        shutil.copytree(root, target, ignore=shutil.ignore_patterns("backups", "logs"))
    shutil.rmtree(workspace, ignore_errors=True)
    (HERE / "PROVENANCE.json").write_text(
        json.dumps(
            {
                "written_by_version": __version__,
                "schema_version": SCHEMA_VERSION,
                "written_at": datetime.now(UTC).isoformat(timespec="seconds"),
                "projects": sorted(built),
                "first_batch_of_split_state": first,
                "second_batch_of_split_state": second,
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {', '.join(sorted(built))} with OMRFlow {__version__} (schema 14)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
