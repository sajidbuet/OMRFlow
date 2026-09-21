"""Workbook-level tests for the Meritwise sheet (Part 2, §10-§12, §25).

What these lock down:
    That ``meritwise`` is a *copy of the completed Rollwise sheet* - the
    institution's heading, the post name, the column widths, the borders, the
    print setup and the logo all present - with only the row membership and
    the row order different. Every claim about preserved formatting is
    asserted against the produced file, never assumed, which is what §25
    asks for in as many words.

    The fixture workbook deliberately carries every feature §25 names, so a
    regression in any one of them fails here rather than in somebody's
    printed merit list.
"""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from omr_scanner.reporting import excel as rx
from omr_scanner.services.report_template import preview_template, read_template

HEADER_ROW = 3
FIRST_DATA_ROW = 4

# (roll, name, mark) - two absentees, a tie on 70, a zero and a full mark.
CANDIDATES: tuple[tuple[str, str, float | None], ...] = (
    ("10001", "Alia Rahman", 70.0),
    ("10002", "Bashir Uddin", None),      # absent
    ("10003", "Chameli Akter", 90.0),     # full marks
    ("10004", "Delwar Hossain", 70.0),    # ties with 10001
    ("10005", "Eshita Nasrin", 0.0),      # zero
    ("10006", "Farid Mia", None),         # absent
)


def _png_bytes() -> bytes:
    """A real, decodable PNG - openpyxl refuses anything else."""
    pillow = pytest.importorskip("PIL.Image", reason="Pillow is required for image support")
    buffer = io.BytesIO()
    pillow.new("RGB", (32, 32), (0, 80, 176)).save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture
def attendance_workbook(tmp_path: Path) -> Path:
    """An attendance sheet shaped like a real examination office's."""
    path = tmp_path / "set10_attendance.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"

    sheet["A1"] = "Bangladesh Submarine Cable Regulatory Authority"
    sheet.merge_cells("A1:E1")
    sheet["A1"].font = Font(bold=True, size=16, color="FF0000CC")
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet["A2"] = "Name of Post: Assistant Engineer (Electrical)"
    sheet.merge_cells("A2:E2")

    for column, text in enumerate(("Sl.No.", "Roll No.", "Name", "Total", "Merit"), start=1):
        cell = sheet.cell(row=HEADER_ROW, column=column, value=text)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFDDDDDD")
        cell.border = Border(bottom=Side(style="thick"))

    for offset, (roll, name, _mark) in enumerate(CANDIDATES):
        row = FIRST_DATA_ROW + offset
        sheet.cell(row=row, column=1, value=offset + 1)
        sheet.cell(row=row, column=2, value=roll)
        sheet.cell(row=row, column=3, value=name).alignment = Alignment(wrap_text=True)
        sheet.cell(row=row, column=4).number_format = "0.00"
        sheet.cell(row=row, column=4).border = Border(left=Side(style="thin"))
        sheet.row_dimensions[row].height = 20 + offset

    # Something printed *below* the table, which must not move.
    footer_row = FIRST_DATA_ROW + len(CANDIDATES) + 1
    sheet.cell(row=footer_row, column=1, value="Signature of the Controller of Examinations")

    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["C"].width = 32
    sheet.freeze_panes = "A4"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = 9
    sheet.page_setup.fitToWidth = 1
    sheet.page_margins.left = 0.25
    sheet.print_title_rows = "1:3"
    sheet.print_area = "A1:E12"
    sheet.oddHeader.center.text = "OFFICIAL"
    sheet.oddFooter.right.text = "Page &P"
    sheet["G1"] = "=COUNTA(B4:B9)"

    try:
        from openpyxl.drawing.image import Image as ExcelImage

        logo = tmp_path / "logo.png"
        logo.write_bytes(_png_bytes())
        sheet.add_image(ExcelImage(str(logo)), "E1")
    except Exception:  # pragma: no cover - no Pillow; asserted separately
        pass

    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture
def generated(attendance_workbook: Path, tmp_path: Path) -> tuple[Path, object]:
    """The attendance workbook, copied, Rollwise-populated and Meritwise-built."""
    output = tmp_path / "Set_10_Result.xlsx"
    rx.copy_into(attendance_workbook, output)

    roster = read_template(
        attendance_workbook, preview_template(attendance_workbook).suggestion.to_mapping()
    )
    decisions = {
        roll: rx.CandidateReportRow(
            roll=roll, is_absent=mark is None,
            final_score=None if mark is None else __import__("fractions").Fraction(mark),
        )
        for roll, _name, mark in CANDIDATES
    }
    rx.populate_rollwise(output, roster, decisions)

    present = [
        (roll, mark) for roll, _name, mark in CANDIDATES if mark is not None
    ]
    # Merit order: score descending, roll ascending on a tie - the same rule
    # the rest of the application ranks by.
    present.sort(key=lambda item: (-item[1], item[0]))
    by_roll = {row.roll: row.sheet_row_number for row in roster.rows}
    ordered = tuple(
        rx.MeritwiseRow(source_row_number=by_roll[roll], roll=roll) for roll, _ in present
    )
    result = rx.build_meritwise_from_rollwise(
        output, roster, ordered,
        serial_column=1, rank_column=5, marks_column=4, scratch_dir=tmp_path,
    )
    return output, result


class TestTheSheetItself:
    def test_it_exists_and_is_named_exactly_meritwise(self, generated) -> None:
        output, _ = generated
        workbook = openpyxl.load_workbook(output)
        assert "meritwise" in workbook.sheetnames
        assert rx.MERITWISE_SHEET_NAME == "meritwise"
        workbook.close()

    def test_the_rollwise_sheet_is_still_there_unchanged(self, generated) -> None:
        output, _ = generated
        workbook = openpyxl.load_workbook(output)
        rollwise = workbook["Attendance"]
        rolls = [rollwise.cell(row=FIRST_DATA_ROW + i, column=2).value for i in range(6)]
        assert rolls == [roll for roll, _n, _m in CANDIDATES]
        workbook.close()


class TestMembershipAndOrder:
    def test_absent_candidates_are_removed(self, generated) -> None:
        output, _ = generated
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.MERITWISE_SHEET_NAME]
        rolls = {
            sheet.cell(row=row, column=2).value
            for row in range(FIRST_DATA_ROW, FIRST_DATA_ROW + 10)
        }
        assert "10002" not in rolls
        assert "10006" not in rolls
        workbook.close()

    def test_present_candidates_all_remain(self, generated) -> None:
        output, _ = generated
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.MERITWISE_SHEET_NAME]
        rolls = [
            sheet.cell(row=FIRST_DATA_ROW + offset, column=2).value for offset in range(4)
        ]
        assert set(rolls) == {"10001", "10003", "10004", "10005"}
        workbook.close()

    def test_rows_are_in_merit_order(self, generated) -> None:
        output, _ = generated
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.MERITWISE_SHEET_NAME]
        ordered = [
            (
                sheet.cell(row=FIRST_DATA_ROW + offset, column=2).value,
                sheet.cell(row=FIRST_DATA_ROW + offset, column=4).value,
            )
            for offset in range(4)
        ]
        assert ordered == [("10003", 90.0), ("10001", 70.0), ("10004", 70.0), ("10005", 0.0)]
        workbook.close()

    def test_a_tie_is_broken_by_roll_ascending_not_by_rank(self, generated) -> None:
        """Both 70s keep the same merit; only their row order is decided."""
        output, _ = generated
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.MERITWISE_SHEET_NAME]
        assert sheet.cell(row=FIRST_DATA_ROW + 1, column=2).value == "10001"
        assert sheet.cell(row=FIRST_DATA_ROW + 2, column=2).value == "10004"
        # RANK.EQ gives tied marks the same rank, which is the whole reason
        # the rank is a formula rather than the row's position. The project's
        # own `rank_formula` wraps it so an absence shows the absence marker
        # instead of a number - that wrapper must come across unchanged.
        for offset in (1, 2):
            formula = sheet.cell(row=FIRST_DATA_ROW + offset, column=5).value
            assert isinstance(formula, str)
            assert "RANK.EQ(" in formula
            # Ranked over this sheet's own rows, not Rollwise's.
            assert "$D$4:$D$7" in formula
        workbook.close()

    def test_the_serial_column_is_renumbered_for_the_merit_listing(self, generated) -> None:
        output, _ = generated
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.MERITWISE_SHEET_NAME]
        serials = [
            sheet.cell(row=FIRST_DATA_ROW + offset, column=1).value for offset in range(4)
        ]
        assert serials == [1, 2, 3, 4]
        workbook.close()

    def test_the_freed_rows_are_blanked_not_left_holding_an_absentee(
        self, generated
    ) -> None:
        output, result = generated
        assert result.rows_removed == 2
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.MERITWISE_SHEET_NAME]
        for offset in (4, 5):
            assert sheet.cell(row=FIRST_DATA_ROW + offset, column=2).value is None
        workbook.close()

    def test_content_printed_below_the_table_does_not_move(self, generated) -> None:
        output, _ = generated
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.MERITWISE_SHEET_NAME]
        footer_row = FIRST_DATA_ROW + len(CANDIDATES) + 1
        assert sheet.cell(row=footer_row, column=1).value == (
            "Signature of the Controller of Examinations"
        )
        workbook.close()


class TestFormattingIsCarriedOver:
    """§25's "verify representative examples" list, one assertion each."""

    @pytest.fixture
    def sheets(self, generated):
        output, _ = generated
        workbook = openpyxl.load_workbook(output)
        yield workbook[rx.MERITWISE_SHEET_NAME], workbook["Attendance"]
        workbook.close()

    def test_the_institution_heading_comes_with_it(self, sheets) -> None:
        merit, _ = sheets
        assert merit["A1"].value == "Bangladesh Submarine Cable Regulatory Authority"

    def test_the_post_name_comes_with_it(self, sheets) -> None:
        merit, _ = sheets
        assert merit["A2"].value == "Name of Post: Assistant Engineer (Electrical)"

    def test_merged_cells(self, sheets) -> None:
        merit, rollwise = sheets
        assert sorted(map(str, merit.merged_cells.ranges)) == sorted(
            map(str, rollwise.merged_cells.ranges)
        )

    def test_fonts(self, sheets) -> None:
        merit, rollwise = sheets
        assert merit["A1"].font.bold == rollwise["A1"].font.bold
        assert merit["A1"].font.size == rollwise["A1"].font.size
        assert merit["A1"].font.color.rgb == rollwise["A1"].font.color.rgb

    def test_fills(self, sheets) -> None:
        merit, rollwise = sheets
        assert merit.cell(row=HEADER_ROW, column=1).fill.fgColor.rgb == (
            rollwise.cell(row=HEADER_ROW, column=1).fill.fgColor.rgb
        )

    def test_borders(self, sheets) -> None:
        merit, rollwise = sheets
        assert merit.cell(row=HEADER_ROW, column=1).border.bottom.style == (
            rollwise.cell(row=HEADER_ROW, column=1).border.bottom.style
        )

    def test_a_data_rows_border_travels_with_the_row(self, sheets) -> None:
        merit, _ = sheets
        assert merit.cell(row=FIRST_DATA_ROW, column=4).border.left.style == "thin"

    def test_number_formats(self, sheets) -> None:
        merit, _ = sheets
        assert merit.cell(row=FIRST_DATA_ROW, column=4).number_format == "0.00"

    def test_alignment(self, sheets) -> None:
        merit, _ = sheets
        assert merit.cell(row=FIRST_DATA_ROW, column=3).alignment.wrap_text is True

    def test_column_widths(self, sheets) -> None:
        merit, rollwise = sheets
        assert merit.column_dimensions["C"].width == rollwise.column_dimensions["C"].width

    def test_row_heights_travel_with_their_row(self, sheets) -> None:
        merit, _ = sheets
        # 10003 was row 6 (offset 2, height 22) and is now the first row.
        assert merit.row_dimensions[FIRST_DATA_ROW].height == 22

    def test_freeze_panes(self, sheets) -> None:
        merit, rollwise = sheets
        assert merit.freeze_panes == rollwise.freeze_panes == "A4"

    def test_page_orientation(self, sheets) -> None:
        merit, _ = sheets
        assert merit.page_setup.orientation == "landscape"

    def test_paper_size_and_scaling(self, sheets) -> None:
        merit, _ = sheets
        assert merit.page_setup.paperSize == 9
        assert merit.page_setup.fitToWidth == 1

    def test_margins(self, sheets) -> None:
        merit, _ = sheets
        assert merit.page_margins.left == 0.25

    def test_print_titles(self, sheets) -> None:
        merit, _ = sheets
        assert merit.print_title_rows == "$1:$3"

    def test_print_area(self, sheets) -> None:
        merit, _ = sheets
        assert merit.print_area is not None
        assert "$A$1:$E$12" in str(merit.print_area)

    def test_headers_and_footers(self, sheets) -> None:
        merit, _ = sheets
        assert merit.oddHeader.center.text == "OFFICIAL"
        assert merit.oddFooter.right.text == "Page &P"

    def test_gridline_setting(self, sheets) -> None:
        merit, _ = sheets
        assert merit.sheet_view.showGridLines is False

    def test_an_unrelated_formula_survives(self, sheets) -> None:
        merit, _ = sheets
        assert merit["G1"].value == "=COUNTA(B4:B9)"


class TestLogo:
    def test_the_logo_is_carried_onto_the_meritwise_sheet(self, generated) -> None:
        pytest.importorskip("PIL.Image", reason="Pillow is required for image support")
        output, result = generated
        assert result.images_copied == 1
        workbook = openpyxl.load_workbook(output)
        assert len(workbook["Attendance"]._images) == 1
        assert len(workbook[rx.MERITWISE_SHEET_NAME]._images) == 1
        workbook.close()

    def test_it_keeps_its_anchor_and_size(self, generated) -> None:
        pytest.importorskip("PIL.Image", reason="Pillow is required for image support")
        output, _ = generated
        workbook = openpyxl.load_workbook(output)
        original = workbook["Attendance"]._images[0]
        copied = workbook[rx.MERITWISE_SHEET_NAME]._images[0]
        assert (copied.anchor._from.col, copied.anchor._from.row) == (
            original.anchor._from.col, original.anchor._from.row
        )
        assert (copied.width, copied.height) == (original.width, original.height)
        workbook.close()


class TestSourceFileIntegrity:
    def test_the_attendance_workbook_is_byte_for_byte_unchanged(
        self, attendance_workbook: Path, tmp_path: Path
    ) -> None:
        """§17/§26: hash before, generate, hash after."""
        import hashlib

        before = hashlib.sha256(attendance_workbook.read_bytes()).hexdigest()

        output = tmp_path / "Result.xlsx"
        rx.copy_into(attendance_workbook, output)
        roster = read_template(
            attendance_workbook,
            preview_template(attendance_workbook).suggestion.to_mapping(),
        )
        rx.populate_rollwise(output, roster, {})
        rx.build_meritwise_from_rollwise(
            output, roster,
            (rx.MeritwiseRow(source_row_number=FIRST_DATA_ROW, roll="10001"),),
            serial_column=1, rank_column=5, marks_column=4, scratch_dir=tmp_path,
        )

        after = hashlib.sha256(attendance_workbook.read_bytes()).hexdigest()
        assert after == before
