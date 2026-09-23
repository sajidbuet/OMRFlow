"""Tests for the synthetic attendance/absentee population and its workbooks.

Scope:
    The half of a synthetic examination that is paperwork rather than pixels:
    who exists, who attended, what the workbook claims, what the scan shows,
    and what Phase 7 should conclude from the disagreement.

    ===== ==========================================================
    Test  Behaviour
    ===== ==========================================================
    A     The authoritative roster: size, identity, sets, names.
    B     Every conflict is staged, and stages exactly what it says.
    C     Conflicts are mutually exclusive.
    D     Determinism, and sensitivity to the seed.
    E     The workbooks OMRFlow's own importer can read.
    F     Ground truth keeps the three states apart.
    G     Rates, profiles and their limits.
    ===== ==========================================================

Why the expected states are asserted against a table here:
    The point of this dataset is to be an answer key for the reconciliation
    engine. If these tests derived the expected state the same way the module
    does, they would assert only that the module is self-consistent. Each case
    below therefore names the state a human decided is correct for that
    situation - which is the thing a future refactor must not quietly change.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from omr_scanner.domain.reconciliation import AttendanceState, ReconciliationStatus
from omr_scanner.evaluation.attendance_dataset import (
    ABSENT_TOKEN,
    HEADERS,
    SHEET_TITLE,
    ConflictKind,
    ConflictProfile,
    ConflictRates,
    Population,
    expected_states,
    plan_population,
    summarise,
    write_ground_truth,
    write_workbooks,
)

SETS = ("10", "11", "12")
SEED = 20260923


@pytest.fixture
def population() -> Population:
    """A population large enough for the quotas to bite, small enough to be fast."""
    return plan_population(count=400, set_codes=SETS, seed=SEED)


def one_with(population: Population, kind: ConflictKind):
    """The first candidate staged with ``kind``."""
    for candidate in population.candidates:
        if candidate.conflict is kind:
            return candidate
    raise AssertionError(f"no candidate was staged with {kind}")


# ----------------------------------------------------------------------
# A - the authoritative roster
# ----------------------------------------------------------------------
class TestARoster:
    def test_the_requested_number_of_candidates_exists(self, population: Population):
        assert len(population.candidates) == 400

    def test_candidate_uids_are_unique(self, population: Population):
        uids = [candidate.candidate_uid for candidate in population.candidates]
        assert len(set(uids)) == len(uids)

    def test_registered_rolls_are_unique(self, population: Population):
        rolls = [candidate.roll for candidate in population.candidates]
        assert len(set(rolls)) == len(rolls)

    def test_rolls_are_zero_padded_to_a_constant_width(self, population: Population):
        widths = {len(candidate.roll) for candidate in population.candidates}
        assert len(widths) == 1

    def test_every_set_is_used(self, population: Population):
        assert set(population.by_set()) == set(SETS)

    def test_names_are_synthetic(self, population: Population):
        """Nothing in a generated dataset may look like a real person."""
        for candidate in population.candidates:
            assert candidate.name.startswith("Candidate ")

    def test_a_population_needs_a_count_and_a_set(self):
        with pytest.raises(ValueError):
            plan_population(count=0, set_codes=SETS, seed=SEED)
        with pytest.raises(ValueError):
            plan_population(count=10, set_codes=[], seed=SEED)


# ----------------------------------------------------------------------
# B - each conflict stages what it claims
# ----------------------------------------------------------------------
class TestBConflictsAreStagedCorrectly:
    def test_every_conflict_kind_is_represented(self, population: Population):
        """The edge-case guarantee: a dataset that omits a case cannot prove it."""
        staged = {candidate.conflict for candidate in population.candidates}
        assert staged == set(ConflictKind)

    def test_a_clean_candidate_is_present_scanned_and_correct(
        self, population: Population
    ):
        candidate = one_with(population, ConflictKind.NONE)
        assert candidate.true_attendance is AttendanceState.PRESENT
        assert candidate.attendance_status is AttendanceState.PRESENT
        assert candidate.scan_present is True
        assert candidate.observed_roll == candidate.roll
        assert candidate.observed_set == candidate.set_code
        assert candidate.expected_status is ReconciliationStatus.MATCHED
        assert candidate.manual_review_required is False

    def test_a_genuine_absentee_has_no_script_and_is_listed_absent(
        self, population: Population
    ):
        candidate = one_with(population, ConflictKind.TRUE_ABSENTEE)
        assert candidate.true_attendance is AttendanceState.ABSENT
        assert candidate.attendance_status is AttendanceState.ABSENT
        assert candidate.scan_present is False
        assert candidate.expected_status is ReconciliationStatus.ABSENT_CONFIRMED
        assert candidate.manual_review_required is False

    def test_present_but_marked_absent_keeps_its_script(self, population: Population):
        candidate = one_with(population, ConflictKind.MARKED_ABSENT_BUT_PRESENT)
        assert candidate.true_attendance is AttendanceState.PRESENT
        assert candidate.attendance_status is AttendanceState.ABSENT
        assert candidate.scan_present is True
        assert candidate.observed_roll == candidate.roll
        assert candidate.expected_status is ReconciliationStatus.ABSENT_WITH_SCRIPT

    def test_absent_but_marked_present_has_no_script(self, population: Population):
        candidate = one_with(population, ConflictKind.MARKED_PRESENT_BUT_ABSENT)
        assert candidate.true_attendance is AttendanceState.ABSENT
        assert candidate.attendance_status is AttendanceState.PRESENT
        assert candidate.scan_present is False
        assert candidate.expected_status is ReconciliationStatus.PRESENT_WITHOUT_SCRIPT

    @pytest.mark.parametrize(
        "kind",
        [
            ConflictKind.BLANK_CANDIDATE_ID,
            ConflictKind.PARTIAL_CANDIDATE_ID,
            ConflictKind.CANDIDATE_ID_MULTIPLE_MARK,
        ],
    )
    def test_an_unusable_identifier_leaves_the_script_unattributed(
        self, population: Population, kind: ConflictKind
    ):
        """The script is real; only the identifier is not usable.

        All three land in the same state on purpose - the engine cannot tell a
        blank field from an over-marked one and should not guess - but which
        defect it was stays recorded in the conflict type.
        """
        candidate = one_with(population, kind)
        assert candidate.scan_present is True
        assert candidate.observed_roll is None
        assert candidate.expected_status is (
            ReconciliationStatus.UNRESOLVED_CANDIDATE_ID
        )
        assert candidate.manual_review_required is True

    def test_a_wrong_identifier_moves_the_script_to_a_stray(
        self, population: Population
    ):
        """The candidate loses their script; a script with no owner appears."""
        candidate = one_with(population, ConflictKind.WRONG_CANDIDATE_ID)
        assert candidate.scan_present is False
        assert candidate.expected_status is ReconciliationStatus.PRESENT_WITHOUT_SCRIPT

        strays = [
            sheet
            for sheet in population.stray_sheets
            if sheet.candidate_uid == f"{candidate.candidate_uid}-STRAY"
        ]
        assert len(strays) == 1
        assert strays[0].scan_present is True
        assert strays[0].observed_roll != candidate.roll

    def test_a_wrong_identifier_never_collides_with_a_real_candidate(
        self, population: Population
    ):
        """Otherwise it would become an accidental duplicate nobody designed."""
        registered = {candidate.roll for candidate in population.candidates}
        for sheet in population.stray_sheets:
            if sheet.conflict is ConflictKind.WRONG_CANDIDATE_ID:
                assert sheet.observed_roll not in registered

    def test_an_unknown_identifier_is_outside_the_roster(
        self, population: Population
    ):
        candidate = one_with(population, ConflictKind.UNKNOWN_CANDIDATE_ID)
        registered = {other.roll for other in population.candidates}
        assert candidate.observed_roll not in registered
        assert candidate.expected_status is ReconciliationStatus.UNKNOWN_ID

    def test_a_duplicate_produces_a_second_script_with_the_same_roll(
        self, population: Population
    ):
        candidate = one_with(population, ConflictKind.DUPLICATE_SCRIPT)
        duplicates = [
            sheet
            for sheet in population.stray_sheets
            if sheet.duplicate_of == candidate.candidate_uid
        ]
        assert len(duplicates) == 1
        assert duplicates[0].observed_roll == candidate.roll
        assert duplicates[0].scan_present is True
        # Both images must survive - the second must not overwrite the first.
        assert duplicates[0].candidate_uid != candidate.candidate_uid
        assert candidate.expected_status is ReconciliationStatus.DUPLICATE_SCRIPT

    def test_a_missing_scan_leaves_a_present_candidate_without_one(
        self, population: Population
    ):
        candidate = one_with(population, ConflictKind.MISSING_SCAN)
        assert candidate.true_attendance is AttendanceState.PRESENT
        assert candidate.attendance_status is AttendanceState.PRESENT
        assert candidate.scan_present is False
        assert candidate.expected_status is ReconciliationStatus.PRESENT_WITHOUT_SCRIPT

    def test_a_wrong_set_marks_a_different_configured_set(
        self, population: Population
    ):
        candidate = one_with(population, ConflictKind.WRONG_SET)
        assert candidate.observed_set in SETS
        assert candidate.observed_set != candidate.set_code
        assert candidate.observed_roll == candidate.roll
        # The script still belongs to its owner; only the key is in doubt.
        assert candidate.expected_status is ReconciliationStatus.MATCHED
        assert candidate.manual_review_required is True

    def test_a_blank_set_marks_nothing(self, population: Population):
        candidate = one_with(population, ConflictKind.BLANK_SET)
        assert candidate.observed_set is None
        assert candidate.observed_roll == candidate.roll
        assert candidate.manual_review_required is True


# ----------------------------------------------------------------------
# C - mutual exclusivity
# ----------------------------------------------------------------------
class TestCConflictsDoNotOverlap:
    def test_each_candidate_carries_exactly_one_conflict(
        self, population: Population
    ):
        """A compound failure that arose by accident has no defensible truth."""
        for candidate in population.candidates:
            assert isinstance(candidate.conflict, ConflictKind)

    def test_an_absent_candidate_never_also_has_a_script(
        self, population: Population
    ):
        for candidate in population.candidates:
            if candidate.true_attendance is AttendanceState.ABSENT:
                assert candidate.scan_present is False

    def test_sheets_to_render_matches_the_scan_flags(self, population: Population):
        rendered = population.sheets_to_render()
        assert all(sheet.scan_present for sheet in rendered)
        expected = sum(
            1
            for entry in (*population.candidates, *population.stray_sheets)
            if entry.scan_present
        )
        assert len(rendered) == expected

    def test_every_rendered_sheet_has_a_unique_identity(
        self, population: Population
    ):
        """Duplicate *rolls* are intended; duplicate file identities are not."""
        uids = [sheet.candidate_uid for sheet in population.sheets_to_render()]
        assert len(set(uids)) == len(uids)


# ----------------------------------------------------------------------
# D - determinism
# ----------------------------------------------------------------------
class TestDDeterminism:
    def test_the_same_seed_reproduces_the_whole_plan(self):
        first = plan_population(count=250, set_codes=SETS, seed=SEED)
        second = plan_population(count=250, set_codes=SETS, seed=SEED)
        assert first.candidates == second.candidates
        assert first.stray_sheets == second.stray_sheets
        assert expected_states(first) == expected_states(second)

    def test_a_different_seed_changes_who_is_affected(self):
        first = plan_population(count=250, set_codes=SETS, seed=SEED)
        other = plan_population(count=250, set_codes=SETS, seed=SEED + 1)
        assert first.candidates != other.candidates

    def test_a_different_seed_keeps_the_same_composition(self):
        """Which candidates differ; how many of each conflict does not.

        Quota-based assignment is what makes this true, and it is why a small
        dataset can be trusted to contain the cases it advertises.
        """
        first = plan_population(count=250, set_codes=SETS, seed=SEED)
        other = plan_population(count=250, set_codes=SETS, seed=SEED + 1)
        assert first.counts() == other.counts()

    def test_the_plan_does_not_depend_on_evaluation_order(self):
        """Nothing here may be decided by the order sheets happen to render.

        Planning twice while consuming the results differently must give the
        same answer, which is the property multi-worker rendering relies on.
        """
        first = plan_population(count=120, set_codes=SETS, seed=SEED)
        list(reversed(first.sheets_to_render()))
        second = plan_population(count=120, set_codes=SETS, seed=SEED)
        assert first.candidates == second.candidates


# ----------------------------------------------------------------------
# E - the workbooks
# ----------------------------------------------------------------------
class TestEWorkbooks:
    def test_one_workbook_per_set(self, population: Population, tmp_path: Path):
        written = write_workbooks(population, tmp_path)
        assert len(written) == len(SETS)
        assert {path.name for path in written} == {
            f"Set_{code}_Attendance.xlsx" for code in SETS
        }

    def test_the_workbook_uses_the_layout_the_importer_expects(
        self, population: Population, tmp_path: Path
    ):
        import openpyxl

        written = write_workbooks(population, tmp_path)
        workbook = openpyxl.load_workbook(written[0])
        assert workbook.sheetnames == [SHEET_TITLE]
        rows = list(workbook[SHEET_TITLE].iter_rows(values_only=True))
        assert tuple(rows[0]) == HEADERS

    def test_omrflow_s_own_importer_reads_every_workbook(
        self, population: Population, tmp_path: Path
    ):
        """The check that matters.

        Not "a file was written" but "OMRFlow can import it". Anything else is
        a format invented to match the generator.
        """
        from omr_scanner.services.candidate_import import read_roster

        grouped = population.by_set()
        for path in write_workbooks(population, tmp_path):
            set_code = path.stem.split("_")[1]
            result = read_roster(path)
            assert result.issues == ()
            assert len(result.candidates) == len(grouped[set_code])

    def test_the_absentees_in_the_workbook_are_the_ones_the_plan_staged(
        self, population: Population, tmp_path: Path
    ):
        """Including the ones the workbook is *wrong* about.

        A candidate wrongly marked absent must appear as absent in the file -
        that disagreement is the test case, and a workbook that quietly told
        the truth would remove it.
        """
        from omr_scanner.services.candidate_import import read_roster

        grouped = population.by_set()
        for path in write_workbooks(population, tmp_path):
            set_code = path.stem.split("_")[1]
            result = read_roster(path)
            listed_absent = {
                candidate.candidate_id
                for candidate in result.candidates
                if candidate.imported_attendance is AttendanceState.ABSENT
            }
            planned_absent = {
                candidate.roll
                for candidate in grouped[set_code]
                if candidate.attendance_status is AttendanceState.ABSENT
            }
            assert listed_absent == planned_absent

    def test_a_wrongly_marked_absentee_really_is_in_the_file(
        self, population: Population, tmp_path: Path
    ):
        from omr_scanner.services.candidate_import import read_roster

        candidate = one_with(population, ConflictKind.MARKED_ABSENT_BUT_PRESENT)
        path = tmp_path / f"Set_{candidate.set_code}_Attendance.xlsx"
        write_workbooks(population, tmp_path)
        result = read_roster(path)
        row = next(
            entry
            for entry in result.candidates
            if entry.candidate_id == candidate.roll
        )
        assert row.imported_attendance is AttendanceState.ABSENT
        assert row.imported_value.strip().upper() == ABSENT_TOKEN

    def test_conflicts_do_not_leak_between_sets(
        self, population: Population, tmp_path: Path
    ):
        """Each workbook holds only its own set's candidates."""
        from omr_scanner.services.candidate_import import read_roster

        grouped = population.by_set()
        for path in write_workbooks(population, tmp_path):
            set_code = path.stem.split("_")[1]
            imported = {entry.candidate_id for entry in read_roster(path).candidates}
            assert imported == {c.roll for c in grouped[set_code]}
            for other_code, members in grouped.items():
                if other_code != set_code:
                    assert imported.isdisjoint({c.roll for c in members})


# ----------------------------------------------------------------------
# F - ground truth
# ----------------------------------------------------------------------
class TestFGroundTruth:
    def test_both_files_are_written(self, population: Population, tmp_path: Path):
        candidates, reconciliation = write_ground_truth(population, tmp_path)
        assert candidates.is_file()
        assert reconciliation.is_file()

    def test_the_candidate_file_holds_only_the_truth(
        self, population: Population, tmp_path: Path
    ):
        """It must not contain the deliberately corrupted attendance.

        A truth file that already carries the errors is not a truth file, and a
        test comparing against it would be comparing the paperwork with itself.
        """
        import csv

        candidates, _ = write_ground_truth(population, tmp_path)
        with candidates.open(encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        assert len(rows) == len(population.candidates)
        by_uid = {row["candidate_uid"]: row for row in rows}
        wrong = one_with(population, ConflictKind.MARKED_ABSENT_BUT_PRESENT)
        assert by_uid[wrong.candidate_uid]["true_attendance"] == "present"

    def test_the_reconciliation_file_separates_the_three_states(
        self, population: Population, tmp_path: Path
    ):
        import csv

        _, reconciliation = write_ground_truth(population, tmp_path)
        with reconciliation.open(encoding="utf-8") as handle:
            rows = {row["candidate_uid"]: row for row in csv.DictReader(handle)}

        wrong = one_with(population, ConflictKind.MARKED_ABSENT_BUT_PRESENT)
        row = rows[wrong.candidate_uid]
        assert row["true_attendance"] == "present"       # what happened
        assert row["attendance_status"] == "absent"      # what the file claims
        assert row["scan_present"] == "true"             # what was scanned
        assert row["expected_reconciliation_state"] == "absent_with_script"

    def test_an_unresolved_identifier_is_blank_never_invented(
        self, population: Population, tmp_path: Path
    ):
        import csv

        _, reconciliation = write_ground_truth(population, tmp_path)
        with reconciliation.open(encoding="utf-8") as handle:
            rows = {row["candidate_uid"]: row for row in csv.DictReader(handle)}
        blank = one_with(population, ConflictKind.BLANK_CANDIDATE_ID)
        assert rows[blank.candidate_uid]["omr_roll"] == ""

    def test_stray_sheets_appear_and_carry_no_workbook_state(
        self, population: Population, tmp_path: Path
    ):
        import csv

        _, reconciliation = write_ground_truth(population, tmp_path)
        with reconciliation.open(encoding="utf-8") as handle:
            rows = {row["candidate_uid"]: row for row in csv.DictReader(handle)}
        for sheet in population.stray_sheets:
            row = rows[sheet.candidate_uid]
            assert row["attendance_status"] == ""
            assert row["attendance_record_present"] == ""

    def test_every_entry_has_an_expected_state(self, population: Population):
        states = expected_states(population)
        assert len(states) == len(population.candidates) + len(
            population.stray_sheets
        )
        assert all(isinstance(state, ReconciliationStatus) for state in states.values())


# ----------------------------------------------------------------------
# G - rates and profiles
# ----------------------------------------------------------------------
class TestGRatesAndProfiles:
    def test_the_none_profile_stages_no_conflicts_at_all(self):
        """The canonical dataset: clean paperwork, for a baseline run."""
        population = plan_population(
            count=50,
            set_codes=SETS,
            seed=SEED,
            rates=ConflictProfile.NONE.rates(),
            include_edge_cases=False,
        )
        assert population.counts()[ConflictKind.NONE.value] == 50
        assert population.stray_sheets == ()

    def test_edge_cases_are_staged_even_when_the_rates_are_zero(self):
        """The guarantee that makes a small dataset a complete test."""
        population = plan_population(
            count=50,
            set_codes=SETS,
            seed=SEED,
            rates=ConflictProfile.NONE.rates(),
            include_edge_cases=True,
        )
        staged = {candidate.conflict for candidate in population.candidates}
        assert staged == set(ConflictKind)

    def test_a_higher_profile_stages_more_conflicts(self):
        low = plan_population(
            count=1000, set_codes=SETS, seed=SEED, rates=ConflictProfile.LOW.rates()
        )
        high = plan_population(
            count=1000, set_codes=SETS, seed=SEED, rates=ConflictProfile.HIGH.rates()
        )
        assert low.counts()[ConflictKind.NONE.value] > (
            high.counts()[ConflictKind.NONE.value]
        )

    def test_impossible_rates_are_refused(self):
        with pytest.raises(ValueError, match="leaves no candidates"):
            ConflictRates(true_absentee=0.9, marked_absent_but_present=0.5).validate()

    def test_a_negative_rate_is_refused(self):
        with pytest.raises(ValueError, match="must not be negative"):
            ConflictRates(true_absentee=-0.1).validate()

    def test_a_tiny_population_still_gets_coverage_without_overflowing(self):
        """Coverage survives; only the proportions give way."""
        population = plan_population(count=12, set_codes=SETS, seed=SEED)
        assert len(population.candidates) == 12
        staged = {candidate.conflict for candidate in population.candidates}
        assert ConflictKind.DUPLICATE_SCRIPT in staged

    def test_the_summary_reports_what_was_staged_not_what_was_asked(
        self, population: Population
    ):
        report = summarise(population)
        assert report["registered_candidates"] == 400
        assert report["true_present"] + report["true_absent"] == 400
        assert sum(report["expected_conflicts"].values()) == 400
        assert report["seed"] == SEED
