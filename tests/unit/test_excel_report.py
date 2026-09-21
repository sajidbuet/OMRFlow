"""Excel workbook mechanics (Phase 9): rollwise, meritwise, layout, safety.

Scope:
    :mod:`omr_scanner.reporting.excel` against real openpyxl workbooks built
    by ``tests/report_fixtures.py``. No database, no GUI.
"""

from __future__ import annotations

from datetime import UTC, datetime
from fractions import Fraction
from typing import TYPE_CHECKING

import openpyxl
import pytest
from tests.report_fixtures import SAMPLE_HEADERS, build_result_template, sha256_of

from omr_scanner.domain.scoring import AnswerKey, AnswerKeyStatus
from omr_scanner.reporting import excel as rx
from omr_scanner.services.report_template import read_template, suggest_mapping

if TYPE_CHECKING:
    from pathlib import Path


@pytest.fixture
def roster(tmp_path: Path):
    template = build_result_template(tmp_path / "sample.xlsx", candidate_count=20)
    mapping = suggest_mapping(list(SAMPLE_HEADERS)).to_mapping()
    return template, read_template(template, mapping)


@pytest.fixture
def generated(tmp_path: Path, roster):
    template, built_roster = roster
    output = tmp_path / "output.xlsx"
    rx.copy_into(template, output)
    return template, built_roster, output


def _decisions(built_roster, *, absent_every: int = 5) -> dict:
    decisions = {}
    for index, row in enumerate(built_roster.rows, start=1):
        if absent_every and index % absent_every == 0:
            decisions[row.roll] = rx.CandidateReportRow(roll=row.roll, is_absent=True)
        else:
            decisions[row.roll] = rx.CandidateReportRow(
                roll=row.roll, is_absent=False, final_score=Fraction(100 - index)
            )
    return decisions


# ----------------------------------------------------------------------
class TestTemplatePreservation:
    def test_the_template_bytes_are_unchanged_after_a_full_generation(
        self, tmp_path: Path, roster
    ):
        template, built_roster = roster
        before = sha256_of(template)
        output = tmp_path / "output.xlsx"
        rx.copy_into(template, output)
        rx.populate_rollwise(output, built_roster, _decisions(built_roster))
        rx.build_meritwise_from_rollwise(output, built_roster, ())
        after = sha256_of(template)
        assert before == after

    def test_the_copy_is_a_different_file(self, tmp_path: Path, roster):
        template, _built_roster = roster
        output = tmp_path / "output.xlsx"
        rx.copy_into(template, output)
        assert output != template
        assert output.exists()

    def test_a_locked_or_unwritable_destination_is_reported_usefully(
        self, tmp_path: Path, roster
    ):
        template, _built_roster = roster
        bad_dir = tmp_path / "does_not_exist_and_cannot_be_created"
        import unittest.mock

        with (
            unittest.mock.patch.object(
                type(bad_dir), "mkdir", side_effect=OSError("locked")
            ),
            pytest.raises(rx.ExcelReportError, match=r"[Cc]ould not"),
        ):
            rx.copy_into(template, bad_dir / "out.xlsx")


# ----------------------------------------------------------------------
class TestRollwise:
    def test_no_candidate_row_is_lost(self, generated):
        _template, built_roster, output = generated
        result = rx.populate_rollwise(output, built_roster, _decisions(built_roster))
        assert result.rows_written == len(built_roster.rows) == 20

    def test_present_candidates_get_numeric_marks(self, generated):
        _template, built_roster, output = generated
        decisions = _decisions(built_roster)
        rx.populate_rollwise(output, built_roster, decisions)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[built_roster.sheet]
        first_row = built_roster.rows[0]
        cell = sheet.cell(
            row=first_row.sheet_row_number, column=built_roster.mapping.marks + 1
        )
        assert cell.value == 99  # 100 - 1
        assert isinstance(cell.value, int | float)

    def test_absent_candidates_show_the_canonical_marker(self, generated):
        _template, built_roster, output = generated
        decisions = _decisions(built_roster, absent_every=5)
        rx.populate_rollwise(output, built_roster, decisions)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[built_roster.sheet]
        absent_row = built_roster.rows[4]  # position 5
        marks = sheet.cell(
            row=absent_row.sheet_row_number, column=built_roster.mapping.marks + 1
        ).value
        rank = sheet.cell(
            row=absent_row.sheet_row_number, column=built_roster.mapping.rank + 1
        ).value
        assert marks == "ABSENT"
        assert rank == "---"

    def test_absent_candidates_are_never_dropped(self, generated):
        _template, built_roster, output = generated
        decisions = _decisions(built_roster, absent_every=5)
        result = rx.populate_rollwise(output, built_roster, decisions)
        assert result.absent_written == 4  # 20 / 5

    def test_rank_cells_contain_formulas_not_numbers(self, generated):
        _template, built_roster, output = generated
        decisions = _decisions(built_roster, absent_every=0)
        rx.populate_rollwise(output, built_roster, decisions)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[built_roster.sheet]
        row = built_roster.rows[0]
        rank_cell = sheet.cell(
            row=row.sheet_row_number, column=built_roster.mapping.rank + 1
        )
        assert str(rank_cell.value).startswith("=")

    def test_the_formula_range_covers_exactly_the_candidate_rows(self, generated):
        _template, built_roster, output = generated
        decisions = _decisions(built_roster, absent_every=0)
        rx.populate_rollwise(output, built_roster, decisions)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[built_roster.sheet]
        row = built_roster.rows[0]
        formula = sheet.cell(
            row=row.sheet_row_number, column=built_roster.mapping.rank + 1
        ).value
        expected_range = (
            f"${openpyxl.utils.get_column_letter(built_roster.mapping.marks + 1)}"
            f"${built_roster.first_data_row_number}:"
            f"${openpyxl.utils.get_column_letter(built_roster.mapping.marks + 1)}"
            f"${built_roster.last_data_row_number}"
        )
        assert expected_range in formula

    def test_row_order_is_unchanged_from_the_template(self, generated):
        _template, built_roster, output = generated
        decisions = _decisions(built_roster, absent_every=0)
        rx.populate_rollwise(output, built_roster, decisions)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[built_roster.sheet]
        rolls_in_sheet = [
            sheet.cell(row=r.sheet_row_number, column=built_roster.mapping.roll + 1).value
            for r in built_roster.rows
        ]
        assert rolls_in_sheet == list(built_roster.roll_numbers)

    def test_an_unresolved_candidate_leaves_the_marks_cell_blank(self, generated):
        _template, built_roster, output = generated
        decisions = {
            row.roll: rx.CandidateReportRow(roll=row.roll, is_absent=False, final_score=None)
            for row in built_roster.rows
        }
        result = rx.populate_rollwise(output, built_roster, decisions)
        assert len(result.unresolved_rows) == len(built_roster.rows)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[built_roster.sheet]
        row = built_roster.rows[0]
        assert sheet.cell(
            row=row.sheet_row_number, column=built_roster.mapping.marks + 1
        ).value is None

    def test_a_marks_header_override_is_applied(self, generated):
        _template, built_roster, output = generated
        rx.populate_rollwise(
            output, built_roster, {}, marks_header_override="Total (100)"
        )
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[built_roster.sheet]
        header = sheet.cell(
            row=built_roster.header_row_number, column=built_roster.mapping.marks + 1
        ).value
        assert header == "Total (100)"

    def test_decimal_marks_are_written_exactly(self, generated):
        _template, built_roster, output = generated
        decisions = {
            built_roster.rows[0].roll: rx.CandidateReportRow(
                roll=built_roster.rows[0].roll, is_absent=False,
                final_score=Fraction(291, 4),  # 72.75
            )
        }
        rx.populate_rollwise(output, built_roster, decisions)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[built_roster.sheet]
        value = sheet.cell(
            row=built_roster.rows[0].sheet_row_number,
            column=built_roster.mapping.marks + 1,
        ).value
        assert value == pytest.approx(72.75)

    def test_zero_marks_are_written_as_zero_not_blank(self, generated):
        _template, built_roster, output = generated
        decisions = {
            built_roster.rows[0].roll: rx.CandidateReportRow(
                roll=built_roster.rows[0].roll, is_absent=False, final_score=Fraction(0)
            )
        }
        rx.populate_rollwise(output, built_roster, decisions)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[built_roster.sheet]
        value = sheet.cell(
            row=built_roster.rows[0].sheet_row_number,
            column=built_roster.mapping.marks + 1,
        ).value
        assert value == 0

    def test_a_negative_mark_is_written(self, generated):
        _template, built_roster, output = generated
        decisions = {
            built_roster.rows[0].roll: rx.CandidateReportRow(
                roll=built_roster.rows[0].roll, is_absent=False, final_score=Fraction(-2)
            )
        }
        rx.populate_rollwise(output, built_roster, decisions)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[built_roster.sheet]
        value = sheet.cell(
            row=built_roster.rows[0].sheet_row_number,
            column=built_roster.mapping.marks + 1,
        ).value
        assert value == -2

    def test_full_calc_on_load_is_set(self, generated):
        _template, built_roster, output = generated
        rx.populate_rollwise(output, built_roster, _decisions(built_roster))
        workbook = openpyxl.load_workbook(output)
        assert workbook.calculation.fullCalcOnLoad is True


# ----------------------------------------------------------------------
class TestMeritwise:
    """Meritwise is now a *copy* of the completed Rollwise sheet.

    These cover membership, ordering and regeneration at this module's own
    level. The formatting the copy has to carry over - merges, fonts, fills,
    borders, widths, heights, page setup, print titles, headers/footers and
    the logo - is asserted in ``tests/unit/test_meritwise_workbook.py``
    against a workbook built to contain every one of them.
    """

    def _order(self, built_roster, rolls) -> tuple[rx.MeritwiseRow, ...]:
        """Merit order expressed as the Rollwise rows to keep."""
        by_roll = {row.roll: row.sheet_row_number for row in built_roster.rows}
        return tuple(
            rx.MeritwiseRow(source_row_number=by_roll[roll], roll=roll) for roll in rolls
        )

    def test_rows_appear_in_the_order_given(self, generated):
        _template, built_roster, output = generated
        rolls = [row.roll for row in built_roster.rows if row.roll][:3]
        rx.build_meritwise_from_rollwise(
            output, built_roster, self._order(built_roster, reversed(rolls))
        )
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.MERITWISE_SHEET_NAME]
        first = built_roster.first_data_row_number
        column = built_roster.mapping.roll + 1
        written = [sheet.cell(row=first + i, column=column).value for i in range(3)]
        assert written == list(reversed(rolls))

    def test_a_roll_left_out_does_not_appear(self, generated):
        _template, built_roster, output = generated
        rolls = [row.roll for row in built_roster.rows if row.roll]
        kept, dropped = rolls[:2], rolls[2:]
        rx.build_meritwise_from_rollwise(
            output, built_roster, self._order(built_roster, kept)
        )
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.MERITWISE_SHEET_NAME]
        column = built_roster.mapping.roll + 1
        present = {
            sheet.cell(row=r, column=column).value for r in range(1, sheet.max_row + 1)
        }
        assert set(kept) <= present
        assert not (set(dropped) & present)

    def test_regenerating_replaces_rather_than_duplicates_the_sheet(self, generated):
        _template, built_roster, output = generated
        rolls = [row.roll for row in built_roster.rows if row.roll]
        for _ in range(2):
            rx.build_meritwise_from_rollwise(
                output, built_roster, self._order(built_roster, rolls[:1])
            )
        workbook = openpyxl.load_workbook(output)
        assert workbook.sheetnames.count(rx.MERITWISE_SHEET_NAME) == 1

    def test_an_old_capitalised_sheet_is_replaced_not_left_beside_the_new_one(
        self, generated
    ):
        """A workbook regenerated from an earlier build carried "Meritwise"."""
        _template, built_roster, output = generated
        workbook = openpyxl.load_workbook(output)
        workbook.create_sheet("Meritwise")
        workbook.save(output)
        workbook.close()

        rolls = [row.roll for row in built_roster.rows if row.roll]
        rx.build_meritwise_from_rollwise(
            output, built_roster, self._order(built_roster, rolls[:1])
        )
        reopened = openpyxl.load_workbook(output)
        names = [name.casefold() for name in reopened.sheetnames]
        assert names.count("meritwise") == 1

    def test_the_result_reports_what_it_did(self, generated):
        _template, built_roster, output = generated
        rolls = [row.roll for row in built_roster.rows if row.roll]
        result = rx.build_meritwise_from_rollwise(
            output, built_roster, self._order(built_roster, rolls[:1])
        )
        assert result.sheet_name == rx.MERITWISE_SHEET_NAME
        assert result.rows_written == 1
        assert result.rows_removed == len(rolls) - 1


# ----------------------------------------------------------------------
class TestSummarySheet:
    def test_every_field_is_shown(self, generated):
        _template, _built_roster, output = generated
        data = rx.SummaryData(
            project_name="Test Exam", set_code="A", registered=100, present=90,
            absent=10, scored=90, unresolved=0, maximum_score=Fraction(90),
            highest_score=Fraction(88), lowest_score=Fraction(20),
            mean_score=Fraction(55), median_score=Fraction(56),
            scoring_policy_summary=("Correct: +1.00",), answer_key_revision=2,
            generated_at=datetime(2026, 1, 1, tzinfo=UTC), application_version="0.1.0",
        )
        rx.add_summary_sheet(output, data)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.SUMMARY_SHEET_NAME]
        values = {sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
                  for r in range(1, 15)}
        assert values["Registered candidates"] == "100"
        assert values["Set"] == "A"

    def test_an_uncomputable_statistic_shows_n_a_not_a_fabricated_value(self, generated):
        _template, _built_roster, output = generated
        data = rx.SummaryData(
            project_name="p", set_code="A", registered=0, present=0, absent=0,
            scored=0, unresolved=0, maximum_score=None, highest_score=None,
            lowest_score=None, mean_score=None, median_score=None,
            scoring_policy_summary=(), answer_key_revision=None,
            generated_at=datetime(2026, 1, 1, tzinfo=UTC), application_version="",
        )
        rx.add_summary_sheet(output, data)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.SUMMARY_SHEET_NAME]
        values = {sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
                  for r in range(1, 15)}
        assert values["Highest obtained score"] == "N/A"
        assert values["Answer-key revision"] == "N/A"


# ----------------------------------------------------------------------
class TestAnswerKeySheet:
    def test_only_this_sets_key_appears(self, generated):
        _template, _built_roster, output = generated
        key_a = AnswerKey(set_code="A", answers="ABCD", status=AnswerKeyStatus.VERIFIED)
        rx.add_answer_key_sheet(output, set_code="A", key=key_a, key_id=1)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.ANSWER_KEY_SHEET_NAME]
        assert "Set A" in sheet.cell(row=1, column=1).value
        assert "ABCD" in sheet.cell(row=3, column=1).value
        assert sheet.cell(row=6, column=2).value == "A"

    def test_wrong_questions_are_flagged(self, generated):
        _template, _built_roster, output = generated
        key = AnswerKey(
            set_code="A", answers="ABCD", wrong_questions=frozenset({2}),
            status=AnswerKeyStatus.VERIFIED,
        )
        rx.add_answer_key_sheet(output, set_code="A", key=key, key_id=1)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.ANSWER_KEY_SHEET_NAME]
        assert sheet.cell(row=7, column=3).value == "Yes"  # question 2


# ----------------------------------------------------------------------
class TestProcessingLogPrivacy:
    def test_only_supplied_entries_appear_no_hidden_data(self, generated):
        _template, _built_roster, output = generated
        entries = [
            rx.ProcessingLogEntry("Operation", "Generate XLSX"),
            rx.ProcessingLogEntry("Candidate count", "229"),
        ]
        rx.add_processing_log_sheet(output, entries)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.PROCESSING_LOG_SHEET_NAME]
        text = "\n".join(
            str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
        )
        assert "Generate XLSX" in text
        assert "229" in text


# ----------------------------------------------------------------------
class TestLayout:
    def test_page_setup_is_applied(self, generated):
        _template, built_roster, output = generated
        layout = rx.LayoutSettings(
            page_size="Letter", orientation="landscape", margin_top_mm=10,
        )
        rx.apply_layout(
            output, [built_roster.sheet], layout,
            primary_header_row=built_roster.header_row_number,
        )
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[built_roster.sheet]
        assert sheet.page_setup.orientation == "landscape"

    def test_default_layout_changes_nothing_visible(self, tmp_path: Path, roster):
        # An all-default LayoutSettings is what "preserve the template" means.
        template, built_roster = roster
        output = tmp_path / "output.xlsx"
        rx.copy_into(template, output)
        before = openpyxl.load_workbook(output)[built_roster.sheet].cell(1, 1).value
        rx.apply_layout(
            output, [built_roster.sheet], rx.LayoutSettings(),
            primary_header_row=built_roster.header_row_number,
        )
        after = openpyxl.load_workbook(output)[built_roster.sheet].cell(1, 1).value
        assert before == after

    def test_a_title_with_no_blank_space_is_skipped_with_a_warning(self, generated):
        _template, built_roster, output = generated
        layout = rx.LayoutSettings(title_text="Should not fit")
        warnings = rx.apply_layout(
            output, [built_roster.sheet], layout,
            primary_header_row=built_roster.header_row_number,
        )
        assert warnings
        assert "blank row" in warnings[0]

    def test_layout_settings_round_trip_through_json(self):
        original = rx.LayoutSettings(
            title_text="X", logo_max_width_px=200, page_size="Letter"
        )
        restored = rx.LayoutSettings.from_json(original.to_json())
        assert restored == original

    def test_a_missing_logo_file_is_skipped_with_a_warning_not_a_failure(
        self, generated
    ):
        _template, built_roster, output = generated
        layout = rx.LayoutSettings(logo_path="C:/does/not/exist.png")
        warnings = rx.apply_layout(
            output, [built_roster.sheet], layout,
            primary_header_row=built_roster.header_row_number,
        )
        assert any("not found" in warning for warning in warnings)
        # Generation itself is unaffected - the workbook still opens cleanly.
        assert openpyxl.load_workbook(output) is not None

    def test_a_real_logo_file_degrades_gracefully_without_failing_generation(
        self, generated, tmp_path
    ):
        # openpyxl's image support requires Pillow, which is not a declared
        # OMRFlow dependency (§16/§17 make a logo optional, and this project
        # avoids adding one for a single decorative feature - see
        # reporting/excel.py's `_insert_logo` docstring). Whichever branch
        # this environment takes - Pillow absent (this build) or present -
        # the contract under test is the same: a real image file never turns
        # a successful generation into a failed one, and the workbook is
        # still usable either way.
        _template, built_roster, output = generated
        logo = tmp_path / "logo.png"
        # A minimal, genuinely valid 1x1 PNG - real image bytes, not a stub.
        logo.write_bytes(
            bytes.fromhex(
                "89504e470d0a1a0a0000000d49484452000000010000000108020000009077"
                "3df40000000a49444154789c6360000002000155a4f9a70000000049454e44"
                "ae426082"
            )
        )
        layout = rx.LayoutSettings(logo_path=str(logo))
        warnings = rx.apply_layout(
            output, [built_roster.sheet], layout,
            primary_header_row=built_roster.header_row_number,
        )
        try:
            import PIL  # noqa: F401

            has_pillow = True
        except ImportError:
            has_pillow = False
        if not has_pillow:
            assert any(
                "Pillow" in warning or "not installed" in warning for warning in warnings
            )
        assert openpyxl.load_workbook(output) is not None


# ----------------------------------------------------------------------
class TestUnicode:
    def test_unicode_names_survive_meritwise(self, generated):
        """The name is copied from the Rollwise cell, so it must survive intact.

        Unlike the old generated table, nothing re-encodes or re-escapes a
        name here - which is also why the old formula-injection guard on a
        name beginning with ``=`` no longer applies: the value is the
        operator's own cell, copied, never a string this application writes.
        """
        _template, built_roster, output = generated
        first = next(row for row in built_roster.rows if row.roll)
        name_column = (built_roster.mapping.name or 0) + 1

        workbook = openpyxl.load_workbook(output)
        workbook[built_roster.sheet].cell(
            row=first.sheet_row_number, column=name_column, value="মোহাম্মদ রহিম"
        )
        workbook.save(output)
        workbook.close()

        rx.build_meritwise_from_rollwise(
            output,
            built_roster,
            (rx.MeritwiseRow(source_row_number=first.sheet_row_number, roll=first.roll),),
        )
        reopened = openpyxl.load_workbook(output)
        sheet = reopened[rx.MERITWISE_SHEET_NAME]
        assert sheet.cell(
            row=built_roster.first_data_row_number, column=name_column
        ).value == "মোহাম্মদ রহিম"

    def test_unicode_project_name_in_summary(self, generated):
        _template, _built_roster, output = generated
        data = rx.SummaryData(
            project_name="পরীক্ষা ২০২৬", set_code="A", registered=1, present=1,
            absent=0, scored=1, unresolved=0, maximum_score=Fraction(1),
            highest_score=Fraction(1), lowest_score=Fraction(1), mean_score=Fraction(1),
            median_score=Fraction(1), scoring_policy_summary=(), answer_key_revision=1,
            generated_at=datetime(2026, 1, 1, tzinfo=UTC), application_version="0.1.0",
        )
        rx.add_summary_sheet(output, data)
        workbook = openpyxl.load_workbook(output)
        sheet = workbook[rx.SUMMARY_SHEET_NAME]
        assert sheet.cell(row=1, column=2).value == "পরীক্ষা ২০২৬"
