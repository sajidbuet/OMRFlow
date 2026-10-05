"""Building a campaign's examination project through production services (revised phase 9).

One project per run (interrupted, control, finite), each built the way an
office builds one: a project, its template, one defined set per paper with
its own attendance workbook (also the result template), a verified answer key
per set, a scoring policy, an open scan session and - for the continuous runs
- one watched source per simulated scanner, attached to the session. Nothing
is written to the database except through these services.
"""

from __future__ import annotations

import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

from omr_scanner.evaluation.intake_qualification.config import OPERATOR, StabilitySettings

if TYPE_CHECKING:  # pragma: no cover - typing only
    from omr_scanner.evaluation.intake_qualification.cohort import CampaignPlan

SESSION_NAME = "Qualification examination"


@dataclass(frozen=True, slots=True)
class ProjectFacts:
    """Identifiers the run needs afterwards."""

    root: Path
    scan_session_id: str
    template_path: Path
    set_ids: dict[str, str]
    source_ids: dict[str, str]


def write_attendance_workbook(
    plan: CampaignPlan, set_code: str, path: Path, *, absent: dict[str, bool] | None = None
) -> Path:
    """One set's attendance workbook (and result template), as an office keeps it.

    ``Sl.No. | Roll No. | Name | Total | Merit`` with the header on row 3; an
    absentee is marked ``ABSENT`` in the Total column - the layout the
    existing golden regression already uses. ``absent`` overrides the planned
    marks for named rolls (a corrected list).
    """
    corrections = absent or {}
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = f"Set {set_code}"
    for column, header in enumerate(("Sl.No.", "Roll No.", "Name", "Total", "Merit"), 1):
        sheet.cell(row=3, column=column, value=header)
    rows = sorted(
        (item for item in plan.candidates if item.set_code == set_code and item.on_roster),
        key=lambda item: item.row,
    )
    for offset, candidate in enumerate(rows):
        row = 4 + offset
        sheet.cell(row=row, column=1, value=offset + 1)
        sheet.cell(row=row, column=2, value=candidate.roll)
        sheet.cell(row=row, column=3, value=candidate.name)
        if corrections.get(candidate.roll, candidate.absent_on_roster):
            sheet.cell(row=row, column=4, value="ABSENT")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def create_campaign_project(
    workspace: Path,
    name: str,
    plan: CampaignPlan,
    template_path: Path,
    inputs: Path,
    *,
    source_roots: dict[str, Path] | None = None,
    stability: StabilitySettings | None = None,
) -> ProjectFacts:
    """Create and configure one campaign project; return its identifiers.

    Args:
        workspace: Parent folder; the project folder is created inside.
        name: Project name.
        plan: The campaign plan (sets, rosters, keys, scoring).
        template_path: The qualification template file.
        inputs: Where the attendance workbooks are written (kept as evidence).
        source_roots: ``{source label: folder}`` - watched sources to create
            and attach (``None``: a finite project, no source).
        stability: The sources' stability policy.
    """
    from omr_scanner.domain.intake import StabilityPolicy
    from omr_scanner.domain.scoring import NegativeMarking, ScoringPolicy
    from omr_scanner.services import (
        candidate_import,
        create_project,
        project_sets,
        scan_sessions,
        scoring_store,
        set_active_template,
        set_attendance,
    )
    from omr_scanner.services import intake as intake_service
    from omr_scanner.services.answer_key import plan_for, read_key
    from omr_scanner.services.template_service import load_template

    template = load_template(template_path)
    session = create_project(workspace, name)
    database = session.database
    try:
        target = session.project.layout.templates_dir / template_path.name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(template_path, target)
        set_active_template(session, target)

        question_plan = plan_for(template)
        set_ids: dict[str, str] = {}
        for code in plan.config.sets:
            set_ids[code] = project_sets.add_set(database, code, f"Paper {code}").set_id
            workbook = write_attendance_workbook(plan, code, inputs / f"attendance_set_{code}.xlsx")
            set_attendance.assign_attendance_workbook(
                database, set_ids[code], workbook, candidate_import.read_roster(workbook),
                imported_by=OPERATOR,
            )
            draft = read_key(
                plan.keys[code], question_plan, code,
                wrong_questions=plan.scoring.wrong_questions.get(code, ()),
            )
            stored = scoring_store.save_key(database, draft.to_key(), created_by=OPERATOR)
            scoring_store.verify_key(database, stored.key_id, verified_by=OPERATOR)
        rules = plan.scoring
        scoring_store.save_policy(
            database,
            ScoringPolicy(
                correct_mark=rules.correct,
                blank_mark=rules.blank,
                incorrect_penalty=rules.incorrect_penalty,
                multiple_penalty=(
                    None if rules.multiple_penalty == rules.incorrect_penalty
                    else rules.multiple_penalty
                ),
                mode=NegativeMarking.FIXED,
                clamp_minimum=True,
                minimum_score=rules.minimum,
            ),
            created_by=OPERATOR,
        )
        info = scan_sessions.create_scan_session(
            database, name=SESSION_NAME, created_by=OPERATOR
        )
        source_ids: dict[str, str] = {}
        settings = stability or StabilitySettings()
        policy = StabilityPolicy(
            min_observations=settings.min_observations,
            quiet_seconds=settings.quiet_seconds,
            max_decode_attempts=settings.max_decode_attempts,
            retry_backoff_seconds=settings.retry_backoff_seconds,
            poll_interval_seconds=settings.poll_interval_seconds,
        )
        for label, root in (source_roots or {}).items():
            created = intake_service.create_source(
                database, label=f"Scanner {label}", root_path=str(root), policy=policy,
                created_by=OPERATOR,
            )
            intake_service.attach_source(
                database, created.source_id, info.scan_session_id, actor=OPERATOR
            )
            source_ids[label] = created.source_id
        return ProjectFacts(
            root=session.root,
            scan_session_id=info.scan_session_id,
            template_path=target,
            set_ids=set_ids,
            source_ids=source_ids,
        )
    finally:
        session.close()



__all__ = ["SESSION_NAME", "ProjectFacts", "create_campaign_project", "write_attendance_workbook"]
