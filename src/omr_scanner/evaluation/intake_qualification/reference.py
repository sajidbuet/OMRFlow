"""The independent reference: what the examination's results must be (revised phase 9).

Purpose:
    Compute, from the campaign plan alone, every value the campaign compares
    OMRFlow's output against: the effective population, each sheet's effective
    Student ID, the attendance list after the office's corrections, every
    candidate's mark (positive, negative, blank, multiple and full-credit
    rules), the rank order, and every compared cell of each set's result
    workbook - plus, at a checkpoint, which review conflicts must be open for
    the sheets committed so far and the decisions committed so far.

Independence:
    Nothing here calls an OMRFlow scoring, reconciliation, population or
    report function. The scorer is a deliberately plain loop over the plan's
    marks and keys; the report expectation is built from the rule the user
    documentation states (Rollwise in list order, Meritwise by descending mark
    then ascending roll, ``RANK.EQ`` merit, ``ABSENT`` / ``---`` for an
    absentee). Where OMRFlow's behaviour is the thing under test, this module
    says what it should be; it never asks OMRFlow.
"""

from __future__ import annotations

import re
from collections import defaultdict
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from fractions import Fraction
from typing import Any

from omr_scanner.evaluation.intake_qualification.cohort import (
    CampaignPlan,
    Category,
    Content,
    ContentKind,
    IdDefect,
    TaskKind,
)

PHASES = ("first_close", "reopen")
ABSENT_MARK = "ABSENT"
ABSENT_RANK = "---"


# ----------------------------------------------------------------------
# Marks
# ----------------------------------------------------------------------
@dataclass(frozen=True, slots=True)
class ExpectedMark:
    """One candidate's mark, computed by the reference scorer."""

    raw: Fraction
    final: Fraction
    correct: int
    incorrect: int
    blank: int
    multiple: int
    wrong_question: int


def reference_score(plan: CampaignPlan, set_code: str, answers: tuple[str, ...]) -> ExpectedMark:
    """Mark ``answers`` against ``set_code``'s key under the plan's rules - by hand.

    Precedence, as the scoring documentation states it: a question withdrawn
    for the set gives full credit; a blank gives the blank mark; two marks give
    the multiple deduction; the key's answer gives the correct mark; anything
    else the incorrect deduction. The total is clamped once, at the end.
    """
    rules = plan.scoring
    key = plan.keys[set_code]
    withdrawn = set(rules.wrong_questions.get(set_code, ()))
    total = Fraction(0)
    correct = incorrect = blank = multiple = wrong = 0
    for number, (given, right) in enumerate(zip(answers, key, strict=True), start=1):
        if number in withdrawn:
            total += rules.correct
            wrong += 1
        elif given == "":
            total += rules.blank
            blank += 1
        elif "-" in given:
            total -= rules.multiple_penalty
            multiple += 1
        elif given == right:
            total += rules.correct
            correct += 1
        else:
            total -= rules.incorrect_penalty
            incorrect += 1
    final = max(total, rules.minimum)
    return ExpectedMark(total, final, correct, incorrect, blank, multiple, wrong)


# ----------------------------------------------------------------------
# The population and identities
# ----------------------------------------------------------------------
def tasks_of(plan: CampaignPlan, kind: TaskKind) -> list[Any]:
    return [task for task in plan.tasks if task.kind is kind]


def corrected_contents(plan: CampaignPlan) -> dict[str, str]:
    """Content -> the Student ID the operator will enter (the written one)."""
    return {task.content: task.value for task in tasks_of(plan, TaskKind.CORRECT_ID)}


def registered_contents(plan: CampaignPlan, *, phase: str = "reopen") -> set[str]:
    """Contents that become sheets.

    Every main-timeline content, plus the late file after a reopen.
    """
    keys = {item.content for item in plan.main_arrivals}
    if phase == "reopen":
        keys |= {item.content for item in plan.late_arrivals}
    return keys


def superseded_contents(plan: CampaignPlan) -> set[str]:
    """Rejected sheets whose rescan was confirmed: folded originals and a chain's first rescan."""
    out = {task.content for task in tasks_of(plan, TaskKind.CONFIRM_REPLACEMENT)}
    return out


def effective_contents(plan: CampaignPlan, *, phase: str = "reopen") -> set[str]:
    """The expected final effective population, by content."""
    return registered_contents(plan, phase=phase) - superseded_contents(plan)


def expected_identifier(content: Content, *, corrected: bool) -> str:
    """The sheet's effective Student ID: the written one once corrected, else the bubbles'."""
    if corrected:
        return content.written_roll
    return content.bubbled_roll


def reliable(value: str) -> bool:
    """An identifier a duplicate check may use (no blank or unresolved position)."""
    return bool(value) and "?" not in value and "_" not in value


# ----------------------------------------------------------------------
# Attendance and results
# ----------------------------------------------------------------------
def attendance_list(plan: CampaignPlan, set_code: str, phase: str) -> dict[str, bool]:
    """``{roll: absent}`` - the office's list for ``set_code`` after ``phase``'s corrections."""
    upto = PHASES[: PHASES.index(phase) + 1]
    absent = {
        item.roll: item.absent_on_roster
        for item in plan.candidates
        if item.set_code == set_code and item.on_roster
    }
    for task in plan.tasks:
        if task.set_code == set_code and task.phase in upto:
            if task.kind is TaskKind.LIST_ABSENT:
                absent[task.candidate] = True
            elif task.kind is TaskKind.LIST_PRESENT:
                absent[task.candidate] = False
    return absent


@dataclass(frozen=True, slots=True)
class ExpectedResult:
    """One roster candidate's expected outcome in one set."""

    roll: str
    name: str
    serial: int
    absent: bool
    mark: ExpectedMark | None
    rank: int | None


def expected_results(plan: CampaignPlan, set_code: str, phase: str) -> list[ExpectedResult]:
    """Every roster candidate of ``set_code`` after ``phase``.

    In list order, with marks and ranks.
    """
    listed = attendance_list(plan, set_code, phase)
    rows = sorted(
        (item for item in plan.candidates if item.set_code == set_code and item.on_roster),
        key=lambda item: item.row,
    )
    marks: dict[str, ExpectedMark] = {}
    for candidate in rows:
        if listed[candidate.roll]:
            continue
        marks[candidate.roll] = reference_score(plan, set_code, candidate.answers)
    finals = [mark.final for mark in marks.values()]
    results = []
    for serial, candidate in enumerate(rows, start=1):
        mark = marks.get(candidate.roll)
        rank = None if mark is None else 1 + sum(1 for other in finals if other > mark.final)
        results.append(ExpectedResult(
            roll=candidate.roll, name=candidate.name, serial=serial,
            absent=listed[candidate.roll], mark=mark, rank=rank,
        ))
    return results


def expected_statuses_before(plan: CampaignPlan, set_code: str, phase: str) -> dict[str, str]:
    """Reconciliation statuses against the list in use *before* ``phase``'s decisions.

    ``first_close``: the original lists, no decision yet. ``reopen``: the list
    the first close left in use, with that close's decisions on it (an
    accidental second scan stays set aside, an unknown script stays
    dismissed) - and the late script now arrived for a candidate that list
    calls absent.
    """
    late = {item.content for item in plan.late_arrivals}
    previous = None if phase == "first_close" else PHASES[PHASES.index(phase) - 1]
    if previous is not None:
        carried = expected_statuses_after(plan, set_code, previous)
        listed_before = attendance_list(plan, set_code, previous)
        for key in late:
            content = plan.content(key)
            if content.set_code == set_code and content.candidate in listed_before:
                carried[content.candidate] = (
                    "absent_with_script" if listed_before[content.candidate] else "matched"
                )
        return carried
    listed = {
        item.roll: item.absent_on_roster
        for item in plan.candidates if item.set_code == set_code and item.on_roster
    }
    scripts: dict[str, int] = defaultdict(int)
    registered = registered_contents(plan, phase=phase)
    for key in registered:
        content = plan.content(key)
        if content.kind is ContentKind.BLANK_PAGE or content.set_code != set_code:
            continue
        if key in superseded_contents(plan) or (phase == "first_close" and key in late):
            continue
        scripts[content.candidate] += 1
    statuses: dict[str, str] = {}
    for roll, absent in listed.items():
        count = scripts.get(roll, 0)
        if absent:
            statuses[roll] = "absent_with_script" if count else "absent_confirmed"
        elif count == 0:
            statuses[roll] = "present_without_script"
        elif count > 1:
            statuses[roll] = "duplicate_script"
        else:
            statuses[roll] = "matched"
    for roll in plan.unknown_rolls:
        candidate = plan.candidate(roll)
        if candidate is not None and candidate.set_code == set_code:
            statuses[roll] = "unknown_id"
    return statuses


def expected_statuses_after(plan: CampaignPlan, set_code: str, phase: str) -> dict[str, str]:
    """Statuses once the office's decisions of ``phase`` are in.

    Unknown ones are dismissed, not gone.
    """
    listed = attendance_list(plan, set_code, phase)
    statuses = {
        roll: ("absent_confirmed" if absent else "matched") for roll, absent in listed.items()
    }
    for roll in plan.unknown_rolls:
        candidate = plan.candidate(roll)
        if candidate is not None and candidate.set_code == set_code:
            statuses[roll] = "unknown_id"
    return statuses


# ----------------------------------------------------------------------
# The result workbook
# ----------------------------------------------------------------------
_RANK = re.compile(
    r"RANK\.EQ\(\s*\$?([A-Z]+)\$?(\d+)\s*,\s*\$?([A-Z]+)\$?(\d+)\s*:\s*\$?([A-Z]+)\$?(\d+)"
)


def evaluate_rank(formula: str, column_values: Mapping[int, Any]) -> int | None:
    """The value Excel computes for the merit formula OMRFlow writes.

    Only the ``RANK.EQ(cell, range, 0)`` the formula contains is evaluated
    (descending: 1 + the number of numeric marks in the range above this one);
    an ``ABSENT``/``ABS`` mark gives ``None`` (the formula's ``---``).
    """
    match = _RANK.search(formula or "")
    if match is None:
        return None
    row = int(match.group(2))
    first, last = int(match.group(4)), int(match.group(6))
    value = column_values.get(row)
    if isinstance(value, str) and value.strip().upper() in ("ABSENT", "ABS"):
        return None
    if not isinstance(value, int | float):
        return None
    numbers = [
        item for number, item in column_values.items()
        if first <= number <= last and isinstance(item, int | float)
    ]
    return 1 + sum(1 for item in numbers if item > value)


@dataclass(frozen=True, slots=True)
class CellMismatch:
    """One workbook cell whose value differs from the expectation."""

    sheet: str
    row: int
    column: str
    expected: Any
    actual: Any

    def describe(self) -> str:
        """The mismatch as one line: the cell, the expected and the actual value."""
        return (
            f"{self.sheet}!{self.column}{self.row}: expected {self.expected!r}, "
            f"got {self.actual!r}"
        )


def compare_workbook(
    plan: CampaignPlan, set_code: str, phase: str, path: str
) -> tuple[int, list[CellMismatch]]:
    """Compare a generated result workbook's values with the expectation.

    Compared: on the set's Rollwise sheet (``Set X``) and on ``meritwise``,
    every data row's Sl.No., Roll No., Name and Total (the mark - numerically
    exact to 1e-9, or ``ABSENT``) and the Merit cell (``---``, or the value its
    ``RANK.EQ`` formula evaluates to); the header row's labels; on the Answer
    Key sheet every question's correct answer and its withdrawn flag. The
    Summary and Processing Log sheets (generation timestamps) are not compared,
    as in the golden regression. Returns ``(cells compared, mismatches)``.
    """
    import openpyxl

    expected = expected_results(plan, set_code, phase)
    workbook = openpyxl.load_workbook(path)
    mismatches: list[CellMismatch] = []
    compared = 0
    try:
        rollwise = workbook[f"Set {set_code}"]
        merit = workbook["meritwise"]

        def check(sheet: str, row: int, column: str, want: Any, have: Any) -> None:
            nonlocal compared
            compared += 1
            same = (
                isinstance(want, Fraction) and isinstance(have, int | float)
                and abs(float(want) - float(have)) < 1e-9
            ) or (not isinstance(want, Fraction) and _text(want) == _text(have))
            if not same:
                mismatches.append(CellMismatch(sheet, row, column, want, have))

        for sheet in (rollwise, merit):
            headers = [sheet.cell(row=3, column=col).value for col in range(1, 4)]
            for col, want in zip("ABC", ("Sl.No.", "Roll No.", "Name"), strict=True):
                check(sheet.title, 3, col, want, headers["ABC".index(col)])

        totals = {row: rollwise.cell(row=row, column=4).value
                  for row in range(4, 4 + len(expected))}
        for offset, item in enumerate(expected):
            row = 4 + offset
            check(rollwise.title, row, "A", item.serial, rollwise.cell(row=row, column=1).value)
            check(rollwise.title, row, "B", item.roll, rollwise.cell(row=row, column=2).value)
            check(rollwise.title, row, "C", item.name, rollwise.cell(row=row, column=3).value)
            if item.absent:
                check(rollwise.title, row, "D", ABSENT_MARK, totals[row])
                check(rollwise.title, row, "E", ABSENT_RANK,
                      rollwise.cell(row=row, column=5).value)
            else:
                assert item.mark is not None
                check(rollwise.title, row, "D", item.mark.final, totals[row])
                check(rollwise.title, row, "E", item.rank,
                      evaluate_rank(str(rollwise.cell(row=row, column=5).value), totals))
        trailing = rollwise.cell(row=4 + len(expected), column=2).value
        check(rollwise.title, 4 + len(expected), "B", None, trailing)

        ordered = sorted(
            (item for item in expected if not item.absent and item.mark is not None),
            key=lambda item: (-item.mark.final if item.mark else 0, item.roll),
        )
        merit_totals = {row: merit.cell(row=row, column=4).value
                        for row in range(4, 4 + len(ordered))}
        for offset, item in enumerate(ordered):
            row = 4 + offset
            assert item.mark is not None
            check(merit.title, row, "A", offset + 1, merit.cell(row=row, column=1).value)
            check(merit.title, row, "B", item.roll, merit.cell(row=row, column=2).value)
            check(merit.title, row, "C", item.name, merit.cell(row=row, column=3).value)
            check(merit.title, row, "D", item.mark.final, merit_totals[row])
            check(merit.title, row, "E", item.rank,
                  evaluate_rank(str(merit.cell(row=row, column=5).value), merit_totals))
        check(merit.title, 4 + len(ordered), "B", None,
              merit.cell(row=4 + len(ordered), column=2).value)

        key_sheet = workbook["Answer Key"]
        withdrawn = set(plan.scoring.wrong_questions.get(set_code, ()))
        for number, answer in enumerate(plan.keys[set_code], start=1):
            row = 5 + number
            check(key_sheet.title, row, "A", number, key_sheet.cell(row=row, column=1).value)
            check(key_sheet.title, row, "B", answer, key_sheet.cell(row=row, column=2).value)
            flagged = key_sheet.cell(row=row, column=3).value not in (None, "")
            check(key_sheet.title, row, "C", number in withdrawn, flagged)
    finally:
        workbook.close()
    return compared, mismatches


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# ----------------------------------------------------------------------
# Conflicts at a checkpoint
# ----------------------------------------------------------------------
@dataclass(frozen=True, slots=True)
class CommittedFacts:
    """What the reference needs about the committed state - decisions, not OMRFlow's conclusions.

    Attributes:
        sheets: ``scan id -> content key`` for every sheet whose recognition
            is committed (completed / warning / failed) in a live batch.
        corrected: Scan ids whose Student ID the operator corrected
            (committed audit events).
        accepted_duplicate: Scan ids whose duplicate-ID conflict the operator
            accepted.
        acknowledged: Scan ids whose sheet conflict the operator acknowledged.
        rejected: Scan ids rejected (and not since un-rejected).
    """

    sheets: Mapping[int, str]
    corrected: frozenset[int]
    accepted_duplicate: frozenset[int]
    acknowledged: frozenset[int]
    rejected: frozenset[int]


def expected_open_conflicts(plan: CampaignPlan, facts: CommittedFacts) -> set[tuple[str, str]]:
    """``{(content key, conflict type)}`` that must be open for the committed population.

    Sheet-local conflicts come from the plan's expected reading of each
    committed, non-rejected sheet and close when the operator decided them.
    Duplicate-ID conflicts: committed, non-rejected sheets are grouped by
    their effective Student ID (the written ID once the operator corrected it,
    the bubbles' otherwise, reliable values only); every member of a group of
    two or more has one, open unless the operator corrected or accepted it.
    """
    expected: set[tuple[str, str]] = set()
    live = {scan: key for scan, key in facts.sheets.items() if scan not in facts.rejected}
    groups: dict[str, list[int]] = defaultdict(list)
    for scan, key in live.items():
        content = plan.content(key)
        corrected = scan in facts.corrected
        for kind in content.expected_conflicts:
            if kind == "registration_failed":
                if scan not in facts.acknowledged:
                    expected.add((key, kind))
            elif not corrected:
                expected.add((key, kind))
        identity = expected_identifier(content, corrected=corrected)
        if reliable(identity):
            groups[identity].append(scan)
    for members in groups.values():
        if len(members) < 2:
            continue
        for scan in members:
            if scan in facts.corrected or scan in facts.accepted_duplicate:
                continue
            expected.add((live[scan], "identifier_duplicate"))
    return expected


def duplicate_groups(plan: CampaignPlan) -> dict[str, list[tuple[str, str]]]:
    """Planted byte-copy groups: content -> every ``(source, name)`` carrying its bytes."""
    groups: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for item in plan.main_arrivals:
        groups[item.content].append((item.source, item.name))
    return {key: value for key, value in groups.items() if len(value) > 1}


def category_counts(plan: CampaignPlan) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    for item in plan.candidates:
        counts[item.category.value] += 1
    return dict(counts)


def defect_contents(plan: CampaignPlan, defect: IdDefect) -> list[str]:
    return [item.key for item in plan.contents if item.id_defect is defect]


def chain_contents(plan: CampaignPlan) -> list[str]:
    return [
        item.key for item in plan.contents
        if plan.candidate(item.candidate) is not None
        and plan.candidate(item.candidate).category is Category.CHAIN  # type: ignore[union-attr]
    ]


def summarize_expected(plan: CampaignPlan) -> dict[str, Any]:
    """The ground-truth summary the report prints beside the measured one."""
    main = plan.main_arrivals
    unique = {item.content for item in main}
    final_effective = effective_contents(plan, phase="reopen")
    scripts = {
        key for key in final_effective
        if plan.content(key).kind is not ContentKind.BLANK_PAGE
    }
    per_set = {}
    for code in plan.config.sets:
        results = expected_results(plan, code, "reopen")
        per_set[code] = {
            "candidates": len(results),
            "absent": sum(1 for item in results if item.absent),
            "scored": sum(1 for item in results if not item.absent),
        }
    return {
        "arrivals_written": len(main),
        "late_arrivals": len(plan.late_arrivals),
        "unique_contents": len(unique),
        "byte_copy_arrivals": len(main) - len(unique),
        "registered_sheets": len(registered_contents(plan)),
        "superseded_sheets": len(superseded_contents(plan)),
        "effective_sheets": len(final_effective),
        "effective_scripts": len(scripts),
        "blank_pages": sum(1 for key in final_effective
                           if plan.content(key).kind is ContentKind.BLANK_PAGE),
        "candidates_on_lists": sum(1 for item in plan.candidates if item.on_roster),
        "scored_candidates": sum(item["scored"] for item in per_set.values()),
        "per_set": per_set,
        "categories": category_counts(plan),
    }


def iter_pairs(values: Iterable[Any]) -> list[Any]:
    return list(values)


__all__ = [
    "ABSENT_MARK",
    "ABSENT_RANK",
    "PHASES",
    "CellMismatch",
    "CommittedFacts",
    "ExpectedMark",
    "ExpectedResult",
    "attendance_list",
    "compare_workbook",
    "corrected_contents",
    "duplicate_groups",
    "effective_contents",
    "evaluate_rank",
    "expected_identifier",
    "expected_open_conflicts",
    "expected_results",
    "expected_statuses_after",
    "expected_statuses_before",
    "reference_score",
    "registered_contents",
    "reliable",
    "summarize_expected",
    "superseded_contents",
]
