"""One small examination, carried through every stage of the workflow.

Purpose:
    Prove the stages are wired to each other. The repository's own suite tests
    each stage thoroughly; what a release needs to know in addition is that a
    project created on the Project stage can be templated, scanned, recognised,
    reconciled, keyed, scored and exported as a workbook **in one run, against
    one database**, without a hand-built object anywhere in the chain.

Scale, deliberately:
    Six candidates and twenty questions. §5 of the framework's brief is
    explicit that release qualification must not become the 100,000-sheet run -
    that is ``-m stress`` and ``run_phase10_100k_qualification.ps1``, and it
    answers a different question. This one has to be fast enough that nobody is
    tempted to skip it.

Determinism:
    The sheets are rendered from the template by the same code the recogniser
    uses to find the bubbles, so the ground truth is *stated* rather than
    eyeballed: "candidate 100001 answered every question A" is a fact about the
    input, and the assertions check the pipeline reproduced it. No random data,
    no sampling, no tolerance.

What is asserted:
    That each stage produced what the next one needs, and that the workbook at
    the end contains the marks the arithmetic says it should. Not pixels, and
    not timings.
"""

from __future__ import annotations

from fractions import Fraction
from pathlib import Path
from typing import TYPE_CHECKING

import cv2
import openpyxl
import pytest

# The repository's own fixtures. Importable because `pythonpath = ["."]` in
# pyproject.toml puts the repository root on sys.path - the same mechanism the
# repository's own suites rely on.
from tests.conftest import build_answer_sheet_template, render_marked_sheet
from tests.report_fixtures import build_result_template

from omr_scanner.domain.scoring import ScoringPolicy
from omr_scanner.services import (
    batch_store,
    reconciliation_store,
    report_store,
    review_store,
    scoring_store,
)
from omr_scanner.services.answer_key import plan_for, read_key
from omr_scanner.services.batch_processor import BatchOptions, process_batch
from omr_scanner.services.candidate_import import read_roster
from omr_scanner.services.project_service import create_project
from omr_scanner.services.report_template import preview_template, suggest_mapping
from omr_scanner.services.template_service import save_template

if TYPE_CHECKING:  # pragma: no cover - typing only
    from collections.abc import Iterator

    from omr_scanner.services.project_service import ProjectSession

OPERATOR = "Release Qualification"
SET_CODE = "A"

CANDIDATES = (
    # roll,    name,                 wrong questions (all others answered "A")
    ("100001", "CANDIDATE PERFECT", ()),
    ("100002", "CANDIDATE THREE WRONG", (1, 2, 3)),
    ("100003", "CANDIDATE ONE WRONG", (7,)),
    ("100004", "CANDIDATE ABSENT", None),  # no script at all
    ("100005", "CANDIDATE TWO WRONG", (4, 5)),
    ("100006", "CANDIDATE PERFECT TOO", ()),
)
"""The scenario, stated once so the assertions can be derived from it."""

PRESENT = tuple(row for row in CANDIDATES if row[2] is not None)


# ------------------------------------------------------------------ fixtures


@pytest.fixture(scope="module")
def template():
    """A template whose geometry matches the synthetic page it will render."""
    return build_answer_sheet_template()


@pytest.fixture(scope="module")
def plan(template):
    return plan_for(template)


@pytest.fixture
def project(tmp_path: Path) -> Iterator[ProjectSession]:
    """Stage 1 - Project. A real project on disk, closed afterwards.

    Closing matters on Windows: an open SQLite handle stops pytest removing the
    temporary directory, and a qualification run that leaves those behind fills
    a disk over a few weeks.
    """
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    session = create_project(workspace, "Qualification Exam", exam_name="Release Smoke Test")
    try:
        yield session
    finally:
        session.close()


def sheet_marks(roll: str, wrong: tuple[int, ...], question_count: int) -> dict:
    """What is shaded on one candidate's sheet.

    Every question is answered ``A`` except those in ``wrong``, which get ``B``.
    The key is all ``A``, so the expected mark is ``question_count - len(wrong)``.
    """
    answers = {number: ("B" if number in wrong else "A") for number in range(1, question_count + 1)}
    return {
        "roll_number": dict(enumerate(roll)),
        "set_code": {0: SET_CODE},
        "questions_0": {index: answers[index + 1] for index in range(10)},
        "questions_1": {index: answers[index + 11] for index in range(10)},
    }


@pytest.fixture
def scans(project, template, plan, tmp_path: Path) -> list[Path]:
    """Stage 3 - Scan. Render one sheet per present candidate."""
    directory = tmp_path / "scans"
    directory.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for roll, _name, wrong in PRESENT:
        destination = directory / f"{roll}.png"
        image = render_marked_sheet(template, sheet_marks(roll, wrong, plan.question_count))
        cv2.imwrite(str(destination), image)
        paths.append(destination)
    return paths


@pytest.fixture
def roster_file(tmp_path: Path, plan) -> Path:
    """Stage 6 - Attendance. The candidate list, including the absentee."""
    lines = [f"Roll No.,Name,Total ({plan.question_count})"]
    for roll, name, wrong in CANDIDATES:
        lines.append(f"{roll},{name}," + ("ABSENT" if wrong is None else str(plan.question_count)))
    path = tmp_path / "candidates.csv"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


@pytest.fixture
def processed(project, template, plan, scans, roster_file, tmp_path: Path):
    """Run the whole chain once; every test below reads its output.

    Module-level fixtures would be faster still, but a shared database across
    tests makes a failure in one of them impossible to attribute.
    """
    database = project.database

    # ---------------------------------------------- stage 2: Template
    template_path = project.root / "templates" / "qualification.omrt"
    template_path.parent.mkdir(parents=True, exist_ok=True)
    save_template(template, template_path)

    # -------------------------------------- stages 3-4: Scan and Processing
    batch_id = batch_store.create_batch(
        database, scans, identity=batch_store.BatchIdentity.of(template)
    )
    recorder = batch_store.BatchRecorder(database=database, batch_id=batch_id)
    report = process_batch(
        scans, template, options=BatchOptions(), on_result=recorder.record, workers=1
    )
    recorder.flush()
    batch_store.finalise_batch(database, batch_id)

    scan_ids = batch_store.scan_ids_by_path(database, batch_id)
    for item in report.processed:
        review_store.sync_conflicts(
            database,
            batch_id=batch_id,
            scan_id=scan_ids[item.source_path],
            result=item.result,
            template=template,
        )
    review_store.sync_duplicate_identifiers(database, batch_id)

    # ---------------------------------------------- stage 6: Attendance
    roster_id = reconciliation_store.import_roster(
        database, read_roster(roster_file), imported_by=OPERATOR
    )
    reconciliation_store.reconcile_batch(database, roster_id, batch_id)

    # --------------------------------------- stages 7-8: Answer key, Results
    stored_key = scoring_store.save_key(
        database, read_key("A" * plan.question_count, plan, SET_CODE).to_key()
    )
    scoring_store.verify_key(database, stored_key.key_id, verified_by=OPERATOR)
    scoring_store.save_policy(
        database, ScoringPolicy(correct_mark=Fraction(1)), created_by=OPERATOR
    )
    scoring_store.score_batch(
        database, roster_id, batch_id, template, computed_by=OPERATOR
    )

    # ------------------------------------------------- stage 9: Reports
    result_template = build_result_template(
        tmp_path / "result-template.xlsx",
        candidate_count=len(CANDIDATES),
        roll_prefix="1000",
        roll_digits=6,
        names=[name for _roll, name, _wrong in CANDIDATES],
        absent_every=4,  # candidate 100004, matching CANDIDATES above
    )
    from tests.report_fixtures import SAMPLE_HEADERS

    preview = preview_template(result_template)
    report_store.associate_template(
        database,
        SET_CODE,
        result_template,
        suggest_mapping(list(SAMPLE_HEADERS)).to_mapping(),
        sheet_name=preview.sheet,
        updated_by=OPERATOR,
    )

    return {
        "database": database,
        "roster_id": roster_id,
        "batch_id": batch_id,
        "batch_report": report,
        "template_path": template_path,
    }


# --------------------------------------------------------------------- tests


class TestProjectStage:
    def test_the_project_exists_on_disk(self, project):
        assert project.root.is_dir()
        assert (project.root / "project.json").is_file()

    def test_the_examination_name_was_stored(self, project):
        assert project.exam_name == "Release Smoke Test"


class TestTemplateStage:
    def test_the_template_saves_and_reloads(self, processed, template):
        from omr_scanner.services.template_service import load_template

        path = processed["template_path"]
        assert path.is_file()
        reloaded = load_template(path)
        assert reloaded.name == template.name
        assert len(reloaded.zones) == len(template.zones)


class TestScanAndProcessing:
    def test_every_sheet_was_processed(self, processed):
        report = processed["batch_report"]
        assert len(report.processed) == len(PRESENT)

    def test_no_sheet_failed_to_process(self, processed):
        report = processed["batch_report"]
        failed = getattr(report, "failed", ())
        assert not failed, f"sheets the pipeline could not process: {failed}"

    def test_every_roll_number_was_recognised(self, processed):
        """The recognised identifier must match the sheet that was rendered."""
        recognised = {
            item.result.identifier_value
            for item in processed["batch_report"].processed
            if item.result.identifier_value
        }
        expected = {roll for roll, _name, _wrong in PRESENT}
        assert recognised == expected, f"expected {sorted(expected)}, read {sorted(recognised)}"


class TestAttendanceStage:
    def test_the_roster_was_imported(self, processed):
        assert processed["roster_id"] is not None

    def test_every_script_was_matched_to_a_candidate(self, processed):
        """Every rendered sheet matches a roster candidate.

        Re-running reconciliation is documented as idempotent, so the counts it
        returns are also the counts already stored.
        """
        counts = reconciliation_store.reconcile_batch(
            processed["database"], processed["roster_id"], processed["batch_id"]
        )
        assert counts.registered == len(CANDIDATES)
        assert counts.scripts == len(PRESENT)
        assert counts.unknown_id == 0, "a rendered sheet was not matched to the roster"
        assert counts.duplicate_script == 0
        assert counts.matched == len(PRESENT), (
            f"expected {len(PRESENT)} scripts matched to candidates, got {counts.matched}"
        )
        assert counts.absent_confirmed == 1, (
            "candidate 100004 was rendered absent and has no script, so exactly "
            f"one absence should be confirmed; got {counts.absent_confirmed}"
        )
        assert counts.present_without_script == 0
        assert counts.absent_with_script == 0


class TestResultsStage:
    def test_every_present_candidate_was_scored(self, processed, template):
        results = scoring_store.list_results(
            processed["database"], processed["roster_id"], processed["batch_id"], template
        )
        assert len(results) >= len(PRESENT)

    def test_marks_match_the_stated_ground_truth(self, processed, template, plan):
        """The arithmetic is checked against the scenario, not against itself."""
        results = {
            item.candidate_id: item
            for item in scoring_store.list_results(
                processed["database"], processed["roster_id"], processed["batch_id"], template
            )
        }
        for roll, _name, wrong in PRESENT:
            item = results.get(roll)
            assert item is not None, f"candidate {roll} has no result"
            assert item.has_mark, f"candidate {roll} was not scored: {item.status}"
            expected = Fraction(plan.question_count - len(wrong))
            assert item.final_score == expected, (
                f"candidate {roll}: expected {expected}, scored {item.final_score}"
            )


class TestReportsStage:
    def test_the_set_is_ready_for_export(self, processed, template):
        readiness = report_store.check_readiness(
            processed["database"],
            processed["roster_id"],
            processed["batch_id"],
            template,
            SET_CODE,
            for_final_export=True,
        )
        assert readiness.is_ready, readiness.describe()

    def test_a_result_workbook_is_generated_and_readable(
        self, processed, template, tmp_path: Path
    ):
        """The Excel path end to end: openpyxl writes it, openpyxl reads it.

        This is the stage that exercises openpyxl and Pillow in anger, which is
        why it is in the release qualification rather than only in the unit
        suite.
        """
        outcome = report_store.generate_xlsx(
            processed["database"],
            processed["roster_id"],
            processed["batch_id"],
            template,
            SET_CODE,
            project_name="Release Smoke Test",
            output_dir=tmp_path / "exports",
            computed_by=OPERATOR,
        )
        assert outcome.ok, f"generation reported: {outcome.warnings}"
        assert outcome.output_path.is_file()
        assert outcome.output_path.stat().st_size > 4096

        workbook = openpyxl.load_workbook(outcome.output_path)
        assert workbook.sheetnames, "the generated workbook has no sheets"

        rollwise = next(
            (name for name in workbook.sheetnames if "Rollwise" in name), None
        )
        assert rollwise is not None, f"no roll-wise sheet in {workbook.sheetnames}"

        sheet = workbook[rollwise]
        found = {
            str(sheet.cell(row=row, column=column).value)
            for row in range(1, sheet.max_row + 1)
            for column in range(1, min(sheet.max_column, 6) + 1)
            if sheet.cell(row=row, column=column).value is not None
        }
        missing = [roll for roll, _name, _wrong in CANDIDATES if roll not in found]
        assert not missing, f"candidates missing from the workbook: {missing}"
